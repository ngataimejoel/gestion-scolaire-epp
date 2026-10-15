"""Reprise du prototype AKWA SCHOOL v6 : référentiels, codes DREN / IEPP, import / export du fichier national
ELEVES (.xls lu sans bibliothèque), rapports officiels sur une seule page, logos sous l'en-tête, rapport CM2.

Toutes les données d'élèves de ces tests sont inventées.
"""
from __future__ import annotations

import base64
import csv
import io
import shutil
import subprocess
from pathlib import Path
from types import SimpleNamespace

import docx
import openpyxl
import pytest
from pypdf import PdfReader

from epp.core.security import system_session
from epp.domain import fiche_scolaire, referentiels, repo
from epp.io import import_eleves19 as ie
from epp.io.lecture_xls import FichierXlsInvalide, lire_xls
from epp.reports import engine as E
from epp.reports import render_pdf, service, store
from epp.reports.builtin import TEMPLATES

S = system_session()
SOFFICE = shutil.which("soffice")

# ------------------------------------------------------------------ données fictives (format ELEVES, 22 colonnes)
LIGNES = [
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "12345678A", "KOFFI", "Ama Bintou", "FEMININ", "3", "7",
     "2016", "IVOIRIENNE", "OUANINOU", "SP000725", "OUANINOU", "KOFFI YAO PAUL", "AKA AYA", "245", "OUANINOU",
     "", "CE1"],
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "23456789B", "TRAORE", "Moussa", "MASCULIN", "14", "11",
     "2015", "Burkinabè", "AUCUN", "SP000001", "AUCUN", "AUCUN", "SANOU MARIAM", "0", "AUCUN", "", "CE2"],
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "34567890C", "DIOMANDE", "Fanta", "FEMININ", "1", "1",
     "2014", "IVOIRIENNE", "TOUBA", "", "Sous-préfecture de Touba", "DIOMANDE LAMINE", "", "1088", "TOUBA",
     "", "CM2"],
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "45678901D", "YAO", "Kouassi", "GARCON", "5", "5", "2019",
     "", "", "", "", "", "", "", "", "", "CP3"],                                            # classe inexistante
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "12345678A", "KOFFI", "Ama Bintou", "FEMININ", "3", "7",
     "2016", "IVOIRIENNE", "", "", "", "", "", "", "", "", "CE1"],                          # matricule répété
    ["35", "000251", "OUANINOU", "E000001", "EPP TEST", "56789012E", "BAMBA", "Issa", "INCONNU", "5", "5", "2017",
     "", "", "", "", "", "", "", "", "", "CP2"],                                            # sexe illisible
]


def _ecrire_csv(chemin: Path, lignes, sep=";"):
    with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f, delimiter=sep).writerows([ie.COLONNES] + lignes)
    return chemin


def _vers_xls(csv_path: Path, dossier: Path) -> Path:
    """Convertit un CSV (UTF-8, ;) en .xls Excel 97-2003 avec LibreOffice ; toutes les colonnes en texte."""
    if not SOFFICE:
        pytest.skip("LibreOffice (soffice) absent : conversion .xls impossible")
    filtre = "Text - txt - csv (StarCalc):59,34,76,1," + "/".join(f"{i}/2" for i in range(1, 23))
    r = subprocess.run([SOFFICE, "--headless", "--infilter=" + filtre, "--convert-to", "xls:MS Excel 97",
                        "--outdir", str(dossier), str(csv_path)], capture_output=True, timeout=180)
    out = dossier / (csv_path.stem + ".xls")
    if not out.exists():
        pytest.skip(f"Conversion .xls impossible : {r.stderr[:200]!r}")
    return out


# =================================================================== 1. référentiels
def test_referentiels_sous_prefectures_et_nationalites():
    sp = referentiels.sous_prefectures()
    nat = referentiels.nationalites()
    assert len(sp) == 519 and ("SP000725", "OUANINOU") in sp and ("SP000726", "TOUBA") in sp
    assert len(nat) == 94 and nat[0] == ("70", "IVOIRIENNE")
    assert referentiels.code_sp("Sous-préfecture de Touba") == "SP000726"
    assert referentiels.code_sp("  ouaninou ") == "SP000725"
    assert referentiels.code_sp("hors cote d'ivoire") == "SP000998"
    assert referentiels.code_sp("sp000725") == "SP000725" and referentiels.code_sp("SP999999") == ""
    assert referentiels.code_sp("Atlantide") == "" and referentiels.code_sp("") == ""
    assert referentiels.libelle_sp("SP000725") == "OUANINOU"
    assert referentiels.nationalite_officielle("Burkinabè") == "BURKINABE"
    assert referentiels.nationalite_officielle("ivoirienne") == "IVOIRIENNE"
    assert referentiels.nationalite_officielle("75") == "BURKINABE"
    assert referentiels.nationalite_officielle("Martienne") == "MARTIENNE"
    assert referentiels.nationalite_officielle("") == ""
    assert referentiels.code_nationalite("Malienne") == "79"


def test_referentiels_livres_avec_l_application():
    racine = Path(__file__).resolve().parents[1]
    for nom in ("sous_prefectures_ci.csv", "nationalites.csv"):
        assert (racine / "epp" / "resources" / "referentiels" / nom).exists()
    assert '"epp" / "resources"' in (racine / "packaging" / "gestion_scolaire.spec").read_text(encoding="utf-8")


# =================================================================== 2. codes DREN / IEPP
def test_codes_dren_iepp_parametres_qr_et_variables(demo):
    db, annee = demo
    e = repo.etablissement(db)
    assert e["code_dren"] == "35" and e["code_iepp"] == "000251"
    repo.save_etablissement(db, S, {"code_dren": "36", "code_iepp": "000999", "ecole_code": "E000001"})
    f = fiche_scolaire.fiches(db, annee, annee.codes[0])[0]
    assert "IEPP: IEPP OUANINOU (code 000999)" in f["qr"] and "DREN: DRENAET DE TOUBA (code 36)" in f["qr"]
    assert "(code E000001)" in f["qr"]
    v = E.variables(db, annee, {})
    assert v["code_dren"] == "36" and v["code_iepp"] == "000999"


def test_base_existante_recoit_les_codes(tmp_path):
    from epp.core import db as dbmod
    path = tmp_path / "ancienne.sqlite"
    d = dbmod.open_database(path)
    d.execute("DELETE FROM parametre WHERE cle IN ('code_dren', 'code_iepp')")
    d.close()
    d = dbmod.open_database(path)
    assert repo.etablissement(d)["code_iepp"] == "000251" and int(d.meta("schema_version")) == dbmod.SCHEMA_VERSION
    d.close()


# =================================================================== 3. lecture .xls / import / export
def test_lecture_xls_biff8(tmp_path):
    """Classeur .xls produit par LibreOffice : chaînes partagées longues (SST + CONTINUE), accents,
    nombres entiers et décimaux, cellules vides."""
    lignes = [["Libellé", "Nombre", "Décimal", "Vide", "Texte long"]]
    for i in range(400):
        lignes.append([f"Élève n°{i} – Ünïcødé ★", str(i * 7), f"{i}.25", "",
                       f"chaîne longue numéro {i:04d} " + "x" * 30])
    src = tmp_path / "essai.csv"
    with open(src, "w", encoding="utf-8", newline="") as f:
        csv.writer(f, delimiter=";").writerows(lignes)
    if not SOFFICE:
        pytest.skip("LibreOffice (soffice) absent")
    subprocess.run([SOFFICE, "--headless", "--infilter=Text - txt - csv (StarCalc):59,34,76", "--convert-to",
                    "xls:MS Excel 97", "--outdir", str(tmp_path), str(src)], capture_output=True, timeout=180)
    xls = tmp_path / "essai.xls"
    if not xls.exists():
        pytest.skip("conversion .xls impossible")
    rows = lire_xls(xls)
    assert len(rows) == 401
    assert rows[0][:3] == ["Libellé", "Nombre", "Décimal"]
    assert rows[1][0] == "Élève n°0 – Ünïcødé ★" and rows[400][0] == "Élève n°399 – Ünïcødé ★"
    assert rows[123][1] == 122 * 7 and isinstance(rows[123][1], int)
    assert rows[10][2] == pytest.approx(9.25)
    assert rows[400][4].startswith("chaîne longue numéro 0399")
    assert all(r[3] == "" for r in rows[1:])
    assert ie.lire_fichier(xls)[5][1] == "28"
    with pytest.raises(FichierXlsInvalide):
        lire_xls(b"pas un classeur")


def test_classe_du_niveau():
    a = SimpleNamespace(codes=["CP1", "CE1 A", "CE1 B", "CM2"])
    assert ie.classe_du_niveau(a, "cp1") == "CP1"
    assert ie.classe_du_niveau(a, "CE1") == "CE1 A"
    assert ie.classe_du_niveau(a, "CE1", {"CE1 A": 40, "CE1 B": 31}) == "CE1 B"     # la moins chargée
    assert ie.classe_du_niveau(a, "CE2") is None and ie.classe_du_niveau(a, "") is None


@pytest.mark.parametrize("fmt", ["xls", "xlsx", "csv", "csv,"])
def test_import_fichier_national(demo, tmp_path, fmt):
    db, annee = demo
    avant = db.scalar("SELECT COUNT(*) FROM inscription WHERE annee_id=?", (annee.id,))
    src = _ecrire_csv(tmp_path / "eleves.csv", LIGNES, sep="," if fmt == "csv," else ";")
    if fmt == "xls":
        src = _vers_xls(src, tmp_path)
        assert src.read_bytes()[:4] == b"\xD0\xCF\x11\xE0"
    elif fmt == "xlsx":
        wb = openpyxl.Workbook()
        for r in [ie.COLONNES] + LIGNES:
            wb.active.append([int(x) if x.isdigit() and len(x) < 5 else x for x in r])   # nombres comme Excel
        src = tmp_path / "eleves.xlsx"
        wb.save(src)
    apercu = ie.analyser(db, annee, src)
    assert db.scalar("SELECT COUNT(*) FROM inscription WHERE annee_id=?", (annee.id,)) == avant   # rien d'écrit
    b = ie.importer(db, S, annee, src)
    assert (b["lignes"], b["importes"], b["reinscrits"], b["doublons"], b["erreurs"]) == (6, 3, 0, 1, 2)
    assert apercu["a_importer"] == 3 and b["sauvegarde"] and Path(b["sauvegarde"]).exists()
    etats = {d["ligne"]: (d["etat"], d["motif"]) for d in b["details"]}
    assert etats[5][0] == "erreur" and "CP3" in etats[5][1]
    assert etats[6] == ("doublon", "matricule en double dans le fichier")
    assert etats[7][0] == "erreur" and "sexe" in etats[7][1]
    assert any("EPP TEST" in a for a in b["avertissements"])          # autre école que celle des paramètres
    assert db.scalar("SELECT COUNT(*) FROM inscription WHERE annee_id=?", (annee.id,)) == avant + 3
    rows = {r["matricule_desps"]: r for r in repo.inscriptions(db, annee) if r["matricule_desps"]}
    a, t, d = rows["12345678A"], rows["23456789B"], rows["34567890C"]
    assert a["nom"] == "KOFFI AMA BINTOU" and a["sexe"] == "F" and a["date_naissance"] == "2016-07-03"
    assert a["classe"] == "CE1" and a["extrait"] == "OUI" and a["acte_num"] == "245"
    assert a["centre_etat_civil"] == "OUANINOU" and a["sous_prefecture"] == "OUANINOU"
    assert a["pere_nom"] == "KOFFI YAO PAUL" and a["mere_nom"] == "AKA AYA" and a["nationalite"] == "IVOIRIENNE"
    assert t["sexe"] == "M" and t["nationalite"] == "BURKINABE" and t["extrait"] == "NON"
    assert t["acte_num"] is None and t["localite"] is None and t["pere_nom"] is None and t["sous_prefecture"] is None
    assert d["sous_prefecture"] == "TOUBA" and d["classe"] == "CM2" and d["date_naissance"] == "2014-01-01"
    assert db.scalar("SELECT COUNT(*) FROM journal WHERE action='IMPORT_ELEVES_NATIONAL'") == 1
    # second import du même fichier : tout est doublon ou erreur, rien n'est ajouté
    b2 = ie.importer(db, S, annee, src)
    assert b2["importes"] == 0 and b2["doublons"] == 4 and b2["erreurs"] == 2
    assert "doublon" in ie.resume(b2)


def test_import_reinscrit_un_eleve_connu(demo, tmp_path):
    db, annee = demo
    ie.importer(db, S, annee, _ecrire_csv(tmp_path / "a.csv", LIGNES[:1]))
    eid = db.scalar("SELECT id FROM eleve WHERE matricule_desps='12345678A'")
    db.execute("DELETE FROM inscription WHERE eleve_id=?", (eid,))       # connu de la base, pas inscrit cette année
    b = ie.importer(db, S, annee, _ecrire_csv(tmp_path / "b.csv", LIGNES[:1]))
    assert b["reinscrits"] == 1 and b["importes"] == 0
    assert db.scalar("SELECT type FROM inscription WHERE eleve_id=?", (eid,)) == "REINSCRIPTION"
    assert db.scalar("SELECT COUNT(*) FROM eleve WHERE matricule_desps='12345678A'") == 1


def test_import_colonnes_manquantes(demo, tmp_path):
    db, annee = demo
    p = tmp_path / "x.csv"
    p.write_text("nom;sexe\nKONE;M\n", encoding="utf-8")
    with pytest.raises(repo.ValidationError, match="prenoms"):
        ie.analyser(db, annee, p)


def test_export_et_modele_format_national(demo, tmp_path):
    db, annee = demo
    ie.importer(db, S, annee, _ecrire_csv(tmp_path / "e.csv", LIGNES))
    repo.save_etablissement(db, S, {"ecole_code": "E000001"})
    out = ie.exporter(db, S, annee, tmp_path / "export.csv")
    with open(out, encoding="utf-8-sig", newline="") as f:
        lignes = list(csv.reader(f, delimiter=";"))
    assert lignes[0] == ie.COLONNES and len(lignes[0]) == 22
    presents = repo.inscriptions(db, annee, statut="PRESENT")
    assert len(lignes) == len(presents) + 1
    par_mat = {l[5]: l for l in lignes[1:]}
    a = par_mat["12345678A"]
    assert a[:5] == ["35", "000251", "OUANINOU", "E000001", "EPP OUANINOU 2"]
    assert a[6:22] == ["KOFFI", "AMA BINTOU", "FEMININ", "3", "7", "2016", "IVOIRIENNE", "OUANINOU", "SP000725",
                       "OUANINOU", "KOFFI YAO PAUL", "AKA AYA", "245", "OUANINOU", "", "CE1"]
    t = par_mat["23456789B"]
    assert t[8] == "MASCULIN" and t[12] == "BURKINABE" and t[18] == "0" and t[19] == "" and t[21] == "CE2"
    assert par_mat["34567890C"][14:16] == ["SP000726", "TOUBA"]
    # l'export se relit tel quel (même format) : dans une autre année, tout serait importable
    assert ie.analyser(db, annee, out)["doublons"] == len(presents)
    # Excel
    x = ie.exporter(db, S, annee, tmp_path / "export.xlsx")
    ws = openpyxl.load_workbook(x).active
    assert [c.value for c in ws[1]] == ie.COLONNES and ws.max_row == len(presents) + 1
    assert ie.lire_fichier(x)[1:] == [l[:len(ie.lire_fichier(x)[i + 1])] for i, l in enumerate(lignes[1:])]
    # modèles vides
    m = ie.modele(tmp_path / "modele.csv")
    assert ie.lire_fichier(m) == [ie.COLONNES]
    m2 = ie.modele(tmp_path / "modele.xlsx", db, exemple=True)
    ex = ie.lire_fichier(m2)
    assert ex[0] == ie.COLONNES and ex[1][:2] == ["35", "000251"] and ex[1][21] == "CE1"


# =================================================================== 4. une seule page
def _doc(db, annee, code, params=None):
    mid = db.scalar("SELECT id FROM modele_rapport WHERE code=?", (code,))
    return E.build_document(db, annee, store.definition(db, mid), params or {})


def test_rapports_officiels_sur_une_page(demo, tmp_path):
    db, annee = demo
    defs = {code: d for code, _, _, _, d in TEMPLATES}
    for code in ("RENTREE", "COMPOSITION", "SYNTHESE", "RESULTATS", "RESULTATS_CM2", "EFFECTIFS"):
        assert defs[code]["page"] == {"taille": "A4", "orientation": "paysage", "une_page": True}, code
    assert defs["MENSUEL"]["page"] == {"taille": "A4", "orientation": "portrait"}
    assert "une_page" not in defs["BULLETIN_ANNUEL"]["page"] and "une_page" not in defs["REGISTRE"]["page"]
    # la démonstration : rapport de rentrée et de composition trop longs pour une page -> réduits
    for code, params in (("RENTREE", {}), ("COMPOSITION", {"compo": 1})):
        doc = _doc(db, annee, code, params)
        ech = render_pdf.mise_a_l_echelle(doc)
        assert len(ech) == 1 and 0.4 <= ech[0] < 1, (code, ech)
        assert len(PdfReader(service.generer(db, S, annee, code, params, tmp_path / f"{code}.pdf")).pages) == 1
        doc.une_page = False
        assert len(PdfReader(io.BytesIO(render_pdf.render_bytes(doc))).pages) == 2     # sans réduction
    # une page par classe : chaque partie tient sur sa page
    r = PdfReader(service.generer(db, S, annee, "RESULTATS", {}, tmp_path / "r.pdf"))
    assert len(r.pages) == len([c for c in annee.classes if c.regle_mga == "STANDARD"])
    # le rapport mensuel reste en portrait, taille normale
    m = PdfReader(service.generer(db, S, annee, "MENSUEL", {"mois": 10}, tmp_path / "m.pdf"))
    assert m.pages[0].mediabox.width < m.pages[0].mediabox.height
    assert render_pdf.mise_a_l_echelle(_doc(db, annee, "MENSUEL", {"mois": 10})) == []
    # bulletins : une page par élève, jamais réduits
    doc = _doc(db, annee, "BULLETIN_ANNUEL", {"classe": "CM2"})
    doc.une_page = True
    assert render_pdf.mise_a_l_echelle(doc) == []
    n = len(repo.inscriptions(db, annee, "CM2", statut="PRESENT"))
    assert len(PdfReader(io.BytesIO(render_pdf.render_bytes(doc))).pages) == n
    # Word et Excel restent disponibles ; Excel : impression sur une page
    assert docx.Document(service.generer(db, S, annee, "RENTREE", {}, tmp_path / "r.docx")).tables
    ws = openpyxl.load_workbook(service.generer(db, S, annee, "RENTREE", {}, tmp_path / "r.xlsx")).active
    assert ws.page_setup.fitToHeight == 1 and ws.page_setup.orientation == "landscape"


def test_partie_trop_longue_garde_sa_taille(demo):
    db, annee = demo
    doc = _doc(db, annee, "REGISTRE")
    doc.blocks = doc.blocks * 6                      # environ 18 pages : illisible si réduit sur une page
    doc.une_page = True
    assert render_pdf.mise_a_l_echelle(doc) == [1.0]
    assert len(PdfReader(io.BytesIO(render_pdf.render_bytes(doc))).pages) > 3


# =================================================================== 5. logos sous l'en-tête
def _png(couleur) -> bytes:
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (120, 80), couleur).save(b, "PNG")
    return b.getvalue()


def _poser_logos(db):
    db.set_param("logo_iepp", base64.b64encode(_png("red")).decode())
    db.set_param("logo_ecole", base64.b64encode(_png("blue")).decode())


def test_logos_sous_les_lignes_de_gauche(demo, tmp_path):
    from reportlab.lib.units import mm

    from epp.reports.fiche_pdf import dessiner_entete
    db, annee = demo
    _poser_logos(db)
    gauche, droite, logos = E.entete_officiel(_doc(db, annee, "EFFECTIFS"))
    assert len(logos) == 2

    class Canevas:
        def __init__(self):
            self.images, self.textes = [], []

        def drawImage(self, img, x, y, w, h, **k):
            self.images.append((x, y, w, h))

        def drawString(self, x, y, t):
            self.textes.append((x, y, t))

        def __getattr__(self, nom):
            return lambda *a, **k: None

    c = Canevas()
    W = 270 * mm
    hauteur = dessiner_entete(c, 0, 0, W, gauche, droite, logos, ("Helvetica", "Helvetica-Bold"))
    assert hauteur == dessiner_entete(None, 0, 0, W, gauche, droite, logos, ("Helvetica", "Helvetica-Bold"))
    bas_texte = min(y for _, y, _ in c.textes)
    assert len(c.images) == 2
    for x, y, w, h in c.images:
        assert x + w < W * 0.56 and x > 0          # colonne de gauche, décalés vers la droite
        assert y + h <= bas_texte                   # sous les lignes de l'en-tête
    assert c.images[0][0] < c.images[1][0]          # IEPP puis école
    assert hauteur > dessiner_entete(None, 0, 0, W, gauche, droite, [], ("Helvetica", "Helvetica-Bold"))
    # PDF, Word (cellule de gauche) et Excel
    assert len(PdfReader(service.generer(db, S, annee, "EFFECTIFS", {}, tmp_path / "e.pdf")).pages[0].images) == 2
    d = docx.Document(service.generer(db, S, annee, "EFFECTIFS", {}, tmp_path / "e.docx"))
    ent = d.tables[0]
    assert ent.cell(0, 0)._tc.xml.count("<pic:pic") == 2 and "<pic:pic" not in ent.cell(0, 2)._tc.xml
    ws = openpyxl.load_workbook(service.generer(db, S, annee, "EFFECTIFS", {}, tmp_path / "e.xlsx")).active
    assert len(ws._images) == 2


# =================================================================== 6. rapport CM2 et trois meilleurs
def test_rapport_cm2_recapitulatif_observations_signature(demo, tmp_path):
    db, annee = demo
    from epp.domain import resultats
    res = [r for r in resultats.compute(db, annee) if r["classe"] == "CM2"]
    admis = sum(1 for r in res if r["decision"] == "ADMIS")
    presents = sum(1 for r in res if r["statut"] == "PRESENT")
    doc = _doc(db, annee, "RESULTATS_CM2")
    recap = next(b for b in doc.blocks if isinstance(b, E.Table) and b.title == "RÉCAPITULATIF")
    lignes = {r.cells[0]: r.cells[1:] for r in recap.rows}
    assert lignes["ADMIS EN 6ème"][2] == admis and lignes["PRÉSENTS"][2] == presents
    assert any(isinstance(b, E.TextZone) for b in doc.blocks) and any(isinstance(b, E.Signature) for b in doc.blocks)
    t = " ".join(PdfReader(service.generer(db, S, annee, "RESULTATS_CM2", {}, tmp_path / "c.pdf"))
                 .pages[0].extract_text().split())
    for attendu in ("RÉSULTATS DES COMPOSITIONS CM2", "RÉCAPITULATIF", "ADMIS EN 6ème", "TAUX D'ADMISSION",
                    "OBSERVATIONS", "LE DIRECTEUR"):
        assert attendu in t, attendu


def test_trois_meilleurs_par_niveau(demo):
    db, annee = demo
    from epp.domain import stats
    par_classe = stats.classement(db, annee, 1, 3)
    par_niveau = stats.classement(db, annee, 1, 3, par_niveau=True)
    assert set(par_niveau) == set(annee.codes)      # une classe par niveau dans la démonstration
    for code in annee.codes:
        assert [x["nom"] for x in par_niveau[code]] == [x["nom"] for x in par_classe[code]]
        moys = [x["moyenne"] for x in par_niveau[code]]
        assert moys == sorted(moys, reverse=True)
    doc = _doc(db, annee, "COMPOSITION", {"compo": 1})
    top = next(b for b in doc.blocks if isinstance(b, E.Table) and "TROIS MEILLEURS" in b.title)
    assert [r.cells[0] for r in top.rows] == list(annee.codes)
    assert top.rows[0].cells[1] == par_niveau[annee.codes[0]][0]["nom"]

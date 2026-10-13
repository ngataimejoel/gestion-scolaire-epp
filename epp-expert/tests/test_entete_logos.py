"""En-tête d'impression modifiable et logos de l'IEPP / de l'école (Paramètres)."""
from __future__ import annotations

import openpyxl
import pytest
from pypdf import PdfReader

from epp.core import seed
from epp.core.security import AccessDenied, Session, system_session
from epp.domain import repo
from epp.reports import service


def _png(w, h, couleur=0x0E6646, alpha=False) -> bytes:
    from PySide6.QtCore import QBuffer, QIODevice
    from PySide6.QtGui import QImage
    img = QImage(w, h, QImage.Format_ARGB32 if alpha else QImage.Format_RGB32)
    img.fill(0x00000000 if alpha else couleur)
    buf = QBuffer()
    buf.open(QIODevice.WriteOnly)
    img.save(buf, "PNG")
    return bytes(buf.data())


def _texte(pdf) -> str:
    return PdfReader(pdf).pages[0].extract_text()


def test_entete_par_defaut(demo, tmp_path):
    db, annee = demo
    e = repo.etablissement(db)
    assert e["entete_gauche"] == seed.ENTETE_GAUCHE_DEFAUT and e["entete_droite"] == seed.ENTETE_DROITE_DEFAUT
    t = _texte(service.generer(db, system_session(), annee, "EFFECTIFS", {}, tmp_path / "e.pdf"))
    for attendu in ("DRENAET DE TOUBA", "IEPP OUANINOU", "EPP OUANINOU 2", "REPUBLIQUE DE COTE D'IVOIRE",
                    f"ANNEE SCOLAIRE : {annee.libelle}"):
        assert attendu in t


def test_entete_modifie_puis_retabli(demo, tmp_path):
    db, annee = demo
    s = system_session()
    repo.save_etablissement(db, s, {"entete_gauche": "MINISTERE TEST\nINSPECTION {iepp}\n\n{ecole}",
                                    "entete_droite": "REPUBLIQUE TEST\nAnnée {annee}"})
    t = _texte(service.generer(db, s, annee, "SYNTHESE", {}, tmp_path / "a.pdf"))
    assert "MINISTERE TEST" in t and "INSPECTION IEPP OUANINOU" in t and f"Année {annee.libelle}" in t
    assert "DRENAET DE TOUBA" not in t                     # la ligne {dren} a été retirée
    ws = openpyxl.load_workbook(service.generer(db, s, annee, "SYNTHESE", {}, tmp_path / "a.xlsx")).active
    valeurs = {c.value for row in ws.iter_rows(max_row=8) for c in row if c.value}
    assert {"MINISTERE TEST", "INSPECTION IEPP OUANINOU", "REPUBLIQUE TEST", f"Année {annee.libelle}"} <= valeurs
    repo.reset_entete(db, s)
    t = _texte(service.generer(db, s, annee, "SYNTHESE", {}, tmp_path / "b.pdf"))
    assert "MINISTERE TEST" not in t and "DRENAET DE TOUBA" in t


def test_logos_iepp_et_ecole_sur_rapports_et_bulletins(demo, tmp_path):
    db, annee = demo
    s = system_session()
    t0 = PdfReader(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "0.pdf"))
    assert not t0.pages[0].images
    repo.set_logo(db, s, "logo_iepp", _png(1600, 800))
    png = repo.get_logo(db, "logo_iepp")
    from PySide6.QtGui import QImage
    img = QImage.fromData(png)
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and (img.width(), img.height()) == (repo.LOGO_MAX_PX, 200)
    assert len(PdfReader(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "1.pdf")).pages[0].images) == 1
    repo.set_logo(db, s, "logo_ecole", _png(300, 300, alpha=True))
    assert QImage.fromData(repo.get_logo(db, "logo_ecole")).hasAlphaChannel()   # transparence conservée
    assert len(PdfReader(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "2.pdf")).pages[0].images) == 2
    # bulletin : 2 logos + aucune photo
    b = PdfReader(service.generer(db, s, annee, "BULLETIN_ANNUEL", {"classe": annee.codes[0]}, tmp_path / "b.pdf"))
    assert len(b.pages[0].images) == 2
    # Excel
    ws = openpyxl.load_workbook(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "e.xlsx")).active
    assert len(ws._images) == 2
    # retrait
    repo.set_logo(db, s, "logo_iepp", None)
    repo.set_logo(db, s, "logo_ecole", None)
    assert repo.get_logo(db, "logo_iepp") is None
    assert not PdfReader(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "3.pdf")).pages[0].images


def test_logo_controles_et_droits(demo):
    db, annee = demo
    s = system_session()
    with pytest.raises(repo.ValidationError):
        repo.set_logo(db, s, "logo_ecole", b"pas une image")
    with pytest.raises(ValueError):
        repo.set_logo(db, s, "logo_autre", _png(10, 10))
    lecteur = Session(5, "l", "Lecteur", "X", "X", {"parametres.lire"})
    with pytest.raises(AccessDenied):
        repo.set_logo(db, lecteur, "logo_ecole", _png(10, 10))
    with pytest.raises(AccessDenied):
        repo.reset_entete(db, lecteur)
    # le texte brut d'un logo ne peut pas être injecté par l'enregistrement des paramètres
    repo.save_etablissement(db, s, {"logo_ecole": "AAAA", "localite": "OUANINOU"})
    assert repo.get_logo(db, "logo_ecole") is None


def test_base_existante_recoit_les_nouveaux_parametres(tmp_path):
    from epp.core import db as dbmod
    path = tmp_path / "ancienne.sqlite"
    d = dbmod.open_database(path)
    d.set_param("dren", "DRENAET PERSONNALISEE")
    d.execute("DELETE FROM parametre WHERE cle IN ('entete_gauche','entete_droite','logo_iepp','logo_ecole')")
    d.close()
    d = dbmod.open_database(path)
    e = repo.etablissement(d)
    assert e["entete_gauche"] == seed.ENTETE_GAUCHE_DEFAUT and e["logo_iepp"] == ""
    assert e["dren"] == "DRENAET PERSONNALISEE"            # valeurs existantes conservées
    d.close()


def test_libelle_retire_quand_la_variable_est_vide(demo, tmp_path):
    db, annee = demo
    s = system_session()
    t = _texte(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "a.pdf"))
    assert "Code :" not in t                               # pas de code école saisi
    repo.save_etablissement(db, s, {"ecole_code": "123456"})
    t = _texte(service.generer(db, s, annee, "EFFECTIFS", {}, tmp_path / "b.pdf"))
    assert "Code : 123456" in t

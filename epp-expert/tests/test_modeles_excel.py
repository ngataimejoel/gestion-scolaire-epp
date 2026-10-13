"""Import / export par modèles Excel : modèles vierges, export rejouable, analyse ligne par ligne, doublons."""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from epp.core import backup
from epp.core.security import AccessDenied, Session, system_session
from epp.domain import repo
from epp.domain.repo import ValidationError
from epp.io import modeles_excel as mx

S = system_session()


def _remplir(chemin, feuille, lignes):
    """Ajoute des lignes (dict titre de colonne → valeur) sous les en-têtes d'un modèle."""
    wb = load_workbook(chemin)
    ws = wb[feuille]
    titres = [str(c.value or "").replace(" *", "") for c in ws[1]]
    for ligne in lignes:
        r = next(i for i in range(2, ws.max_row + 2) if not any(ws.cell(i, j).value for j in range(1, len(titres) + 1)))
        for k, v in ligne.items():
            ws.cell(r, titres.index(k) + 1, v)
    wb.save(chemin)


@pytest.mark.parametrize("kind", ["eleves", "personnel", "classes", "absences_eleves", "absences_personnel", "notes"])
def test_modele_vierge(demo, tmp_path, kind):
    db, annee = demo
    f = mx.creer_modele(db, annee, kind, tmp_path / f"{kind}.xlsx", num=1 if kind == "notes" else None)
    wb = load_workbook(f)
    assert wb[mx.META].sheet_state == "hidden" and wb[mx.LISTES].sheet_state == "hidden"
    assert wb.active.title not in (mx.META, mx.LISTES)
    meta = {r[0]: r[1] for r in wb[mx.META].iter_rows(values_only=True)}
    assert meta["type"] == kind and meta["annee"] == annee.libelle
    ws = wb.active
    assert any(str(c.value).endswith(" *") for c in ws[1])            # colonnes obligatoires marquées
    assert ws.data_validations.dataValidation                          # listes déroulantes / contrôles
    if kind == "notes":
        assert set(annee.codes) <= set(wb.sheetnames)
        assert ws.cell(2, 1).value                                     # élèves de la classe pré-remplis
    else:
        assert not any(v for row in ws.iter_rows(min_row=2, values_only=True) for v in row)
    # un modèle vierge ne contient rien à importer
    an = mx.analyser(db, annee, f)
    assert not an.erreurs_fichier
    assert not an.importable


@pytest.mark.parametrize("kind", ["eleves", "personnel", "classes", "absences_eleves", "notes"])
def test_export_puis_reimport_sans_erreur(demo, tmp_path, kind):
    db, annee = demo
    num = 1 if kind == "notes" else None
    f = mx.exporter(db, S, annee, kind, tmp_path / "export.xlsx", num=num)
    an = mx.analyser(db, annee, f)
    assert not an.erreurs_fichier
    assert an.lignes and not [(l.numero, l.erreurs) for l in an.lignes if l.erreurs]
    if kind == "absences_eleves":
        assert all(l.action == "ignorer" for l in an.lignes)        # déjà enregistrées : pas de doublon
        return
    assert {l.action for l in an.lignes} == {"modifier"}
    avant = _instantane(db, annee)
    bilan = mx.appliquer(db, S, annee, an)
    assert bilan["enregistrees"] == len(an.lignes) and bilan["refusees"] == 0
    assert _instantane(db, repo.load_annee(db, annee.id)) == avant  # l'aller-retour ne change rien


def _instantane(db, annee):
    return (
        [(r["matricule_ecole"], r["nom"], r["sexe"], r["date_naissance"], r["classe"], r["redoublant"], r["statut"],
          r["pere_tel"]) for r in repo.inscriptions(db, annee)],
        [(p["matricule"], p["nom"], p["fonction"], p["classe_tenue"], p["telephone"]) for p in repo.personnel(db, annee)],
        [(c.code, c.seuil, c.bareme, c.grille_id, c.niveau_suivant) for c in annee.classes],
        [tuple(r) for r in db.query("SELECT inscription_id, matiere_id, valeur FROM note ORDER BY 1, 2")],
        [tuple(r) for r in db.query("SELECT inscription_id, eval_num, present FROM presence_eval ORDER BY 1, 2")],
    )


def test_import_eleves_creation_modification_et_erreurs(demo, tmp_path):
    db, annee = demo
    existant = repo.inscriptions(db, annee, "CP2")[0]
    f = mx.creer_modele(db, annee, "eleves", tmp_path / "e.xlsx")
    _remplir(f, "Élèves", [
        {"Nom et prénoms": "KOFFI Ange Michel", "Sexe": "M", "Date de naissance": date(2019, 5, 3), "Classe": "CP1",
         "Téléphone du père": 707070707},
        {"Matricule école": existant["matricule_ecole"], "Nom et prénoms": existant["nom"], "Sexe": existant["sexe"],
         "Classe": "CP2", "Téléphone de la mère": "0101010101"},
        {"Nom et prénoms": "BAMBA Issa", "Sexe": "M", "Classe": "CM9"},
        {"Nom et prénoms": "DIALLO Fanta", "Sexe": "F", "Date de naissance": "31/02/2018", "Classe": "CE1"},
        {"Nom et prénoms": "KOFFI Ange Michel", "Sexe": "M", "Date de naissance": "03/05/2019", "Classe": "CP1"},
        {"Nom et prénoms": existant["nom"], "Sexe": existant["sexe"], "Date de naissance": existant["date_naissance"],
         "Classe": "CP2"},
        {"Nom et prénoms": "SANS Sexe", "Classe": "CP1"},
    ])
    an = mx.analyser(db, annee, f)
    l = {x.numero: x for x in an.lignes}
    assert l[2].action == "creer" and not l[2].erreurs
    assert l[3].action == "modifier" and not l[3].erreurs
    assert "Classe inconnue : CM9." in l[4].erreurs
    assert any("date invalide" in e for e in l[5].erreurs)
    assert any("Doublon" in e and "ligne 2" in e for e in l[6].erreurs)
    assert any("déjà inscrit" in e and existant["matricule_ecole"] in e for e in l[7].erreurs)
    assert any("Sexe" in e for e in l[8].erreurs)
    assert an.compte() == {"creer": 1, "modifier": 1, "inscrire": 0, "ignorer": 0, "erreurs": 5}

    rapport = mx.ecrire_rapport(an, tmp_path / "rapport.xlsx")
    rws = load_workbook(rapport).active
    assert [c.value for c in rws[1]][:3] == ["Feuille", "Ligne", "Action prévue"]
    assert [r[2] for r in rws.iter_rows(min_row=2, values_only=True)].count("Refusée") == 5

    n = len(repo.inscriptions(db, annee))
    bilan = mx.appliquer(db, S, annee, an)
    assert bilan["enregistrees"] == 2 and bilan["refusees"] == 5
    assert bilan["sauvegarde"] and backup.verifier(bilan["sauvegarde"])
    rows = repo.inscriptions(db, annee)
    assert len(rows) == n + 1
    nouveau = next(r for r in rows if r["nom"] == "KOFFI Ange Michel")
    assert nouveau["classe"] == "CP1" and nouveau["date_naissance"] == "2019-05-03"
    assert nouveau["pere_tel"] == "707070707" and nouveau["matricule_ecole"].startswith("CP1-")
    modifie = next(r for r in rows if r["inscription_id"] == existant["inscription_id"])
    assert modifie["mere_tel"] == "0101010101"
    assert modifie["date_naissance"] == existant["date_naissance"]   # case vide : valeur conservée
    assert modifie["matricule_ecole"] == existant["matricule_ecole"]


def test_import_eleves_meme_nom_sans_date_est_un_avertissement(demo, tmp_path):
    db, annee = demo
    f = mx.creer_modele(db, annee, "eleves", tmp_path / "e.xlsx")
    _remplir(f, "Élèves", [{"Nom et prénoms": "TRAORE Ali", "Sexe": "M", "Classe": "CE1"},
                           {"Nom et prénoms": "Traoré  Ali", "Sexe": "M", "Classe": "CE2"}])
    an = mx.analyser(db, annee, f)
    assert all(not l.erreurs for l in an.lignes)
    assert "Doublon possible" in an.lignes[1].avertissements[0]


def test_import_notes_absents_et_hors_limites(demo, tmp_path):
    db, annee = demo
    cl = annee.classe("CM2")
    g = annee.grilles[cl.grille_id]
    mats = g.matieres[2]
    f = mx.creer_modele(db, annee, "notes", tmp_path / "n.xlsx", num=2, classes=["CM2"])
    wb = load_workbook(f)
    ws = wb["CM2"]
    titres = [str(c.value).replace(" *", "") for c in ws[1]]
    col = {m["id"]: titres.index(f"{m['libelle']} (/{repo.max_note(g, m):g})") + 1 for m in mats}
    m1, m2 = mats[0], mats[1]
    mx1 = repo.max_note(g, m1)
    ws.cell(2, col[m1["id"]], mx1 / 2)
    ws.cell(2, col[m2["id"]], "12,5" if repo.max_note(g, m2) >= 12.5 else "1,5")
    ws.cell(3, titres.index("Présent") + 1, "NON")
    ws.cell(4, col[m1["id"]], mx1 + 1)                              # hors barème
    ws.cell(5, col[m1["id"]], "abs")                                # absent
    ws.cell(6, col[m1["id"]], "dix")
    ws.cell(7, 1, "XX-999-26")
    ws.cell(7, col[m1["id"]], 1)
    wb.save(f)
    an = mx.analyser(db, annee, f)
    assert an.kind == "notes" and an.evaluation == 2
    l = {x.numero: x for x in an.lignes}
    assert not l[2].erreurs and not l[3].erreurs and not l[5].erreurs
    assert any("hors limites" in e for e in l[4].erreurs)
    assert any("nombre attendu" in e for e in l[6].erreurs)
    assert any("inconnu en CM2" in e for e in l[7].erreurs)
    iid = {r["matricule_ecole"]: r["inscription_id"] for r in repo.inscriptions(db, annee, "CM2")}
    mats_ids = [m["id"] for m in mats]
    mx.appliquer(db, S, annee, an)
    _, lignes = repo.notes_grille(db, annee, cl.grille_id, 2, "CM2")
    par_iid = {r["inscription_id"]: r for r in lignes}
    r2 = par_iid[iid[ws.cell(2, 1).value]]
    assert r2["notes"][m1["id"]] == mx1 / 2
    assert not par_iid[iid[ws.cell(3, 1).value]]["present"]
    assert not par_iid[iid[ws.cell(3, 1).value]]["notes"]
    assert not par_iid[iid[ws.cell(5, 1).value]]["present"]
    assert set(r2["notes"]) <= set(mats_ids)


def test_import_absences_et_doublons(demo, tmp_path):
    db, annee = demo
    eleves = repo.inscriptions(db, annee)
    f = mx.creer_modele(db, annee, "absences_eleves", tmp_path / "a.xlsx")
    _remplir(f, "Absences élèves", [
        {"Date": date(2026, 10, 5), "Matricule école": eleves[0]["matricule_ecole"], "Nature": "ABSENCE",
         "Jours d'absence": 0.5, "Motif": "Maladie", "Justifié": "OUI"},
        {"Date": "05/10/2026", "Matricule école": eleves[0]["matricule_ecole"], "Nature": "absence",
         "Jours d'absence": 1},
        {"Date": date(2026, 10, 6), "Matricule école": eleves[1]["matricule_ecole"], "Nature": "RETARD",
         "Retard (minutes)": 15},
        {"Date": date(2025, 1, 6), "Matricule école": eleves[1]["matricule_ecole"], "Nature": "RETARD",
         "Retard (minutes)": 15},
        {"Date": date(2026, 10, 7), "Matricule école": "ZZ-1", "Nature": "RETARD", "Retard (minutes)": 5},
        {"Date": date(2026, 10, 8), "Matricule école": eleves[2]["matricule_ecole"], "Nature": "ABSENCE"},
    ])
    an = mx.analyser(db, annee, f)
    l = {x.numero: x for x in an.lignes}
    assert not l[2].erreurs and not l[4].erreurs
    assert any("Même absence que la ligne 2" in e for e in l[3].erreurs)
    assert any("année scolaire" in e for e in l[5].erreurs)
    assert any("inconnu" in e for e in l[6].erreurs)
    assert any("jours" in e for e in l[7].erreurs)
    n = len(repo.absences_eleves(db, annee))
    assert mx.appliquer(db, S, annee, an)["enregistrees"] == 2
    assert len(repo.absences_eleves(db, annee)) == n + 2
    # le même fichier importé une seconde fois ne crée pas de doublon
    an2 = mx.analyser(db, annee, f)
    assert {x.action for x in an2.lignes if not x.erreurs} == {"ignorer"}
    assert mx.appliquer(db, S, annee, an2)["enregistrees"] == 0
    assert len(repo.absences_eleves(db, annee)) == n + 2


def test_import_personnel_et_classes(demo, tmp_path):
    db, annee = demo
    f = mx.creer_modele(db, annee, "personnel", tmp_path / "p.xlsx")
    _remplir(f, "Personnel", [
        {"Matricule": "999999Z", "Nom et prénoms": "YAO Konan", "Sexe": "M", "Fonction": "INSTITUTEUR ADJOINT",
         "Classe tenue": "CE1", "Téléphone": 505050505},
        {"Matricule": "999999Z", "Nom et prénoms": "Autre", "Sexe": "F"},
        {"Matricule": "888888Y", "Nom et prénoms": "KONE Awa", "Sexe": "F", "Classe tenue": "CM7"},
    ])
    an = mx.analyser(db, annee, f)
    assert [bool(l.erreurs) for l in an.lignes] == [False, True, True]
    mx.appliquer(db, S, annee, an)
    p = next(x for x in repo.personnel(db, annee) if x["matricule"] == "999999Z")
    assert p["classe_tenue"] == "CE1" and p["telephone"] == "505050505"

    f = mx.creer_modele(db, annee, "classes", tmp_path / "c.xlsx")
    grille = annee.grilles[annee.classe("CM2").grille_id].code
    _remplir(f, "Classes", [
        {"Code": "cm2", "Seuil d'admission": 9, "Barème": 20, "Grille de notes": grille},
        {"Code": "CM3", "Seuil d'admission": 12, "Barème": 10, "Grille de notes": grille},
        {"Code": "CM4", "Seuil d'admission": 10, "Barème": 20, "Grille de notes": "INCONNUE",
         "Classe suivante": "6EME"},
    ])
    an = mx.analyser(db, annee, f)
    l = an.lignes
    assert l[0].action == "modifier" and not l[0].erreurs
    assert any("seuil" in e for e in l[1].erreurs)
    assert any("Grille" in e for e in l[2].erreurs) and any("suivante" in e for e in l[2].erreurs)
    mx.appliquer(db, S, annee, an)
    cm2 = repo.load_annee(db, annee.id).classe("CM2")
    assert cm2.seuil == 9 and cm2.libelle == annee.classe("CM2").libelle   # le reste est conservé


def test_fichiers_refuses(demo, tmp_path):
    db, annee = demo
    f = mx.creer_modele(db, annee, "eleves", tmp_path / "e.xlsx")
    an = mx.analyser(db, annee, f, kind="personnel")
    assert an.erreurs_fichier and "Élèves" in an.erreurs_fichier[0]
    # colonne obligatoire supprimée
    wb = load_workbook(f)
    wb["Élèves"].delete_cols(3)
    wb.save(f)
    an = mx.analyser(db, annee, f)
    assert any("Nom et prénoms" in e for e in an.erreurs_fichier)
    with pytest.raises(ValidationError):
        mx.appliquer(db, S, annee, an)
    # fichier d'une autre année
    wb = load_workbook(mx.creer_modele(db, annee, "eleves", tmp_path / "e2.xlsx"))
    wb[mx.META]["B2"] = "2019-2020"
    wb.save(tmp_path / "e2.xlsx")
    assert "2019-2020" in mx.analyser(db, annee, tmp_path / "e2.xlsx").erreurs_fichier[0]
    # fichier qui n'est pas un classeur
    (tmp_path / "x.xlsx").write_text("pas un classeur")
    assert "illisible" in mx.analyser(db, annee, tmp_path / "x.xlsx").erreurs_fichier[0]


def test_annee_cloturee_refusee(demo, tmp_path):
    db, annee = demo
    f = mx.exporter(db, S, annee, "eleves", tmp_path / "e.xlsx")
    db.execute("UPDATE annee_scolaire SET cloturee=1 WHERE id=?", (annee.id,))
    an = mx.analyser(db, repo.load_annee(db, annee.id), f)
    assert any("clôturée" in e for e in an.erreurs_fichier)


def test_droits(demo, tmp_path):
    db, annee = demo
    secretariat = Session(2, "sec", "Secrétariat", "SECRETARIAT", "Secrétariat", {"export.excel", "eleves.lire"})
    enseignant = Session(3, "ens", "Enseignant", "ENSEIGNANT", "Enseignant", {"notes.lire"})
    assert mx.exporter(db, secretariat, annee, "eleves", tmp_path / "e.xlsx").exists()
    with pytest.raises(AccessDenied):
        mx.exporter(db, enseignant, annee, "eleves", tmp_path / "e2.xlsx")
    f = mx.creer_modele(db, annee, "eleves", tmp_path / "v.xlsx")
    _remplir(f, "Élèves", [{"Nom et prénoms": "NOUVEL Élève", "Sexe": "F", "Classe": "CP1"}])
    an = mx.analyser(db, annee, f)
    with pytest.raises(AccessDenied):
        mx.appliquer(db, secretariat, annee, an)
    assert not any(r["nom"] == "NOUVEL Élève" for r in repo.inscriptions(db, annee))


def test_echec_en_cours_d_import_ne_laisse_rien(demo, tmp_path, monkeypatch):
    db, annee = demo
    f = mx.creer_modele(db, annee, "eleves", tmp_path / "e.xlsx")
    _remplir(f, "Élèves", [{"Nom et prénoms": "PREMIER Élève", "Sexe": "F", "Classe": "CP1"},
                           {"Nom et prénoms": "SECOND Élève", "Sexe": "M", "Classe": "CP1"}])
    an = mx.analyser(db, annee, f)
    vrai = repo.inscrire
    appels = []

    def inscrire(*a, **k):
        appels.append(1)
        if len(appels) == 2:
            raise ValidationError("incident simulé")
        return vrai(*a, **k)
    monkeypatch.setattr(repo, "inscrire", inscrire)
    with pytest.raises(ValidationError, match="Rien n'a été importé"):
        mx.appliquer(db, S, annee, an)
    noms = {r["nom"] for r in repo.inscriptions(db, annee)}
    assert "PREMIER Élève" not in noms
    assert not db.scalar("SELECT COUNT(*) FROM eleve WHERE nom LIKE '%Élève'")

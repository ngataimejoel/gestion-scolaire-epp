"""Moteur de rapports : génération PDF / Excel de chaque modèle, mise en page, filtres,
regroupements, totaux, lots."""
from __future__ import annotations

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from epp.core.security import system_session
from epp.reports import engine, model, render_pdf, service, store

PARAMS = {"MENSUEL": {"mois": 10}, "COMPOSITION": {"compo": 1}, "FEUILLE_NOTES": {"classe": "CE1", "compo": 1},
          "ABS_ELEVES": {"mois": 10}, "ABS_PERSONNEL": {"mois": 10}}


def _textes(pdf):
    return [(p.extract_text() or "") for p in PdfReader(pdf).pages]


def test_tous_les_modeles_integres_generent_pdf_et_excel(idb, annee, tmp_path):
    s = system_session()
    modeles = store.lister(idb)
    assert len([m for m in modeles if m["integre"]]) >= 20
    for m in modeles:
        p = PARAMS.get(m["code"], {})
        pdf = service.generer(idb, s, annee, m["code"], p, tmp_path / f"{m['code']}.pdf")
        pages = _textes(pdf)
        assert pages, m["code"]
        # aucune page vide : chaque page porte au moins un contenu en plus du pied de page
        for i, t in enumerate(pages, 1):
            assert len(t.replace("Page", "").strip()) > 60, f"{m['code']} : page {i} vide"
            assert f"Page {i} / {len(pages)}" in t, f"{m['code']} : numérotation absente page {i}"
        xl = service.generer(idb, s, annee, m["code"], p, tmp_path / f"{m['code']}.xlsx")
        ws = load_workbook(xl).active
        assert ws.max_row > 3, m["code"]


def _largeurs_ok(defn, idb, annee, params):
    """Vérifie qu'aucune colonne n'est plus étroite que son plus long mot (pas de texte coupé)
    et que la largeur totale ne dépasse pas la page (pas de colonnes superposées)."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    doc = engine.build_document(idb, annee, defn, params)
    size = landscape(A4) if doc.orientation == "paysage" else A4
    avail = size[0] - render_pdf.MARGE_G - render_pdf.MARGE_D
    render_pdf._register_fonts()
    for b in doc.blocks:
        if not isinstance(b, engine.Table):
            continue
        fl = render_pdf._table(b, avail, "FCFA")
        tbls = [f for f in fl if isinstance(f, render_pdf.RLTable)]
        for f in fl:
            if isinstance(f, render_pdf.KeepTogether):
                tbls += [x for x in f._content if isinstance(x, render_pdf.RLTable)]
        for t in tbls:
            widths = t._colWidths
            assert sum(widths) <= avail + 0.5, "tableau plus large que la page"
            spanned = set()
            for cmd in t._spanCmds:
                (c0, r0), (c1, r1) = cmd[1], cmd[2]
                c1 = len(widths) - 1 if c1 < 0 else c1
                if c1 > c0:
                    for r in range(r0, r1 + 1):
                        spanned.add((r, c0))
            for r, row in enumerate(t._cellvalues):
                for i, cell in enumerate(row):
                    if not hasattr(cell, "text") or (r, i) in spanned:
                        continue
                    font = cell.style.fontName
                    for w in render_pdf._words(cell.text):
                        assert pdfmetrics.stringWidth(w, font, cell.style.fontSize) <= widths[i] - 5, \
                            f"mot coupé : {w!r} (colonne {i + 1})"


@pytest.mark.parametrize("code", ["REGISTRE", "RESULTATS", "EFFECTIFS", "SYNTHESE", "COMPOSITION", "MENSUEL",
                                  "PERSONNEL", "ETAT_CLASSE", "RENTREE"])
def test_mise_en_page_sans_debordement(idb, annee, code):
    defn = store.definition(idb, code)
    _largeurs_ok(defn, idb, annee, PARAMS.get(code, {}))


def test_valeurs_du_rapport_de_composition(idb, annee, excel):
    doc = engine.build_document(idb, annee, store.definition(idb, "COMPOSITION"), {"compo": 1})
    tab = next(b for b in doc.blocks if isinstance(b, engine.Table))
    ws = excel["RAPPORT DE COMPOSITION"]
    for i, row in enumerate(tab.rows):
        for j, v in enumerate(row.cells[1:], start=2):
            assert abs((ws.cell(15 + i, j).value or 0) - (v or 0)) < 1e-9


def test_observation_mensuelle_identique_au_classeur(idb, annee, excel):
    doc = engine.build_document(idb, annee, store.definition(idb, "MENSUEL"), {"mois": 10})
    tab = next(b for b in doc.blocks if isinstance(b, engine.Table))
    ws = excel["RAPPORT MENSUEL"]
    obs = {r.cells[0]: r.cells[-1] for r in tab.rows}
    for r in range(13, 23):
        cl = ws[f"A{r}"].value
        attendu = (ws[f"K{r}"].value or "").replace("22.5", "22,5").replace("3.5", "3,5")
        assert (obs.get(cl) or "") == attendu, cl


def _defn(sections, **kw):
    d = model.nouveau_modele()
    d["sections"] = sections
    d.update(kw)
    return d


def test_filtre_tri_regroupement_et_totaux(idb, annee):
    d = _defn([{"type": "tableau", "source": "eleves",
                "colonnes": [{"champ": "numero"}, {"champ": "nom"}, {"champ": "classe"}, {"champ": "age"}],
                "filtres": [{"champ": "sexe", "op": "=", "valeur": "F"},
                            {"champ": "classe", "op": "dans", "valeur": "CP1;CP2"}],
                "tri": [{"champ": "nom", "ordre": "asc"}], "regroupement": "classe",
                "agregats": {"age": "moyenne", "numero": ""}, "sous_totaux": True, "total_general": True}])
    doc = engine.build_document(idb, annee, d, {})
    t = next(b for b in doc.blocks if isinstance(b, engine.Table))
    kinds = [r.kind for r in t.rows]
    assert kinds.count("group") == 2 and kinds.count("subtotal") == 2 and kinds[-1] == "total"
    assert t.rows[0].cells[0].startswith("Classe : CP1")
    data = [r for r in t.rows if r.kind == "data"]
    assert len(data) == 11       # 6 filles en CP1 + 5 en CP2
    cp1 = [r.cells[1] for r in t.rows[1:7]]
    assert cp1 == sorted(cp1, key=str.lower)
    ages = [r.cells[3] for r in data]
    assert t.rows[-1].cells[3] == round(sum(ages) / len(ages), 2)   # moyenne arrondie à 2 décimales


def test_filtre_lie_a_un_parametre(idb, annee):
    d = _defn([{"type": "tableau", "source": "eleves", "colonnes": [{"champ": "nom"}],
                "filtres": [{"champ": "classe", "op": "=", "parametre": "classe"}]}], parametres=["classe"])
    tout = engine.build_document(idb, annee, d, {})
    cm2 = engine.build_document(idb, annee, d, {"classe": "CM2"})
    n = lambda doc: len(next(b for b in doc.blocks if isinstance(b, engine.Table)).rows)   # noqa: E731
    assert n(tout) == 64 and n(cm2) == 9


def test_variables_dans_les_titres(idb, annee):
    d = _defn([{"type": "texte", "texte": "{ecole} – {annee} – {mois_annee} – {compo_libelle}"}],
              titre="Rapport {classe}", parametres=["classe", "mois", "compo"])
    doc = engine.build_document(idb, annee, d, {"classe": "CE1", "mois": 10, "compo": 1})
    assert doc.titre == "Rapport CE1"
    para = next(b for b in doc.blocks if isinstance(b, engine.Paragraph))
    assert para.text.startswith("EPP LIGUIYO – 2026-2027 – Octobre 2026 – ")


def test_modele_vide_ne_produit_pas_de_page_blanche(idb, annee, tmp_path):
    d = _defn([{"type": "tableau", "source": "eleves", "colonnes": [{"champ": "nom"}],
                "filtres": [{"champ": "nom", "op": "=", "valeur": "PERSONNE"}]}])
    pdf = tmp_path / "vide.pdf"
    render_pdf.render(engine.build_document(idb, annee, d, {}), pdf)
    pages = _textes(pdf)
    assert len(pages) == 1 and "Aucune donnée" in pages[0]


def test_lot_dossier_fin_annee(idb, annee, tmp_path):
    p = service.generer_lot_nomme(idb, system_session(), annee, "DOSSIER_FIN_ANNEE", {}, tmp_path)
    txt = "\n".join(_textes(p))
    for titre in ("TABLEAU DE BORD", "TABLEAU DES EFFECTIFS", "SYNTHÈSE", "RAPPORT DE RENTRÉE", "LISTE DU PERSONNEL"):
        assert titre in txt


def test_un_fichier_par_classe(idb, annee, tmp_path):
    out = service.un_fichier_par_classe(idb, system_session(), annee, "EX_LISTE_CLASSE", {}, tmp_path)
    assert len(out) == 6 and all(p.exists() for p in out)
    assert "CP1" in out[0].name


def test_orientation_paysage(idb, annee, tmp_path):
    pdf = service.generer(idb, system_session(), annee, "REGISTRE", {}, tmp_path / "r.pdf")
    box = PdfReader(pdf).pages[0].mediabox
    assert float(box.width) > float(box.height)


def test_export_excel_entetes_fusionnes_et_formats(idb, annee, tmp_path):
    xl = service.generer(idb, system_session(), annee, "SYNTHESE", {}, tmp_path / "s.xlsx")
    ws = load_workbook(xl).active
    assert ws.merged_cells.ranges
    assert ws.page_setup.orientation == "landscape"
    pct = [c for row in ws.iter_rows() for c in row if c.number_format == "0.0%"]
    assert pct

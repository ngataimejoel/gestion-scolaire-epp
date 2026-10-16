"""Fiche scolaire de l'élève (QR code), en-tête « fiche scolaire » et parents sur les documents de l'élève."""
from __future__ import annotations

import docx
from pypdf import PdfReader

from epp.core import seed
from epp.core.security import system_session
from epp.domain import autorisations as au
from epp.domain import fiche_scolaire, repo
from epp.reports import service


def _texte(pdf, page=0) -> str:
    return " ".join(PdfReader(pdf).pages[page].extract_text().split())


def _eleve(db, annee, **parents):
    r = repo.inscriptions(db, annee, annee.codes[0])[0]
    e = repo.get_eleve(db, r["eleve_id"])
    e.update({"pere_nom": "KOUAME Kouadio Jean", "pere_profession": "Planteur", "pere_residence": "OUANINOU",
              "pere_tel": "07 11 22 33 44", "mere_nom": "BAMBA Awa", "mere_profession": "Commerçante",
              "groupe_sanguin": "O+", "allergies": "Arachide", **parents})
    repo.save_eleve(db, system_session(), r["eleve_id"], e, annee)
    return repo.inscriptions(db, annee, annee.codes[0])[0]


def test_nom_prenoms_niveau_et_ligne_parents():
    assert fiche_scolaire.separer_nom("KOUAME Kouadio Jean") == ("KOUAME", "Kouadio Jean")
    assert fiche_scolaire.separer_nom(None) == ("", "")
    assert fiche_scolaire.niveau("CE1 A") == "CE1" and fiche_scolaire.niveau("CM2") == "CM2"
    r = {"pere_nom": "KOUAME Jean", "pere_profession": "Planteur", "mere_nom": "BAMBA Awa"}
    assert fiche_scolaire.ligne_parents(r) == "Père : KOUAME Jean (Planteur)  •  Mère : BAMBA Awa"
    assert fiche_scolaire.ligne_parents({"tuteur_nom": "YAO Paul"}) == "Tuteur : YAO Paul"


def test_qr_code_lisible():
    png = fiche_scolaire.qr_png("ECOLE: EPP OUANINOU 2\nELEVE: KONE Awa")
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    m = fiche_scolaire.qr_matrice("ECOLE: EPP OUANINOU 2\nELEVE: KONE Awa")
    n = len(m)
    assert n >= 21 and (n - 17) % 4 == 0 and all(len(row) == n for row in m)
    assert all(m[i][j] for i, j in ((0, 0), (0, 6), (6, 0), (6, 6)))     # repères de position


def test_fiche_scolaire_pdf_un_eleve(demo, tmp_path):
    db, annee = demo
    s = system_session()
    r = _eleve(db, annee)
    d = fiche_scolaire.fiches(db, annee, inscription_id=r["inscription_id"])[0]
    for attendu in ("ECOLE: EPP OUANINOU 2", "IEPP: IEPP OUANINOU", f"ELEVE: {r['nom']}",
                    f"MATRICULE ECOLE: {r['matricule_ecole']}", "PÈRE: KOUAME Kouadio Jean - Planteur",
                    "MÈRE: BAMBA Awa - Commerçante"):
        assert attendu in d["qr"]
    assert d["scolarite"][-1]["annee"] == annee.libelle and d["scolarite"][-1]["effectif"]
    pdf = service.generer(db, s, annee, "FICHE_SCOLAIRE", {"classe": r["classe"], "eleve": r["inscription_id"]},
                          tmp_path / "f.pdf")
    rd = PdfReader(pdf)
    assert len(rd.pages) == 1
    t = _texte(pdf)
    for attendu in ("FICHE SCOLAIRE", "REPUBLIQUE DE COTE D'IVOIRE", "IEPP OUANINOU", "DRENAET DE TOUBA",
                    "KOUAME", "Kouadio Jean", "Planteur", "BAMBA", "Commerçante", "07 11 22 33 44", "O+",
                    "Arachide", "PÈRE", "MÈRE", "TUTEUR / TUTRICE", "Le Directeur de l'école", "1ère", "8ème",
                    annee.libelle.replace("-", "/"), r["matricule_ecole"]):
        assert attendu in t, attendu


def test_fiches_d_une_classe_une_page_par_eleve(demo, tmp_path):
    db, annee = demo
    s = system_session()
    cl = annee.codes[1]
    n = len(repo.inscriptions(db, annee, cl))
    rd = PdfReader(service.generer(db, s, annee, "FICHE_SCOLAIRE", {"classe": cl}, tmp_path / "c.pdf"))
    assert len(rd.pages) == n


def test_fiche_scolaire_word(demo, tmp_path):
    db, annee = demo
    s = system_session()
    r = _eleve(db, annee)
    f = service.generer(db, s, annee, "FICHE_SCOLAIRE", {"eleve": r["inscription_id"]}, tmp_path / "f.docx")
    d = docx.Document(f)
    texte = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    texte += "\n".join(p.text for p in d.paragraphs)
    for attendu in ("FICHE SCOLAIRE", "KOUAME", "Planteur", "BAMBA", "Commerçante", "MINISTERE", "PÈRE"):
        assert attendu in texte, attendu
    assert d.inline_shapes and len(d.inline_shapes) >= 1             # code QR


def test_en_tete_style_fiche_sur_les_rapports(demo, tmp_path):
    db, annee = demo
    t = _texte(service.generer(db, system_session(), annee, "EFFECTIFS", {}, tmp_path / "e.pdf"))
    assert seed.MINISTERE_DEFAUT in t and "DRENAET DE TOUBA" in t and "Union - Discipline - Travail" in t
    assert "-" * 20 in t                                              # tirets de séparation


def test_parents_sur_les_bulletins(demo, tmp_path):
    db, annee = demo
    s = system_session()
    r = _eleve(db, annee)
    for code in ("BULLETIN_ANNUEL", "BULLETIN_EVAL"):
        t = _texte(service.generer(db, s, annee, code, {"classe": r["classe"], "eleve": r["inscription_id"],
                                                        "compo": 1}, tmp_path / f"{code}.pdf"))
        assert "Père : KOUAME Kouadio Jean (Planteur)" in t and "Mère : BAMBA Awa (Commerçante)" in t, code
        assert "REPUBLIQUE DE COTE D'IVOIRE" in t
    w = docx.Document(service.generer(db, s, annee, "BULLETIN_ANNUEL", {"eleve": r["inscription_id"]},
                                      tmp_path / "b.docx"))
    texte = "\n".join(c.text for t in w.tables for row in t.rows for c in row.cells)
    assert "KOUAME Kouadio Jean (Planteur)" in texte


def test_parents_sur_la_demande_d_autorisation_eleve(demo, tmp_path):
    db, annee = demo
    s = system_session()
    r = _eleve(db, annee)
    aid = au.enregistrer(db, s, annee, None, {"kind": "eleve", "inscription_id": r["inscription_id"],
                                              "date_demande": "2026-11-02", "date_debut": "2026-11-03",
                                              "date_fin": "2026-11-03", "motif": "Maladie"})
    t = _texte(au.imprimer(db, s, annee, au.get(db, aid)["numero"], tmp_path / "a.pdf"))
    assert "KOUAME Kouadio Jean – Planteur" in t and "BAMBA Awa – Commerçante" in t


def test_migration_v7(tmp_path):
    from epp.core import db as dbmod
    path = tmp_path / "ancienne.sqlite"
    d = dbmod.open_database(path)
    d.set_param("ministere", seed.ANCIEN_MINISTERE)
    d.set_meta("schema_version", 6)
    d.close()
    d = dbmod.open_database(path)
    assert {"groupe_sanguin", "allergies"} <= {r[1] for r in d.query("PRAGMA table_info(eleve)")}
    assert repo.etablissement(d)["ministere"] == seed.MINISTERE_DEFAUT
    d.close()


def test_modele_excel_eleves_avec_sante(demo, tmp_path):
    import openpyxl
    from epp.io import modeles_excel
    db, annee = demo
    _eleve(db, annee)
    f = modeles_excel.exporter(db, system_session(), annee, "eleves", tmp_path / "e.xlsx")
    ws = openpyxl.load_workbook(f)["Élèves"]
    entetes = [str(c.value or "") for c in ws[1]]
    assert any(e.startswith("Groupe sanguin") for e in entetes)
    assert any(e.startswith("Allergies") for e in entetes)

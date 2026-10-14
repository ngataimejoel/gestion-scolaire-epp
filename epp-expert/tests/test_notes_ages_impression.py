"""Demandes du 14/10/2026 : motifs d'absence, total en CP, Tab et correction protégée dans la saisie des notes,
âges par niveau et par sexe (avec / sans extrait), orientation des pages à l'impression."""
from __future__ import annotations

import os
import sqlite3

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["EPP_SANS_ANIMATION"] = "1"
os.environ["EPP_SANS_VIDEO"] = "1"

from openpyxl import load_workbook                                      # noqa: E402
from PySide6.QtCore import QCoreApplication, QEvent, Qt                 # noqa: E402
from PySide6.QtWidgets import QAbstractItemDelegate, QAbstractItemView, QApplication, QMessageBox   # noqa: E402

from epp.core import licence, security                                  # noqa: E402
from epp.core.security import system_session                            # noqa: E402
from epp.domain import calculs, repo, stats                             # noqa: E402
from epp.io import modeles_excel as mx                                  # noqa: E402

S = system_session()


@pytest.fixture(scope="module")
def app():
    a = QApplication.instance() or QApplication([])
    yield a
    for w in QApplication.topLevelWidgets():
        w.close()
        w.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    a.processEvents()


@pytest.fixture
def messages(monkeypatch):
    journal = []

    def boite(kind):
        def f(parent, titre, texte, *a, **k):
            journal.append((kind, titre, texte))
            return QMessageBox.Ok
        return staticmethod(f)
    for k in ("critical", "warning", "information", "question"):
        monkeypatch.setattr(QMessageBox, k, boite(k))
    monkeypatch.setattr(QMessageBox, "exec", lambda self: journal.append(("box", self.windowTitle(), self.text())))
    monkeypatch.setattr(QMessageBox, "clickedButton", lambda self: self.buttons()[0] if self.buttons() else None)
    monkeypatch.setattr(licence, "etat", lambda *a, **k: licence.Etat("essai", jours_restants=5))
    return journal


@pytest.fixture
def page_notes(app, demo, messages):
    from epp.ui.pages.notes import PageNotes
    from epp.ui.widgets import Contexte
    db, annee = demo
    p = PageNotes(Contexte(db, S, annee, licence.Etat("essai", jours_restants=5)))
    p.show()
    p.actualiser()
    yield p
    p.hide()
    p.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


# =================================================================== motifs d'absence
def test_motifs_d_absence(db, tmp_path):
    motifs = repo.liste(db, "MOTIF_ABSENCE")
    assert "Travaux champêtres" not in motifs
    assert {"Opération bancaire", "Congé de maternité"} <= set(motifs) and motifs[-1] == "Autre"
    # base existante (version 5) : motif retiré, nouveaux motifs ajoutés avant « Autre »
    from epp.core.db import SCHEMA_VERSION, open_database
    chemin = tmp_path / "ancienne.sqlite"
    open_database(chemin).close()
    con = sqlite3.connect(chemin)
    con.execute("DELETE FROM liste_valeur WHERE valeur IN ('Opération bancaire', 'Congé de maternité')")
    con.execute("INSERT INTO liste_valeur(liste, valeur, ordre) VALUES('MOTIF_ABSENCE', 'Travaux champêtres', 5)")
    con.execute("UPDATE meta SET valeur='5' WHERE cle='schema_version'")
    con.commit()
    con.close()
    d = open_database(chemin)
    try:
        assert d.meta("schema_version") == str(SCHEMA_VERSION)
        motifs = repo.liste(d, "MOTIF_ABSENCE")
        assert "Travaux champêtres" not in motifs
        assert {"Opération bancaire", "Congé de maternité"} <= set(motifs) and motifs[-1] == "Autre"
    finally:
        d.close()


# =================================================================== notes
def test_total_en_cp():
    assert calculs.moyenne_evaluation([(7, 1), (8, 1), (None, 1)], "COEF", 10) == (15, 7.5)
    assert calculs.moyenne_evaluation([(6, 2), (9, 1)], "COEF", 10) == (21, 7)
    assert calculs.moyenne_evaluation([(6, 1)], "COEF", 10, present=False) == (0, 0)
    assert calculs.moyenne_evaluation([(None, 1)], "COEF", 10) == (None, None)


def test_saisie_tab_total_cp_et_correction_protegee(app, page_notes, messages, monkeypatch):
    p, m, v = page_notes, page_notes.modele, page_notes.vue
    assert p.classe.currentData() == "CP1" and m.grille.mode == "COEF"
    cols = [c for c in range(m.columnCount()) if m._col(c)[0] == "note"]
    c_total = cols[-1] + 1
    assert m.headerData(c_total, Qt.Horizontal) == "Total"
    r0 = 0
    iid, mat = m.lignes[r0]["inscription_id"], m.mats[0]
    p.db.execute("DELETE FROM note WHERE inscription_id=? AND matiere_id=?", (iid, mat["id"]))
    m._recharger()
    demandes = []
    reponse = [None]

    def faux_mdp(parent, db, s, texte, motif):
        demandes.append(texte)
        return reponse[0]
    monkeypatch.setattr("epp.ui.pages.notes.mot_de_passe_admin", faux_mdp)

    idx = m.index(r0, cols[0])
    assert m.setData(idx, "7,5") and not demandes          # première saisie : pas de mot de passe
    assert m.data(m.index(r0, c_total)) not in ("", None)   # total affiché en CP
    assert m.lignes[r0]["total"] == sum(x for x in m.lignes[r0]["notes"].values())
    assert m.setData(idx, "7,5") and not demandes          # même valeur : rien à corriger
    assert not m.setData(idx, "8") and len(demandes) == 1  # correction refusée sans mot de passe
    assert m.lignes[r0]["notes"][mat["id"]] == 7.5
    reponse[0] = "admin"
    assert m.setData(idx, "8") and len(demandes) == 2
    assert m.lignes[r0]["notes"][mat["id"]] == 8
    assert m.setData(idx, "") and len(demandes) == 2       # déjà autorisé pour cette classe et cette évaluation
    assert mat["id"] not in m.lignes[r0]["notes"]
    assert p.db.scalar("SELECT COUNT(*) FROM journal WHERE action='NOTE_CORRIGEE'") == 2
    # autre classe : le mot de passe est redemandé
    p.classe.setCurrentIndex(p.classe.findData("CP2"))
    p._remplir_evals()
    assert p._correction is None

    # Tab : matière suivante, puis premier élément de l'élève suivant ; Maj+Tab : en sens inverse
    p.classe.setCurrentIndex(p.classe.findData("CP1"))
    p._remplir_evals()
    v.setCurrentIndex(m.index(0, cols[0]))
    assert v._voisine(m.index(0, cols[0]), 1) == m.index(0, cols[1])
    assert v._voisine(m.index(0, cols[-1]), 1) == m.index(1, cols[0])
    assert v._voisine(m.index(1, cols[0]), -1) == m.index(0, cols[-1])
    assert v.moveCursor(QAbstractItemView.CursorAction.MoveNext, Qt.NoModifier) == m.index(0, cols[1])
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QLineEdit
    p.db.execute("DELETE FROM note WHERE inscription_id IN (?, ?)",
                 (m.lignes[0]["inscription_id"], m.lignes[1]["inscription_id"]))
    m._recharger()

    def saisir(texte, touche):
        ed = v.viewport().findChildren(QLineEdit)
        assert ed, "pas de saisie ouverte"
        ed[-1].setText(texte)
        QTest.keyClick(ed[-1], touche)
        app.processEvents()

    v.setFocus()
    v.setCurrentIndex(m.index(0, cols[0]))
    v.edit(m.index(0, cols[0]))
    saisir("6", Qt.Key_Tab)                                   # Tab après une saisie : matière suivante, ouverte
    assert m.lignes[0]["notes"][m.mats[0]["id"]] == 6
    assert v.currentIndex() == m.index(0, cols[1]) and v.state() == QAbstractItemView.State.EditingState
    saisir("7", Qt.Key_Return)                                # Entrée : élève suivant, même matière
    assert m.lignes[0]["notes"][m.mats[1]["id"]] == 7
    assert v.currentIndex() == m.index(1, cols[1])
    v.setCurrentIndex(m.index(0, cols[-1]))
    v.edit(m.index(0, cols[-1]))
    saisir("5", Qt.Key_Tab)                                   # bout de ligne : élève suivant, 1re matière
    assert v.currentIndex() == m.index(1, cols[0])
    saisir("4", Qt.Key_Backtab)                               # Maj+Tab : retour à la dernière matière
    assert m.lignes[1]["notes"][m.mats[0]["id"]] == 4
    assert v.currentIndex() == m.index(0, cols[-1])


def test_verifier_admin(db):
    security.create_user(db, "chef", "Chef", "Secret-2026", "ADMIN")
    assert security.verifier_admin(db, S, "Secret-2026", "test") == "chef"
    assert security.verifier_admin(db, S, "faux", "test") is None
    assert security.verifier_admin(db, S, "", "test") is None
    actions = [r["action"] for r in db.query("SELECT action FROM journal ORDER BY id")]
    assert actions.count("AUTORISATION_ADMIN") == 1 and actions.count("AUTORISATION_ADMIN_ECHEC") == 2


def test_import_qui_corrige_des_notes(demo, tmp_path):
    db, annee = demo
    cl = annee.classe("CP1")
    num = next(n for n in annee.eval_nums()
               if any(r["notes"] for r in repo.notes_grille(db, annee, cl.grille_id, n, "CP1")[1]))
    f = mx.creer_modele(db, annee, "notes", tmp_path / "n.xlsx", num=num, classes=["CP1"])
    an = mx.analyser(db, annee, f)
    assert an.corrections == 0                               # notes réimportées à l'identique
    wb = load_workbook(f)
    ws = wb["CP1"]
    _, lignes = repo.notes_grille(db, annee, cl.grille_id, num, "CP1")
    par_mat = {r["matricule_ecole"]: r for r in lignes}
    titres = [str(c.value).replace(" *", "") for c in ws[1]]
    for row in range(2, ws.max_row + 1):
        r = par_mat.get(ws.cell(row, 1).value)
        if r and r["notes"]:
            mid, val = next(iter(r["notes"].items()))
            g = annee.grilles[cl.grille_id]
            mat = next(x for x in g.matieres[num] if x["id"] == mid)
            c = titres.index(f"{mat['libelle']} (/{repo.max_note(g, mat):g})") + 1
            ws.cell(row, c, 0 if val else 1)
            break
    wb.save(f)
    an = mx.analyser(db, annee, f)
    assert an.corrections == 1
    assert any("mot de passe" in a for l in an.lignes for a in l.avertissements)


# =================================================================== âges
def test_repartition_des_ages(demo, tmp_path):
    db, annee = demo
    ins = repo.inscriptions(db, annee)
    tous = stats.repartition_ages(db, annee)
    avec = stats.repartition_ages(db, annee, "avec")
    sans = stats.repartition_ages(db, annee, "sans")
    assert tous["n"][("TOTAL", "TOTAL", "T")] == len(ins)
    assert avec["n"][("TOTAL", "TOTAL", "T")] == sum(1 for r in ins if r["extrait"] == "OUI")
    assert sans["n"].get(("TOTAL", "TOTAL", "T"), 0) == sum(1 for r in ins if r["extrait"] != "OUI")
    assert sans["n"].get(("TOTAL", "TOTAL", "T"), 0) > 0
    for cle, n in tous["n"].items():
        assert avec["n"].get(cle, 0) + sans["n"].get(cle, 0) == n, cle
    for cl in annee.codes:
        attendu = sum(1 for r in ins if r["classe"] == cl and r["sexe"] == "F")
        assert tous["n"].get(("TOTAL", cl, "F"), 0) == attendu
        for ag in tous["ages"]:
            assert tous["n"].get((ag, cl, "M"), 0) + tous["n"].get((ag, cl, "F"), 0) == tous["n"].get((ag, cl, "T"), 0)
    r = next(x for x in ins if x["age"] is not None)
    assert tous["n"][(r["age"], r["classe"], r["sexe"])] >= 1
    assert tous["ages"] == sorted(tous["ages"], key=lambda a: (isinstance(a, str), a if isinstance(a, int) else 0))
    with pytest.raises(ValueError):
        stats.repartition_ages(db, annee, "autre")
    # état imprimable : trois tableaux en paysage
    from epp.reports import service
    from pypdf import PdfReader
    out = service.generer(db, S, annee, "AGES", {}, tmp_path / "ages.pdf", "pdf")
    pdf = PdfReader(str(out))
    texte = " ".join(pg.extract_text() for pg in pdf.pages)
    for t in ("ENSEMBLE DES ÉLÈVES", "AVEC EXTRAIT DE NAISSANCE", "SANS EXTRAIT DE NAISSANCE", "CP1", "CM2"):
        assert t in texte, t
    pg = pdf.pages[0].mediabox
    assert float(pg.width) > float(pg.height)


def test_page_statistiques_des_ages(app, demo, messages):
    from epp.ui.pages.statistiques import PageStatistiques
    from epp.ui.widgets import Contexte
    db, annee = demo
    p = PageStatistiques(Contexte(db, S, annee, licence.Etat("essai", jours_restants=5)))
    p.show()
    p.actualiser()
    p.periode.setCurrentIndex(1)
    assert p.periode.currentData()[0] == "ages"
    p.extrait.setCurrentIndex(p.extrait.findData("sans"))
    p._calculer()
    assert not p.extrait.isHidden() and p.sexe.isHidden()
    t = p.table.modele
    total = t.lignes[-1]
    assert total["age"] == "TOTAL"
    assert int(total["TOTAL_T"]) == stats.repartition_ages(db, annee, "sans")["n"][("TOTAL", "TOTAL", "T")]
    p.periode.setCurrentIndex(0)
    p._calculer()
    assert not p.sexe.isHidden() and p.extrait.isHidden()
    p.hide()
    p.deleteLater()


# =================================================================== impression
def test_impression_selon_l_orientation_du_document(app, tmp_path):
    from PySide6.QtPdf import QPdfDocument
    from PySide6.QtPrintSupport import QPrinter
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas

    from epp.ui.dialogs.export import imprimer_pages
    src = tmp_path / "mixte.pdf"
    c = canvas.Canvas(str(src), pagesize=landscape(A4))
    c.drawString(100, 100, "paysage")
    c.showPage()
    c.setPageSize(A4)
    c.drawString(100, 100, "portrait")
    c.showPage()
    c.setPageSize(landscape(A4))
    c.drawString(100, 100, "paysage")
    c.showPage()
    c.save()
    doc = QPdfDocument()
    doc.load(str(src))
    imp = QPrinter(QPrinter.HighResolution)
    imp.setOutputFormat(QPrinter.PdfFormat)
    sortie = tmp_path / "imprime.pdf"
    imp.setOutputFileName(str(sortie))
    assert imprimer_pages(imp, doc, 1, 3)
    from pypdf import PdfReader
    tailles = [(float(p.mediabox.width), float(p.mediabox.height)) for p in PdfReader(str(sortie)).pages]
    assert [w > h for w, h in tailles] == [True, False, True]

"""Export Word (.docx) des rapports, des bulletins et des listes ; listes en PDF."""
from __future__ import annotations

import base64
import io
import shutil
import subprocess

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pypdf import PdfReader

from epp.core.security import AccessDenied, Session, ensure_roles, load_session, system_session
from epp.domain import repo
from epp.reports import service
from epp.reports.builtin import TEMPLATES

S = system_session()


def _texte(docx_path) -> str:
    d = Document(docx_path)
    morceaux = [p.text for p in d.paragraphs]

    def cellules(tables):
        for t in tables:
            for row in t.rows:
                for c in row.cells:
                    morceaux.extend(p.text for p in c.paragraphs)
                    cellules(c.tables)
    cellules(d.tables)
    return "\n".join(morceaux)


def _params(code, annee):
    if code in ("FEUILLE_NOTES", "COMPOSITION"):
        return {"classe": "CM2", "compo": 2}
    if code == "BULLETIN_EVAL":
        return {"classe": "CE1", "compo": 2}
    if code == "BULLETIN_ANNUEL":
        return {"classe": "CM2"}
    return {}


@pytest.mark.parametrize("code", [t[0] for t in TEMPLATES])
def test_tous_les_rapports_en_word(demo, tmp_path, code):
    db, annee = demo
    f = service.generer(db, S, annee, code, _params(code, annee), tmp_path / "r", "docx")
    assert f.suffix == ".docx"
    d = Document(f)
    texte = _texte(f)
    doc = service.document(db, annee, code, _params(code, annee))
    if not code.startswith("BULLETIN"):
        assert doc.titre.split("–")[0].strip()[:20] in texte
        assert "MINISTERE" in texte.upper()                       # en-tête officiel
        pied = d.sections[0].footer._element.xml
        assert "PAGE" in pied and "NUMPAGES" in pied               # pagination calculée par Word
    for t in d.tables:
        if t.rows and t.rows[0]._tr.trPr is not None and t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None:
            break                                                  # en-tête répété sur chaque page
    else:
        if any(type(b).__name__ == "Table" for b in doc.blocks):
            pytest.fail("aucun en-tête de tableau répété")
    # les valeurs du tableau sont celles du PDF (même document abstrait)
    for b in doc.blocks:
        if type(b).__name__ == "Table" and b.rows:
            premiere = next((r for r in b.rows if r.kind == "data"), None)
            if premiere:
                from epp.reports.engine import fmt
                attendu = [fmt(v, b.types[i]) for i, v in enumerate(premiere.cells) if v not in (None, "")]
                assert all(x.replace(" ", " ") in texte.replace(" ", " ") for x in attendu[:6])
            break


def test_bulletins_word_une_page_par_eleve(demo, tmp_path):
    db, annee = demo
    repo.save_etablissement(db, S, {"bulletin_titre": "École {ecole}", "bulletin_pied": "Le savoir est une lumière"})
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (60, 60), "green").save(buf, "PNG")
    repo.set_logo(db, S, "logo_ecole", buf.getvalue())
    eleves = [r for r in repo.inscriptions(db, annee, "CM2")]
    f = service.generer(db, S, annee, "BULLETIN_ANNUEL", {"classe": "CM2"}, tmp_path / "b.docx")
    texte = _texte(f)
    for e in eleves:
        assert e["nom"] in texte
    cl = annee.classe("CM2")
    for m in {m["libelle"] for ms in annee.grilles[cl.grille_id].matieres.values() for m in ms}:
        assert (m[:1] + m[1:].lower()) in texte
    for attendu in ("ÉCOLE EPP OUANINOU 2", "Le savoir est une lumière", "BULLETIN ANNUEL",
                    "Décision du conseil des maîtres", "Félicitations", "Observations des parents"):
        assert attendu in texte
    xml = Document(f).element.xml
    assert xml.count('w:type="page"') == len(eleves) - 1           # un saut de page entre deux élèves
    assert "<pic:pic" in xml                                        # logo de l'école
    if shutil.which("soffice"):
        r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(f)],
                           capture_output=True, timeout=180)
        assert r.returncode == 0
        assert len(PdfReader(tmp_path / "b.pdf").pages) == len(eleves)


def test_bulletin_classique_word(demo, tmp_path):
    db, annee = demo
    from epp.reports import store
    m = store.get(db, "BULLETIN_EVAL")
    defn = store.definition(db, m["id"])
    for sec in defn["sections"]:
        if sec.get("type") == "bulletin":
            sec["style"] = "classique"
    store.creer(db, S, "Bulletin classique", defn, "Bulletins", code="BUL_CLASSIQUE")
    f = service.generer(db, S, annee, "BUL_CLASSIQUE", {"classe": "CE1", "compo": 1}, tmp_path / "c.docx")
    texte = _texte(f)
    assert "BULLETIN DE NOTES" in texte and "MOYENNE" in texte and "RANG" in texte
    assert repo.inscriptions(db, annee, "CE1")[0]["nom"] in texte


def test_lot_et_fichier_par_classe_en_word(demo, tmp_path):
    db, annee = demo
    f = service.generer_lot_nomme(db, S, annee, "DOSSIER_FIN_ANNEE", {}, tmp_path, "docx")
    assert f.suffix == ".docx"
    texte = _texte(f)
    for titre in ("TABLEAU DE BORD", "LISTE DU PERSONNEL"):
        assert titre in texte
    fichiers = service.un_fichier_par_classe(db, S, annee, "EX_LISTE_CLASSE", {}, tmp_path / "classes", "docx")
    assert len(fichiers) == len(annee.codes) and all(x.suffix == ".docx" for x in fichiers)


@pytest.mark.parametrize("kind", ["eleves", "personnel", "classes", "absences_eleves", "absences_personnel", "notes"])
@pytest.mark.parametrize("fmt", ["xlsx", "pdf", "docx"])
def test_listes_en_excel_pdf_et_word(demo, tmp_path, kind, fmt):
    db, annee = demo
    num = 2 if kind == "notes" else None
    f = service.exporter_liste(db, S, annee, kind, tmp_path / f"{kind}.{fmt}", fmt, num=num)
    assert f.exists() and f.suffix == f".{fmt}"
    premier = {"eleves": repo.inscriptions(db, annee)[0]["nom"], "personnel": repo.personnel(db, annee)[0]["nom"],
               "classes": "CP1", "absences_eleves": repo.absences_eleves(db, annee)[0]["nom"],
               "absences_personnel": None, "notes": repo.inscriptions(db, annee, "CP1")[0]["nom"]}[kind]
    if fmt == "xlsx":
        texte = " ".join(str(c.value) for ws in load_workbook(f).worksheets for row in ws.iter_rows() for c in row)
    elif fmt == "pdf":
        texte = " ".join(p.extract_text() for p in PdfReader(f).pages)
    else:
        texte = _texte(f)
    if premier:
        assert premier in texte
    if kind == "notes" and fmt != "xlsx":
        for code in annee.codes:
            assert f"FEUILLE DE NOTES – {code}" in texte               # toutes les classes, à la suite


def test_droits_export_word(demo, tmp_path):
    db, annee = demo
    enseignant = Session(3, "ens", "Enseignant", "ENSEIGNANT", "Enseignant", {"rapports.generer", "notes.lire"})
    assert service.generer(db, enseignant, annee, "REGISTRE", {}, tmp_path / "r.pdf").exists()
    with pytest.raises(AccessDenied):
        service.generer(db, enseignant, annee, "REGISTRE", {}, tmp_path / "r.docx")
    with pytest.raises(AccessDenied):
        service.exporter_liste(db, enseignant, annee, "eleves", tmp_path / "e", "docx")
    sec = Session(2, "sec", "Secrétariat", "SECRETARIAT", "Secrétariat", {"rapports.generer", "export.word"})
    assert service.generer(db, sec, annee, "REGISTRE", {}, tmp_path / "r.docx").exists()


def test_mise_a_jour_donne_export_word_aux_roles_prevus(db):
    """Base créée avant l'export Word : le droit est ajouté une fois au directeur et au secrétariat."""
    from epp.core.security import create_user
    for code in ("DIRECTEUR", "SECRETARIAT"):
        rid = db.scalar("SELECT id FROM role WHERE code=?", (code,))
        db.execute("DELETE FROM role_permission WHERE role_id=? AND permission='export.word'", (rid,))
    db.execute("DELETE FROM meta WHERE cle='permission_ajoutee:export.word'")
    ensure_roles(db)
    for code, attendu in (("DIRECTEUR", True), ("SECRETARIAT", True), ("ENSEIGNANT", False)):
        uid = create_user(db, f"u_{code.lower()}", code, "MotDePasse-2026!", code)
        assert load_session(db, uid).can("export.word") is attendu


def test_graphique_en_image(demo):
    from epp.reports import engine as E
    from epp.reports.render_docx import image_graphique
    png = image_graphique(E.Chart("Effectifs", ["CP1", "CP2"], [("Garçons", [3, 5]), ("Filles", [4, 0])]))
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    png = image_graphique(E.Chart("Taux", ["CP1"], [("Admis", [0.75])], percent=True))
    assert base64.b64encode(png)

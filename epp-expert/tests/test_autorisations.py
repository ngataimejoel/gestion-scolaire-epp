"""Demandes d'autorisation d'absence : préremplissage, report automatique dans les absences, impression."""
from __future__ import annotations

import pytest
from docx import Document
from pypdf import PdfReader

from epp.core.db import SCHEMA_VERSION, AnneeCloturee
from epp.core.security import AccessDenied, Session, system_session
from epp.domain import autorisations as au
from epp.domain import repo
from epp.domain.repo import ValidationError
from epp.reports import service

S = system_session()


def _agent(db, annee):
    return repo.personnel(db, annee)[1]


def test_preremplissage_par_matricule_ou_recherche(demo):
    db, annee = demo
    p = _agent(db, annee)
    f = au.preremplir(db, annee, p["matricule"])
    assert f["kind"] == "personnel" and f["nom"] == p["nom"] and f["fonction"] == p["fonction"]
    assert f["classe"] == p["classe_tenue"]
    e = repo.inscriptions(db, annee, "CE1")[1]
    f = au.preremplir(db, annee, e["matricule_ecole"].lower())
    assert f["kind"] == "eleve" and f["classe"] == "CE1" and f["demandeur"] and f["enseignant"]
    assert au.preremplir(db, annee, e["nom"])["nom"] == e["nom"]
    with pytest.raises(ValidationError, match="Aucun"):
        au.preremplir(db, annee, "ZZZ-INCONNU")
    with pytest.raises(ValidationError, match="Précisez"):
        au.preremplir(db, annee, "CE1")                     # toute une classe : plusieurs personnes


def test_accord_reporte_dans_les_absences_par_mois(demo):
    db, annee = demo
    f = au.preremplir(db, annee, _agent(db, annee)["matricule"])
    n = len(repo.absences_personnel(db, annee))
    aid = au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-10-29", "date_fin": "2026-11-03",
                                                  "motif": "Deuil familial", "precisions": "Décès du père",
                                                  "interim": "Adjoint", "statut": "ACCORDEE"})
    a = au.get(db, aid)
    assert a["numero"] == "AA-2026-001" and a["jours"] == 4 and a["decide_par"] == S.login
    abs_ = [r for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid]
    assert len(repo.absences_personnel(db, annee)) == n + 2
    assert [(r["date"], r["jours"]) for r in abs_] == [("2026-10-29", 2.0), ("2026-11-02", 2.0)]
    assert all(r["motif"] == "Deuil familial" and r["justifie"] == "OUI" and "AA-2026-001" in r["observations"]
               for r in abs_)
    # le rapport mensuel d'octobre ne compte que les jours d'octobre
    from epp.domain import stats
    oct_ = [r for r in stats.assiduite_personnel(db, annee, 10) if r["nom"] == f["nom"]]
    assert oct_ and oct_[0]["jours"] == 2
    # modifier la période met les absences à jour ; refuser les retire ; supprimer aussi
    au.enregistrer(db, S, annee, aid, a | {"date_fin": "2026-10-30"})
    assert [r["jours"] for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid] == [2.0]
    au.decider(db, S, annee, aid, "REFUSEE", "Période d'évaluation")
    assert not [r for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid]
    au.decider(db, S, annee, aid, "ACCORDEE")
    assert au.get(db, aid)["avis"] == "Période d'évaluation"
    au.supprimer(db, S, aid)
    assert len(repo.absences_personnel(db, annee)) == n


def test_eleve_demi_journee_et_controles(demo):
    db, annee = demo
    f = au.preremplir(db, annee, repo.inscriptions(db, annee, "CP2")[0]["matricule_ecole"])
    aid = au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-11-10", "demi_journee": "MATIN",
                                                  "motif": "Maladie", "statut": "ACCORDEE"})
    r = [x for x in repo.absences_eleves(db, annee) if x["autorisation_id"] == aid]
    assert len(r) == 1 and r[0]["jours"] == 0.5 and r[0]["motif"] == "Maladie"
    base = f | {"motif": "Voyage"}
    with pytest.raises(ValidationError, match="couvre déjà"):
        au.enregistrer(db, S, annee, None, base | {"date_debut": "2026-11-09", "date_fin": "2026-11-12"})
    with pytest.raises(ValidationError, match="postérieure"):
        au.enregistrer(db, S, annee, None, base | {"date_debut": "2026-11-20", "date_fin": "2026-11-19"})
    with pytest.raises(ValidationError, match="année scolaire"):
        au.enregistrer(db, S, annee, None, base | {"date_debut": "2025-11-20"})
    with pytest.raises(ValidationError, match="aucun jour de classe"):
        au.enregistrer(db, S, annee, None, base | {"date_debut": "2026-11-14", "date_fin": "2026-11-15"})
    with pytest.raises(ValidationError, match="motif"):
        au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-11-20", "motif": " "})
    with pytest.raises(ValidationError, match="seul jour"):
        au.enregistrer(db, S, annee, None, base | {"date_debut": "2026-11-20", "date_fin": "2026-11-23",
                                                   "demi_journee": "MATIN"})


def test_secretariat_enregistre_le_directeur_decide(demo):
    db, annee = demo
    sec = Session(2, "sec", "Secrétariat", "SECRETARIAT", "Secrétariat", {"absences.modifier", "absences.lire"})
    f = au.preremplir(db, annee, _agent(db, annee)["matricule"])
    aid = au.enregistrer(db, sec, annee, None, f | {"date_debut": "2026-12-01", "motif": "Formation",
                                                    "statut": "ACCORDEE"})
    assert au.get(db, aid)["statut"] == "EN_ATTENTE"         # sans le droit d'accorder : en attente
    assert not [r for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid]
    with pytest.raises(AccessDenied):
        au.decider(db, sec, annee, aid, "ACCORDEE")
    au.decider(db, S, annee, aid, "ACCORDEE", "Accord")
    assert [r for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid]
    with pytest.raises(ValidationError, match="déjà été traitée"):
        au.enregistrer(db, sec, annee, aid, au.get(db, aid) | {"motif": "Autre"})
    with pytest.raises(AccessDenied):
        au.supprimer(db, sec, aid)
    ens = Session(3, "ens", "Enseignant", "ENSEIGNANT", "Enseignant", {"absences.lire"})
    with pytest.raises(AccessDenied):
        au.enregistrer(db, ens, annee, None, f | {"date_debut": "2026-12-08", "motif": "Formation"})


def test_impression_par_numero_matricule_ou_recherche(demo, tmp_path):
    db, annee = demo
    agent = _agent(db, annee)
    f = au.preremplir(db, annee, agent["matricule"])
    au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-10-12", "motif": "Maladie", "statut": "ACCORDEE"})
    au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-11-16", "date_fin": "2026-11-18",
                                            "motif": "Convocation IEPP", "precisions": "Conférence pédagogique"})
    el = repo.inscriptions(db, annee, "CM1")[0]
    au.enregistrer(db, S, annee, None, au.preremplir(db, annee, el["matricule_ecole"]) |
                   {"date_debut": "2026-11-17", "motif": "Voyage", "lieu": "Abidjan"})
    # par numéro : une fiche
    t = " ".join(p.extract_text() for p in PdfReader(au.imprimer(db, S, annee, "AA-2026-002", tmp_path / "a.pdf")).pages)
    assert "DEMANDE D'AUTORISATION D'ABSENCE" in t and "AA-2026-002" in t and "Conférence pédagogique" in t
    assert agent["nom"] in t and "MINISTERE" in t.upper() and "IEPP" in t
    # personnel : lettre au modèle de l'école, remplie avec la fiche de l'agent
    t1 = " ".join(t.split())
    for attendu in ("Monsieur/Madame le Chef de Circonscription", "Objet : Demande d'autorisation d'absence",
                    f"Je soussigné(e), {agent['nom']}, exerçant en qualité de {agent['fonction']} à l'EPP OUANINOU 2",
                    "s'étend du 16/11/2026 au 18/11/2026, soit une durée de 3 jour(s) de classe",
                    "Convocation IEPP – Conférence pédagogique", "SIGNATURES ET AVIS", "LE DEP / CPPP / CPI",
                    "LE CHEF DE CIRCONSCRIPTION", "DRENA : TOUBA", "IEPP : OUANINOU",
                    "Demande d'autorisation d'absence — Document administratif"):
        assert attendu in t1, attendu
    repo.save_etablissement(db, S, {"autorisation_destinataire": "Monsieur l'Inspecteur",
                                    "autorisation_avis": "LE DIRECTEUR"})
    d = au.imprimer(db, S, annee, "AA-2026-002", tmp_path / "a.docx")
    x = Document(d)
    texte = "\n".join([p.text for p in x.paragraphs] + [c.text for tb in x.tables for row in tb.rows for c in row.cells])
    assert "Monsieur l'Inspecteur" in texte and "LE DIRECTEUR" in texte and agent["nom"] in texte
    assert "SIGNATURES ET AVIS" in texte and "Chef de Circonscription" not in texte
    # par matricule : toutes les demandes de l'agent, une page chacune
    r = PdfReader(au.imprimer(db, S, annee, agent["matricule"], tmp_path / "b.pdf"))
    assert len(r.pages) == 2 and "MINISTERE" in r.pages[1].extract_text().upper()
    # par recherche (motif, date) et en Word
    d = au.imprimer(db, S, annee, "abidjan", tmp_path / "c.docx")
    texte = "\n".join(c.text for t in Document(d).tables for row in t.rows for c in row.cells)
    assert el["nom"] in texte and "Parent / tuteur demandeur" in texte
    assert len(au.trouver(db, annee, "16/11/2026")) == 1
    with pytest.raises(ValidationError, match="Aucune demande"):
        au.imprimer(db, S, annee, "rien du tout", tmp_path / "d.pdf")
    # registre de toutes les demandes
    reg = service.generer(db, S, annee, "AUTORISATIONS", {}, tmp_path / "reg.pdf")
    t = PdfReader(reg).pages[0].extract_text()
    assert "REGISTRE DES AUTORISATIONS D'ABSENCE" in t and "AA-2026-003" in t and "En attente" in t


def test_annee_cloturee(demo):
    db, annee = demo
    f = au.preremplir(db, annee, _agent(db, annee)["matricule"])
    au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-10-12", "motif": "Maladie", "statut": "ACCORDEE"})
    db.execute("UPDATE annee_scolaire SET cloturee=1 WHERE id=?", (annee.id,))
    with pytest.raises(AnneeCloturee):
        au.enregistrer(db, S, annee, None, f | {"date_debut": "2026-10-20", "motif": "Maladie"})


def test_migration_d_une_base_version_3(tmp_path):
    """Une base créée avant les autorisations reçoit la table et les colonnes de liaison."""
    import sqlite3

    from epp.core.db import _AUTORISATION_SQL, open_database
    from epp.core.security import system_session as sys_s
    from epp.demo import creer_ecole_demo
    chemin = tmp_path / "v3.sqlite"
    db = open_database(chemin)
    annee = creer_ecole_demo(db, sys_s())
    db.close()
    con = sqlite3.connect(chemin)
    con.execute("DROP TABLE autorisation_absence")
    for t in ("absence_eleve", "absence_personnel"):
        con.execute(f"ALTER TABLE {t} DROP COLUMN autorisation_id")
    con.execute("UPDATE meta SET valeur='3' WHERE cle='schema_version'")
    con.commit()
    con.close()
    assert _AUTORISATION_SQL
    db = open_database(chemin)
    assert db.meta("schema_version") == str(SCHEMA_VERSION)
    annee = repo.load_annee(db, annee.id)
    f = au.preremplir(db, annee, repo.personnel(db, annee)[0]["matricule"])
    aid = au.enregistrer(db, sys_s(), annee, None, f | {"date_debut": "2026-10-12", "motif": "Maladie",
                                                        "statut": "ACCORDEE"})
    assert [r for r in repo.absences_personnel(db, annee) if r["autorisation_id"] == aid]
    db.close()

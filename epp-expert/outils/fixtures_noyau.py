"""Données de test du noyau Java du téléphone (lecture des tableaux, licence), comparées au code Python.

    python outils/fixtures_noyau.py preparer <dossier>   # fichiers d'essai et licences signées
    java ... TestNoyau <dossier>                          # le téléphone lit et vérifie, écrit ses résultats
    python outils/fixtures_noyau.py verifier <dossier>   # résultats Java == lecture Python (import_eleves19)

Les fichiers .xlsx et .csv sont créés ici ; le .xls (format de l'export national) est un fichier d'essai du
dépôt (android/test-jvm/donnees, noms fictifs) produit par LibreOffice, que l'intégration continue n'a pas.
"""
from __future__ import annotations

import base64
import os
import json
import shutil
import sys
from datetime import date, datetime, timezone
from pathlib import Path

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE))
XLS = RACINE / "android" / "test-jvm" / "donnees" / "eleves_essai.xls"
COLONNES_ELEVES = ["codedren", "codeiep", "iep", "codeecole", "ecole", "matricule", "nom", "prenoms", "sexe", "jour",
                   "mois", "annee", "nationalite", "localite", "codesp", "sp", "pere", "mere", "nacte", "lieuacte",
                   "residence", "niveau"]


def _b64u(b: bytes) -> str:
    return base64.urlsafe_b64encode(b).decode().rstrip("=")


def _empreinte(ident: str) -> str:
    import hashlib
    return hashlib.sha256(f"AKWA-ANDROID|{ident}".encode()).hexdigest()


def preparer(dossier: Path) -> None:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from openpyxl import Workbook

    dossier.mkdir(parents=True, exist_ok=True)
    # ------------------------------------------------------------------ tableaux
    wb = Workbook()
    ws = wb.active
    ws.title = "Élèves"
    ws.append(["Nom", "Prénoms", "Sexe", "Date de naissance", "Classe", "Moyenne", "Redoublant", "Vide"])
    for i in range(60):
        ws.append([f"KOUASSI n°{i}", "Aïcha « Marie »" if i % 2 else "Jean;Baptiste", "F" if i % 2 else "M",
                   date(2015, 1 + i % 12, 1 + i % 28), "CE1" if i % 3 else "CM2", 7.25 + i if i % 4 else 12,
                   bool(i % 5 == 0), None])
    ws.append([])                                  # ligne vide : retirée
    ws.append(["fin", None, None, None, None, None, None, ""])
    wb.create_sheet("Autre").append(["ne pas lire"])
    wb.save(dossier / "essai.xlsx")
    lignes = [["Nom", "Téléphone", "Observation"], ["DIABATÉ Ibrahim", "0707070707", 'dit "le grand"; CP1'],
              ["KONÉ Fatoumata", "05 05 05", "ligne\nsur deux lignes"], ["", "", ""], ["YAO", "", ""]]
    import csv
    with open(dossier / "essai_utf8.csv", "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f, delimiter=";").writerows(lignes)
    with open(dossier / "essai_cp1252.csv", "w", encoding="cp1252", newline="") as f:
        csv.writer(f, delimiter=",").writerows([[c.replace("\n", " ") for c in r] for r in lignes])
    shutil.copy(XLS, dossier / "eleves_essai.xls")

    # ------------------------------------------------------------------ licences (clé de test)
    cle = Ed25519PrivateKey.generate()
    pub = base64.b64encode(cle.public_key().public_bytes(serialization.Encoding.Raw,
                                                         serialization.PublicFormat.Raw)).decode()
    machine = _empreinte("identifiant-android-test")
    maintenant = datetime(2026, 10, 15, 12, 0, tzinfo=timezone.utc)

    def signer(contenu: dict) -> tuple[str, str]:
        brut = json.dumps(contenu, ensure_ascii=False).encode("utf-8")
        return _b64u(brut), _b64u(cle.sign(brut))

    base = {"v": 1, "produit": "GESTION SCOLAIRE EPP EXPERT", "code": "ABCD-EFGH-JKLM", "machine": machine,
            "ecole": "EPP « ESSAI »", "active_le": maintenant.isoformat()}
    lic1, sig1 = signer({**base, "type": "1an", "expire_le": "2027-10-15T12:00:00.000Z"})
    lic_vie, sig_vie = signer({**base, "type": "vie", "expire_le": None})
    lic_echue, sig_echue = signer({**base, "type": "1an", "expire_le": "2026-10-01T00:00:00.000Z"})
    lic_autre, sig_autre = signer({**base, "type": "vie", "expire_le": None, "machine": _empreinte("autre")})
    falsifiee = _b64u(json.dumps({**base, "type": "vie", "expire_le": None}).encode())
    (dossier / "licence.txt").write_text("\n".join([
        pub, machine, str(int(maintenant.timestamp() * 1000)),
        lic1, sig1, lic_vie, sig_vie, lic_echue, sig_echue, lic_autre, sig_autre, falsifiee, sig1,
    ]) + "\n", encoding="utf-8")
    # ------------------------------------------------------------------ activation sans Internet (même clé de test)
    from epp.core import hors_ligne as HL
    graine = cle.private_bytes(serialization.Encoding.Raw, serialization.PrivateFormat.Raw,
                               serialization.NoEncryption())

    def code_hl(empreinte: str, type_: str, expire: date | None, serie: int) -> str:
        d15 = HL.nettoyer(HL.demande(empreinte))[:15]
        charge = HL.charge_utile(type_, expire, serie)
        return HL.grouper(HL.coder(charge + cle.sign(HL.message(d15, charge))))

    (dossier / "hors_ligne.txt").write_text("\n".join([
        pub, machine, str(int(maintenant.timestamp() * 1000)), HL.demande(machine), graine.hex(),
        code_hl(machine, "1an", date(2027, 10, 15), 42), code_hl(machine, "ensvie", None, 7),
        code_hl(_empreinte("autre"), "vie", None, 8),
        HL.ecrire_cles_courtes({1: HL.secret_court(graine, 1)}, masquees=True),
        HL.coder_court(HL.secret_court(graine, 1), 1, HL.demande(machine), "1an", date(2026, 10, 15)),
        HL.coder_court(HL.secret_court(graine, 1), 1, HL.demande(_empreinte("autre")), "vie", date(2026, 10, 15)),
    ]) + "\n", encoding="utf-8")
    # ------------------------------------------------------------------ classeur Excel du téléphone
    (dossier / "classeur.json").write_text(json.dumps({"feuilles": [
        {"nom": "ELEVES", "style": "national", "colonnes": [{"titre": c} for c in COLONNES_ELEVES],
         "lignes": [["35", "000251", "OUANINOU", "E0001", "EPP ESSAI", "00012345A", "KOUASSI", "AÏCHA « MARIE »",
                     "FEMININ", "1", "1", "2014"] + [""] * 9 + ["CP1"],
                    [""] * 6 + ["YAO", "JEAN", "MASCULIN"] + [""] * 12 + ["CE1"]]},
        {"nom": "Résultats CP1", "style": "rapport", "entete": ["MINISTERE DE L'EDUCATION NATIONALE", "IEPP OUANINOU"],
         "droite": ["REPUBLIQUE DE COTE D'IVOIRE", "Année scolaire 2026-2027"], "titre": "RÉSULTATS ANNUELS",
         "sous_titre": "Classe CP1", "paysage": True,
         "colonnes": [{"titre": "N°", "type": "entier"}, {"titre": "Nom et prénoms", "nom": True},
                      {"titre": "Né(e) le", "type": "date"}, {"titre": "MGA", "type": "nombre"},
                      {"titre": "Décision", "type": "centre"}],
         "lignes": [[1, "KOUASSI AÏCHA", "2014-01-01", 6.25, "ADMIS"], [2, "YAO <JEAN> & Cie", None, "4,5", "REDOUBLE"]]},
    ]}, ensure_ascii=False), encoding="utf-8")

    # ------------------------------------------------------------------ synchronisation par fichier
    from epp.synchro import fichier as F
    from epp.synchro import protocole as P
    cle_sync = os.urandom(32)
    (dossier / "synchro.txt").write_text(f"{P.b64(cle_sync)}\n12\nserveur-essai\n", encoding="utf-8")
    rep = {"resultats": [{"uid": "t1", "ok": True}], "instantane": {"ecole": {"ecole_nom": "EPP « ESSAI »"}}}
    msg = P.sceller(cle_sync, P.json_octets(rep), P.aad_reponse(12, 7))
    (dossier / "reponse.akwasync").write_bytes(F._enveloppe("reponse", "serveur-essai", 12, 7, ecole="EPP « ESSAI »",
                                                            message=P.b64(msg)))
    (dossier / "reponse_rejeu.akwasync").write_bytes(F._enveloppe(
        "reponse", "serveur-essai", 12, 5, erreur="Requête déjà reçue (rejeu refusé).", code="REJEU"))
    assert _empreinte("x") != machine
    print(f"données de test écrites dans {dossier}")


def verifier(dossier: Path) -> None:
    from epp.io.import_eleves19 import lire_fichier
    erreurs = 0
    for f in ("essai.xlsx", "essai_utf8.csv", "essai_cp1252.csv", "eleves_essai.xls"):
        attendu = lire_fichier(dossier / f)
        obtenu = json.loads((dossier / (f + ".java.json")).read_text(encoding="utf-8"))
        if obtenu != attendu:
            erreurs += 1
            print(f"ÉCHEC {f} : {len(obtenu)} lignes lues par le téléphone, {len(attendu)} par l'ordinateur")
            for i, (a, b) in enumerate(zip(attendu, obtenu)):
                if a != b:
                    print(f"  ligne {i + 1} :\n    ordinateur {a}\n    téléphone  {b}")
                    break
        else:
            print(f"ok {f} : {len(attendu)} lignes identiques")
    erreurs += _verifier_classeur(dossier) + _verifier_synchro(dossier) + _verifier_hors_ligne(dossier)
    sys.exit(1 if erreurs else 0)


def _verifier_hors_ligne(dossier: Path) -> int:
    """Code signé par le générateur Android (Java), vérifié par l'application Windows (Python)."""
    from epp.core import hors_ligne as HL
    lignes = (dossier / "hors_ligne.txt").read_text(encoding="utf-8").split("\n")
    c = HL.verifier((dossier / "hors_ligne.java.txt").read_text(encoding="utf-8").strip(), lignes[1], [lignes[0]])
    if c != {"type": "vie", "expire_le": None, "code": "HL-123456"}:
        print(f"ÉCHEC code hors ligne signé par le téléphone : {c}")
        return 1
    court = (dossier / "hors_ligne_court.java.txt").read_text(encoding="utf-8").strip()
    c = HL.verifier_court(court, lignes[1], HL.lire_cles_courtes(lignes[8], masquees=True))
    if not c or c["type"] != "ensvie" or len(court) != 29:
        print(f"ÉCHEC code court du générateur Android : {court} → {c}")
        return 1
    print("ok codes hors ligne du téléphone (long et court) acceptés par l'ordinateur")
    return 0


def _verifier_classeur(dossier: Path) -> int:
    """Le classeur du téléphone s'ouvre dans Excel (openpyxl) avec la présentation de l'application Windows."""
    from openpyxl import load_workbook
    from epp.io.import_eleves19 import COLONNES, lire_fichier
    wb = load_workbook(dossier / "classeur.xlsx")
    el, res = wb["ELEVES"], wb["Résultats CP1"]
    pb = []
    if [c.value for c in el[1]] != COLONNES:
        pb.append("en-tête ELEVES")
    if el["A1"].fill.fgColor.rgb != "FF0E6646" or not el["A1"].font.b or el["F2"].number_format != "@":
        pb.append("présentation ELEVES")
    if el["F2"].value != "00012345A" or el["H2"].value != "AÏCHA « MARIE »" or el.freeze_panes != "A2":
        pb.append("valeurs ELEVES")
    lu = lire_fichier(dossier / "classeur.xlsx")
    if len(lu) != 3 or lu[1][5] != "00012345A":
        pb.append("relecture par l'import de l'ordinateur")
    cel = {c.value: c for row in res.iter_rows() for c in row if c.value is not None}
    if "RÉSULTATS ANNUELS" not in cel or cel["RÉSULTATS ANNUELS"].font.color.rgb != "FF1F4E79":
        pb.append("titre du rapport")
    if cel["Nom et prénoms"].fill.fgColor.rgb != "FF1F4E79" or not cel["KOUASSI AÏCHA"].font.b:
        pb.append("en-têtes et noms du rapport")
    if cel[6.25].number_format != "0.00" or cel[4.5].number_format != "0.00" or "YAO <JEAN> & Cie" not in cel:
        pb.append("nombres du rapport")
    if res.page_setup.orientation != "landscape" or res.print_title_rows is None:
        pb.append("mise en page")
    import datetime as dt
    if dt.datetime(2014, 1, 1) not in cel or cel[dt.datetime(2014, 1, 1)].number_format != "dd/mm/yyyy":
        pb.append("date")
    print(("ÉCHEC classeur Excel : " + ", ".join(pb)) if pb else "ok classeur Excel identique à l'application Windows")
    return 1 if pb else 0


def _verifier_synchro(dossier: Path) -> int:
    from epp.synchro import protocole as P
    cle = P.unb64((dossier / "synchro.txt").read_text().split()[0])
    pb = []
    if json.loads((dossier / "reponse.java.json").read_text(encoding="utf-8"))["instantane"]["ecole"]["ecole_nom"] \
            != "EPP « ESSAI »":
        pb.append("réponse lue par le téléphone")
    e = json.loads((dossier / "demande.akwasync").read_text(encoding="utf-8"))
    req = json.loads(P.ouvrir(cle, P.unb64(e["message"]), P.aad_requete(12)))
    if (e["sens"], e["seq"], e["serveur_id"], e["appareil"], req["seq"], req["ops"][0]["valeur"]) != \
            ("demande", 8, "serveur-essai", "TECNO « Spark »", 8, 7.5):
        pb.append(f"demande du téléphone {e} {req}")
    print(("ÉCHEC synchronisation par fichier : " + ", ".join(pb)) if pb else
          "ok synchronisation par fichier : demande et réponse compatibles")
    return 1 if pb else 0


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("preparer", "verifier"):
        sys.exit(__doc__)
    (preparer if sys.argv[1] == "preparer" else verifier)(Path(sys.argv[2]))

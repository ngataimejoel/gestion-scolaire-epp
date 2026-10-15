"""Fiche scolaire de l'élève en Word (.docx) : même contenu que le PDF, modifiable dans Word."""
from __future__ import annotations

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from epp.domain import calculs, fiche_scolaire
from epp.reports import engine as E
from epp.reports.render_docx import (_bordures, _bordures_cellule, _ecrire, _format_para, _hauteur, _image,
                                     _largeurs, _ombre, _para, _run, entete_officiel_word)

ORDINAUX = ("1ère", "2ème", "3ème", "4ème", "5ème", "6ème", "7ème", "8ème")
POINTS = "………………"
C = WD_ALIGN_PARAGRAPH.CENTER
G = WD_ALIGN_PARAGRAPH.LEFT
D = WD_ALIGN_PARAGRAPH.RIGHT


def _date(v) -> str:
    d = calculs.to_date(v) if v else None
    return d.strftime("%d/%m/%Y") if d else ""


def _case(coche: bool) -> str:
    return "☒" if coche else "☐"


class FicheWord:
    def __init__(self, rendu, b: E.FicheScolaire):
        self.r, self.d, self.o, self.photo = rendu, b.data, b.options, b.photo
        self.w = rendu.d
        self.avail = rendu.avail

    def _champs(self, p, couples, taille=10):
        """Plusieurs « LIBELLÉ : valeur » sur une ligne (valeur absente : pointillés)."""
        for i, (lab, val) in enumerate(couples):
            if i:
                _run(p, "   ", taille)
            _run(p, f"{lab} ", taille, gras=True)
            _run(p, val if val else POINTS, taille)

    def rendre(self) -> None:
        d, w, avail = self.d, self.w, self.avail
        doc = self.r.doc
        entete_officiel_word(w, avail, *E.entete_officiel(doc))
        _para(w, "", apres=2)
        # titre encadré avec ombre
        t = w.add_table(1, 1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures(t, None)
        _largeurs(t, [90])
        cell = t.cell(0, 0)
        _ecrire(cell, self.o.get("titre_fiche") or "FICHE SCOLAIRE", 22, gras=True, align=C)
        cell.paragraphs[0].runs[0].underline = True
        _bordures_cellule(cell, top=(16, "000000"), left=(16, "000000"), bottom=(36, "A6A6A6"),
                          right=(36, "A6A6A6"))
        _para(w, "", apres=4)

        # photo | sexe, classe, matricule | n° de fiche + QR
        qr = self.o.get("qr", True)
        t = w.add_table(1, 3)
        _bordures(t, None)
        _largeurs(t, [38, avail - 38 - 48, 48])
        c0, c1, c2 = t.rows[0].cells
        p = _ecrire(c0, align=C)
        if not (self.o.get("photo", True) and self.photo and _image(p, self.photo, 33, 40.5)):
            _run(p, "PHOTO\n", 11, gras=True, couleur="6B6B6B")
            _run(p, "33 × 40,5 mm", 8, italique=True, couleur="6B6B6B")
        _bordures_cellule(c0, top=(8, "000000"), bottom=(8, "000000"), left=(8, "000000"), right=(8, "000000"))
        _hauteur(t.rows[0], 42)
        p = _ecrire(c1, valign="haut")
        _run(p, "Sexe : ", 10.5, gras=True)
        _run(p, f"{_case(d.get('sexe') == 'M')} Masculin   {_case(d.get('sexe') == 'F')} Féminin", 10.5, gras=True)
        p = c1.add_paragraph()
        _format_para(p, 8, 0)
        _run(p, "Classe : ", 10.5, gras=True)
        _run(p, "  ".join(f"{_case(d.get('niveau') == n)} {n}" for n in fiche_scolaire.NIVEAUX), 9.5, gras=True)
        p = c1.add_paragraph()
        _format_para(p, 8, 0)
        _run(p, f"Classe : {d.get('classe', '')}    Matricule école : {d.get('matricule_ecole', '')}", 9.5)
        p = c1.add_paragraph()
        _format_para(p, 8, 0)
        _run(p, "NUMÉRO MATRICULE : ", 10.5, gras=True)
        _run(p, d.get("matricule_desps") or POINTS, 11, gras=True)
        p = _ecrire(c2, f"FICHE N° {d.get('fiche_no', '')}", 10, gras=True, align=C, valign="haut")
        if qr:
            p = c2.add_paragraph()
            _format_para(p, 4, 0, C)
            _image(p, fiche_scolaire.qr_png(d.get("qr") or ""), 28, 28)
            p = c2.add_paragraph()
            _format_para(p, 0, 0, C)
            _run(p, "Code QR : école et élève", 7, italique=True, couleur="6B6B6B")
        _para(w, "", apres=4)

        # identité
        for couples in ((("NOM :", d.get("nom_famille")),), (("PRÉNOM(S) :", d.get("prenoms")),),
                        (("Né(e) le :", _date(d.get("date_naissance"))), ("à :", d.get("lieu_naissance")),
                         ("S/P :", d.get("sous_prefecture"))),
                        (("Nationalité :", d.get("nationalite")), ("Groupe sanguin :", d.get("groupe_sanguin")),
                         ("Allergies / handicap :", d.get("allergies"))),
                        (("Acte de naissance / Jugement supplétif N° :", d.get("acte_num")),
                         ("Du :", _date(d.get("acte_date"))), ("Délivré à :", d.get("centre_etat_civil")))):
            p = _para(w, "", apres=5)
            self._champs(p, couples, 10.5 if len(couples) == 1 else 10)

        # parents
        parents = d.get("parents") or []
        t = w.add_table(6, 3)
        _bordures(t, None)
        _largeurs(t, [avail / 3] * 3)
        for j, par in enumerate(parents[:3]):
            p = _ecrire(t.cell(0, j), align=C)
            r = _run(p, par["titre"], 10.5, gras=True)
            r.underline = True
            for i, (lab, cle) in enumerate((("NOM :", "nom"), ("PRÉNOM(S) :", "prenoms"), ("DOMICILE :", "domicile"),
                                           ("CONTACT :", "contact"), ("PROFESSION :", "profession")), start=1):
                p = _ecrire(t.cell(i, j))
                self._champs(p, ((lab, par.get(cle)),), 9.5)
        for row in t.rows:
            _hauteur(row, 6.5)
        _para(w, "", apres=4)

        # scolarité
        sco = list(d.get("scolarite") or [])[-8:]
        t = w.add_table(3, 9)
        _bordures(t, 6, "000000")
        _largeurs(t, [42] + [(avail - 42) / 8] * 8)
        _ecrire(t.cell(0, 0), "ANNÉE SCOLAIRE", 8.5, gras=True)
        _ecrire(t.cell(1, 0), "Année de scolarité", 8.5, italique=True, align=C)
        _ecrire(t.cell(2, 0), "Niveau de cours\nMoyenne annuelle\nEffectif de la classe", 8, gras=True)
        for k in range(8):
            s = sco[k] if k < len(sco) else None
            _ecrire(t.cell(0, k + 1), s["annee"].replace("-", "/") if s else "20…/20…", 8, gras=True, align=C)
            _ecrire(t.cell(1, k + 1), ORDINAUX[k], 8.5, italique=True, align=C)
            if s:
                moy = E.fmt(s["moyenne"], "decimal") if s.get("moyenne") is not None else "–"
                vals = f"{s['niveau'] or s['classe']}\n{moy}\n{s.get('effectif') or '–'}"
            else:
                vals = "……\n……\n……"
            _ecrire(t.cell(2, k + 1), vals, 8, align=C)
        for j in range(9):
            _ombre(t.cell(0, j), "A6A6A6")
        _para(w, "", apres=4)

        # école
        w3 = (("ÉCOLE :", d.get("ecole")), ("VILLE / SOUS-PRÉFECTURE :", d.get("ville")))
        self._champs(_para(w, "", apres=5), w3, 10.5)
        self._champs(_para(w, "", apres=5), (("Fréquentée du :", _date(d.get("frequente_du"))),
                                             ("au :", _date(d.get("frequente_au"))),
                                             ("Matricule :", d.get("matricule_ecole"))), 10.5)
        self._champs(_para(w, "", apres=5), (("Motif de départ :", d.get("motif_depart")),), 10.5)
        v = doc.variables
        _para(w, f"Fait à {d.get('ville') or v.get('localite', '')}, le {v.get('date_jour', '')}", 10.5,
              italique=True, align=D, apres=4)

        # signatures
        t = w.add_table(1, 3)
        _bordures(t, None)
        _largeurs(t, [avail / 3] * 3)
        noms = ["", d.get("enseignant") or "", d.get("directeur") or v.get("directeur", "")]
        for j, titre in enumerate(("Le Parent / Tuteur", "L'Enseignant(e)", "Le Directeur de l'école")):
            cell = t.cell(0, j)
            p = _ecrire(cell, align=C, valign="haut")
            r = _run(p, titre, 9.5, gras=True)
            r.underline = True
            if noms[j]:
                p = cell.add_paragraph()
                _format_para(p, 50, 0, C)
                _run(p, noms[j], 8.5)
            _bordures_cellule(cell, top=(8, "000000"), bottom=(8, "000000"), left=(8, "000000"),
                              right=(8, "000000"))
        _hauteur(t.rows[0], 26)

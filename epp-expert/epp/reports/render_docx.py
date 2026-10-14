"""Rendu Word (.docx, python-docx) d'un document produit par le moteur de rapports.

Même contenu que le PDF : en-tête officiel avec logos, titres, tableaux (en-têtes fusionnés répétés sur
chaque page, groupes, sous-totaux, totaux), indicateurs, graphiques (image), zones à rédiger, signatures
et bulletins (modèle coloré de l'école ou modèle classique, une page par élève). Le document reste
modifiable dans Word ; le pied de page porte le texte du modèle, la pagination et la date d'édition.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from epp.reports import engine as E

POLICE = "Arial"
BLEU = "1F4E79"
BLEU_CLAIR = "DDEBF7"
GRIS = "F2F2F2"
GRIS_FONCE = "D9D9D9"
ROSE = "FCE4EC"
GRIS_TRANSF = "E7E6E6"
BORD = "808080"
MARGE_G = MARGE_D = 8
MARGE_H = MARGE_B = 12
ALIGN = {"gauche": WD_ALIGN_PARAGRAPH.LEFT, "centre": WD_ALIGN_PARAGRAPH.CENTER,
         "droite": WD_ALIGN_PARAGRAPH.RIGHT}


# =================================================================== outils python-docx
def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#").upper())


def _ombre(cell, hexa: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for x in tcpr.findall(qn("w:shd")):
        tcpr.remove(x)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa.lstrip("#").upper())
    tcpr.append(shd)


def _bordures_cellule(cell, **cotes) -> None:
    """cotes : top/bottom/left/right = (taille en 1/8 pt, couleur) ou None pour aucune bordure."""
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcpr.append(b)
    for cote, val in cotes.items():
        el = b.find(qn(f"w:{cote}"))
        if el is None:
            el = OxmlElement(f"w:{cote}")
            b.append(el)
        if val is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val[0]))
            el.set(qn("w:color"), val[1].lstrip("#").upper())


def _bordures(table, taille: int | None = 4, couleur: str = BORD, interieur: bool = True) -> None:
    tblpr = table._tbl.tblPr
    for x in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(x)
    b = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{cote}")
        if taille is None or (not interieur and cote.startswith("inside")):
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(taille))
            el.set(qn("w:color"), couleur)
        b.append(el)
    tblpr.append(b)


def _marges_cellules(table, haut=0.4, gauche=1.2) -> None:
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for cote, v in (("top", haut), ("bottom", haut), ("left", gauche), ("right", gauche)):
        el = OxmlElement(f"w:{cote}")
        el.set(qn("w:w"), str(int(v * 56.7)))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tblpr.append(m)


def _largeurs(table, largeurs_mm: list[float]) -> None:
    table.autofit = False
    tblpr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), largeurs_mm):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for cell, w in zip(row.cells, largeurs_mm):
            cell.width = Mm(w)


def _entete_repetee(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _insecable(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _hauteur(row, mm_: float, exacte: bool = False) -> None:
    row.height = Mm(mm_)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exacte else WD_ROW_HEIGHT_RULE.AT_LEAST


def _para(conteneur, texte="", taille=9, gras=False, italique=False, couleur=None, align=None,
          avant=0, apres=0, police=None):
    p = conteneur.add_paragraph()
    _format_para(p, avant, apres, align)
    if texte != "":
        _run(p, texte, taille, gras, italique, couleur, police)
    return p


def _format_para(p, avant=0, apres=0, align=None, interligne=1.0):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(avant), Pt(apres)
    pf.line_spacing = interligne
    if align is not None:
        p.alignment = align


def _run(p, texte, taille=9, gras=False, italique=False, couleur=None, police=None):
    lignes = str(texte).split("\n")
    r = None
    for i, l in enumerate(lignes):
        r = p.add_run(l)
        r.font.size = Pt(taille)
        r.bold, r.italic = gras, italique
        if couleur:
            r.font.color.rgb = _rgb(couleur)
        if police:
            r.font.name = police
            r._element.rPr.rFonts.set(qn("w:eastAsia"), police)
            r._element.rPr.rFonts.set(qn("w:hAnsi"), police)
            r._element.rPr.rFonts.set(qn("w:cs"), police)
        if i < len(lignes) - 1:
            r.add_break()
    return r


def _ecrire(cell, texte="", taille=9, gras=False, italique=False, couleur=None, align=None, valign="centre"):
    """Remplace le contenu d'une cellule par un paragraphe."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    _format_para(p, 0, 0, align)
    if texte not in (None, ""):
        _run(p, texte, taille, gras, italique, couleur)
    cell.vertical_alignment = {"haut": WD_ALIGN_VERTICAL.TOP, "bas": WD_ALIGN_VERTICAL.BOTTOM}.get(
        valign, WD_ALIGN_VERTICAL.CENTER)
    return p


def _image(conteneur_p, data: bytes, max_w: float, max_h: float) -> bool:
    """Ajoute une image (dans un paragraphe) en conservant ses proportions ; False si illisible."""
    try:
        from PIL import Image as PILImage
        with PILImage.open(io.BytesIO(data)) as im:
            iw, ih = im.size
    except Exception:
        return False
    k = min(max_w / iw, max_h / ih)
    try:
        conteneur_p.add_run().add_picture(io.BytesIO(data), width=Mm(iw * k), height=Mm(ih * k))
    except Exception:
        return False
    return True


def _champ(p, instr: str, taille=7.5, couleur="404040"):
    """Champ Word (PAGE, NUMPAGES) : la valeur est calculée à l'ouverture ou à l'impression."""
    r = p.add_run()
    r.font.size = Pt(taille)
    r.font.color.rgb = _rgb(couleur)
    for kind, texte in (("begin", None), (None, f" {instr} "), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText" if instr in texte else "w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = texte
        r._element.append(el)


def _saut_de_page(d) -> None:
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# =================================================================== rendu
class _Rendu:
    def __init__(self, doc: E.Document):
        self.doc = doc
        self.monnaie = doc.variables.get("monnaie", "FCFA")
        self.d = DocxDocument()
        st = self.d.styles["Normal"]
        st.font.name = POLICE
        st.font.size = Pt(9)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
        st.paragraph_format.space_after = Pt(0)
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        if doc.orientation == "paysage":
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Mm(297), Mm(210)
        sec.left_margin, sec.right_margin = Mm(MARGE_G), Mm(MARGE_D)
        sec.top_margin, sec.bottom_margin = Mm(MARGE_H), Mm(MARGE_B)
        sec.header_distance, sec.footer_distance = Mm(5), Mm(4)
        self.avail = (297 if doc.orientation == "paysage" else 210) - MARGE_G - MARGE_D
        cp = self.d.core_properties
        cp.title = doc.titre or doc.nom
        cp.author = doc.variables.get("ecole", "")
        cp.subject = doc.sous_titre or ""
        cp.comments = "GESTION SCOLAIRE EPP EXPERT"
        self._premier = True        # le premier paragraphe vide du document est réutilisé

    # ------------------------------------------------------------------ pied de page
    def pied(self) -> None:
        pied = self.doc.pied or {}
        v = self.doc.variables
        f = self.d.sections[0].footer
        t = f.add_table(1, 3, Mm(self.avail))
        _bordures(t, None)
        _largeurs(t, [self.avail * 0.45, self.avail * 0.2, self.avail * 0.35])
        _ecrire(t.cell(0, 0), E.subst(pied.get("texte") or "", v)[:120], 7.5, couleur="404040")
        if pied.get("numerotation", True):
            p = _ecrire(t.cell(0, 1), "Page ", 7.5, couleur="404040", align=WD_ALIGN_PARAGRAPH.CENTER)
            _champ(p, "PAGE")
            _run(p, " / ", 7.5, couleur="404040")
            _champ(p, "NUMPAGES")
        if pied.get("date_edition", True):
            _ecrire(t.cell(0, 2), f"Édité le {v.get('date_jour', '')}", 7.5, couleur="404040",
                    align=WD_ALIGN_PARAGRAPH.RIGHT)
        if f.paragraphs and not f.paragraphs[0].text:
            f.paragraphs[0]._element.getparent().remove(f.paragraphs[0]._element)

    # ------------------------------------------------------------------ en-tête officiel
    def officiel(self, force: bool = False) -> None:
        doc = self.doc
        if not force and not doc.entete.get("officiel", True):
            return
        gauche, droite, logos = E.entete_officiel(doc)
        centre = max(self.avail * 0.16, min(self.avail * 0.34, 24 * len(logos) + 6))
        cote = (self.avail - centre) / 2
        t = self.d.add_table(1, 3)
        _bordures(t, None)
        _largeurs(t, [cote, centre, cote])
        for cell, lignes in ((t.cell(0, 0), gauche), (t.cell(0, 2), droite)):
            _ecrire(cell, "\n".join(lignes), 8, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER, valign="haut")
        p = _ecrire(t.cell(0, 1), align=WD_ALIGN_PARAGRAPH.CENTER, valign="haut")
        for i, data in enumerate(logos):
            if i:
                p.add_run("  ")
            _image(p, data, centre / max(len(logos), 1) - 3, 18)
        if doc.entete.get("texte"):
            _para(self.d, E.subst(doc.entete["texte"], doc.variables), 8, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(self.d, "", apres=4)

    # ------------------------------------------------------------------ blocs simples
    def titre(self, titre: str, sous: str) -> None:
        if titre:
            p = _para(self.d, titre, 13, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
        if sous:
            p = _para(self.d, sous, 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
        _para(self.d, "", apres=4).paragraph_format.keep_with_next = True

    def paragraphe(self, b: E.Paragraph) -> None:
        if b.style == "encadre":
            t = self.d.add_table(1, 1)
            _bordures(t, 4, "808080")
            _largeurs(t, [self.avail])
            _ombre(t.cell(0, 0), GRIS)
            _ecrire(t.cell(0, 0), b.text, 9)
            _para(self.d, "", apres=2)
            return
        styles = {"titre2": dict(taille=10, gras=True, couleur=BLEU), "note": dict(taille=8, italique=True)}
        p = _para(self.d, b.text, apres=4, **styles.get(b.style, {}))
        if b.style == "titre2" or (b.style == "normal" and len(b.text) < 160):
            p.paragraph_format.keep_with_next = True

    def indicateurs(self, b: E.KeyValues) -> None:
        if b.title:
            _para(self.d, b.title, 10, gras=True, couleur=BLEU, apres=3).paragraph_format.keep_with_next = True
        if not b.items:
            return
        par_ligne = 4
        n = (len(b.items) + par_ligne - 1) // par_ligne
        t = self.d.add_table(n, par_ligne)
        _bordures(t, 4, GRIS_FONCE)
        _largeurs(t, [self.avail / par_ligne] * par_ligne)
        for i, (k, val) in enumerate(b.items):
            cell = t.cell(i // par_ligne, i % par_ligne)
            _ombre(cell, "FAFCFE")
            p = _ecrire(cell, val, 14, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_before = Pt(3)
            _para(cell, k, 7.5, couleur="404040", align=WD_ALIGN_PARAGRAPH.CENTER, apres=3)
        for row in t.rows:
            _insecable(row)
        _para(self.d, "", apres=4)

    def graphique(self, b: E.Chart) -> None:
        png = image_graphique(b, min(self.avail, 170))
        p = _para(self.d, align=WD_ALIGN_PARAGRAPH.CENTER, apres=4)
        p.paragraph_format.keep_together = True
        p.add_run().add_picture(io.BytesIO(png), width=Mm(min(self.avail, 170)))

    def zone(self, b: E.TextZone) -> None:
        _para(self.d, b.title, 9, gras=True, apres=2).paragraph_format.keep_with_next = True
        t = self.d.add_table(1, 1)
        _bordures(t, 4, "808080")
        _largeurs(t, [self.avail])
        _ecrire(t.cell(0, 0), b.text or "", 9, valign="haut")
        if not b.text:
            _hauteur(t.rows[0], max(b.lines * 4.2, 5))
        _insecable(t.rows[0])
        _para(self.d, "", apres=3)

    def signature(self, b: E.Signature) -> None:
        _para(self.d, "", apres=4).paragraph_format.keep_with_next = True
        t = self.d.add_table(1, 2)
        _bordures(t, None)
        _largeurs(t, [self.avail * 0.55, self.avail * 0.45])
        _ecrire(t.cell(0, 0), b.lieu_date, 9, valign="haut")
        c = t.cell(0, 1)
        _ecrire(c, b.fonction, 9, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER, valign="haut")
        _para(c, "", apres=0).paragraph_format.space_before = Mm(14)
        _para(c, b.nom, 9, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _insecable(t.rows[0])

    # ------------------------------------------------------------------ lettre administrative
    def lettre(self, b: E.Lettre) -> None:
        def segments(p, segs, taille=9.5):
            for s, gras in segs:
                _run(p, s, taille, gras)

        for i, ligne in enumerate(b.entete):
            _para(self.d, ligne, 9 if i < 2 else 8.5, gras=i != 1, align=WD_ALIGN_PARAGRAPH.CENTER,
                  apres=8 if i == 1 else 2)
        _para(self.d, "", apres=6)
        if b.references:
            t = self.d.add_table(len(b.references), 2)
            _bordures(t, None)
            _largeurs(t, [self.avail * 0.62, self.avail * 0.38])
            for r, paire in enumerate(b.references):
                for c, x in enumerate(paire):
                    p = _ecrire(t.cell(r, c), "", 8.5)
                    if x:
                        segments(p, [(x.split(":")[0].rstrip() + " :", True), (":".join(x.split(":")[1:]), False)], 8.5)
            _para(self.d, "", apres=10)
        _para(self.d, b.titre, 15, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER, apres=8)
        if b.destinataire:
            t = self.d.add_table(1, 1)
            _bordures(t, 6, "9FB6CD")
            _largeurs(t, [self.avail])
            c = t.cell(0, 0)
            _ombre(c, "F2F6FA")
            _ecrire(c, b.destinataire[0], 9.5, gras=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            for x in b.destinataire[1:]:
                _para(c, x, 9.5)
            _para(self.d, "", apres=8)
        for style, segs in b.corps:
            if style == "intitule":
                p = _para(self.d, "", avant=2, apres=3)
                segments(p, [(s, True) for s, _ in segs])
            elif style == "droite":
                p = _para(self.d, "", align=WD_ALIGN_PARAGRAPH.RIGHT, avant=4, apres=6)
                segments(p, segs)
            else:
                p = _para(self.d, "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, apres=6)
                p.paragraph_format.line_spacing = 1.15
                segments(p, segs)
        if b.signatures:
            n = len(b.signatures)
            p = _para(self.d, b.titre_signatures, 10, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER,
                      avant=4, apres=4)
            p.paragraph_format.keep_with_next = True
            t = self.d.add_table(2, n)
            _bordures(t, 4, "9FB6CD")
            _largeurs(t, [self.avail / n] * n)
            for i, (tete, lignes) in enumerate(b.signatures):
                _ecrire(t.cell(0, i), tete, 7.5, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER)
                _ombre(t.cell(0, i), "F2F6FA")
                c = t.cell(1, i)
                _ecrire(c, lignes[0] if lignes else "", 7.5, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                for x in lignes[1:]:
                    _para(c, x, 7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            _hauteur(t.rows[0], 9)
            _hauteur(t.rows[1], 34)
            for row in t.rows:
                _insecable(row)

    # ------------------------------------------------------------------ tableaux
    def tableau(self, t: E.Table) -> None:
        ncols = len(t.types)
        from epp.reports.render_pdf import _grid
        hgrid, spans = _grid(t.header, ncols)
        corps = []
        for row in t.rows:
            if row.kind == "group":
                corps.append([str(row.cells[0] or "")] + [""] * (ncols - 1))
                continue
            rtype = row.style.get("type")
            cells = []
            for i, v in enumerate(row.cells):
                typ = t.types[i]
                if rtype and i > 0 and not isinstance(v, str):
                    typ = rtype
                cells.append(E.fmt(v, typ, self.monnaie))
            corps.append(cells)
        taille = 9 if ncols <= 8 else 8 if ncols <= 12 else 7 if ncols <= 18 else 6.5
        largeurs = self._largeurs_colonnes(t, hgrid, spans, corps, taille)
        if t.title:
            _para(self.d, t.title, 10, gras=True, couleur=BLEU, apres=3).paragraph_format.keep_with_next = True
        morceaux: list[list[int]] = [[]]
        for k, row in enumerate(t.rows):
            if t.page_break_groups and row.kind == "group" and morceaux[-1]:
                morceaux.append([])
            morceaux[-1].append(k)
        for n, indices in enumerate(morceaux):
            if n:
                _saut_de_page(self.d)
            self._table_docx(t, hgrid, spans, corps, indices, largeurs, taille)
        _para(self.d, "", apres=4)

    def _largeurs_colonnes(self, t: E.Table, hgrid, spans, corps, taille) -> list[float]:
        marge = 2.6                     # marges intérieures de la cellule (mm)
        nat, mini = [], []
        for i in range(len(t.types)):
            textes = [r[i] for r, row in zip(corps, t.rows) if row.kind != "group"]
            larges = {(c0, r0) for (c0, r0), (c1, _) in spans if c1 > c0}    # en-têtes sur plusieurs colonnes
            entetes = [hgrid[r][i] for r in range(len(hgrid)) if (i, r) not in larges]
            long_ = max([_mesure(x, taille) for x in textes] + [0])
            mots = [w for x in textes for w in _mots(x)] + [w for x in entetes for w in str(x).split()]
            mot = max([_mesure(w, taille, gras=True) for w in mots] + [3])
            tete = max([_mesure(x, taille, gras=True) for x in entetes] + [0])
            nat.append(max(long_, tete * 0.6) + marge)
            mini.append(mot + marge)
        fixe = [w for w in t.widths]
        out = [fixe[i] if fixe[i] else max(nat[i], mini[i]) for i in range(len(nat))]
        total = sum(out)
        if total <= self.avail:
            libres = [i for i in range(len(out)) if not fixe[i]] or list(range(len(out)))
            bonus = (self.avail - total) / len(libres) if total < self.avail * 0.7 else 0
            out = [w + (bonus if i in libres else 0) for i, w in enumerate(out)]
            return out
        # trop large : on réduit au prorata en gardant au moins la largeur du mot le plus long
        exces = total - self.avail
        souple = [max(out[i] - mini[i], 0) for i in range(len(out))]
        s = sum(souple)
        if s >= exces:
            return [out[i] - souple[i] * exces / s for i in range(len(out))]
        return [w * self.avail / total for w in mini] if sum(mini) > self.avail else \
            [mini[i] + (self.avail - sum(mini)) * (out[i] - mini[i]) / max(total - sum(mini), 1e-6)
             for i in range(len(out))]

    def _table_docx(self, t, hgrid, spans, corps, indices, largeurs, taille) -> None:
        ncols = len(t.types)
        nh = len(hgrid)
        tab = self.d.add_table(nh + len(indices), ncols)
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures(tab, 4, BORD)
        _marges_cellules(tab, 0.3, 1.0)
        _largeurs(tab, largeurs)
        for r in range(nh):
            row = tab.rows[r]
            _entete_repetee(row)
            _insecable(row)
            for c in range(ncols):
                cell = row.cells[c]
                _ombre(cell, BLEU)
                _ecrire(cell, hgrid[r][c], taille, gras=True, couleur="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        for (c0, r0), (c1, r1) in spans:
            a, b = tab.cell(r0, c0), tab.cell(r1, c1)
            texte = hgrid[r0][c0]
            m = a.merge(b)
            _ecrire(m, texte, taille, gras=True, couleur="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            for p in m.paragraphs[1:]:
                p._element.getparent().remove(p._element)
        for n, k in enumerate(indices):
            row, txt = t.rows[k], corps[k]
            tr = tab.rows[nh + n]
            _insecable(tr)
            fond = {"subtotal": GRIS, "total": GRIS_FONCE}.get(row.kind)
            if row.kind == "data":
                fond = {"ABANDON": ROSE, "TRANSFERE": GRIS_TRANSF}.get(row.style.get("statut"))
            if row.kind == "group":
                m = tr.cells[0].merge(tr.cells[-1])
                _ombre(m, BLEU_CLAIR)
                _ecrire(m, txt[0], taille, gras=True)
                for p in m.paragraphs[1:]:
                    p._element.getparent().remove(p._element)
                continue
            gras = row.kind in ("subtotal", "total")
            for i, x in enumerate(txt):
                cell = tr.cells[i]
                if " " in x and len(_mots(x)) == 1:
                    x = x.replace(" ", "\u00a0")          # numéro ou nombre : jamais coupé
                _ecrire(cell, x, taille, gras=gras or (t.first_col_header and i == 0),
                        align=ALIGN.get(t.aligns[i], WD_ALIGN_PARAGRAPH.LEFT))
                if fond:
                    _ombre(cell, fond)
                elif t.first_col_header and i == 0:
                    _ombre(cell, GRIS)

    # ------------------------------------------------------------------ bulletins
    def bulletin(self, b: E.Bulletin) -> None:
        if b.options.get("style", "colore") == "colore":
            BulletinWord(self, b).rendre()
        else:
            self._bulletin_classique(b)

    def _bulletin_classique(self, b: E.Bulletin) -> None:
        d, o, doc = b.data, b.options, self.doc
        avail = self.avail
        self.officiel(force=True)
        annuel = d["type"] == "annuel"
        titre = o.get("titre_bulletin") or ("BULLETIN ANNUEL DE NOTES" if annuel else "BULLETIN DE NOTES")
        sous = f"Année scolaire {doc.variables.get('annee', '')}" if annuel else \
            f"{d['evaluation']}" + (f" du {E.fmt(d.get('date_evaluation'), 'date')}" if d.get("date_evaluation")
                                    else "") + f" – Année scolaire {doc.variables.get('annee', '')}"
        _para(self.d, titre, 14, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(self.d, sous, 9.5, align=WD_ALIGN_PARAGRAPH.CENTER, apres=6)
        ident = [("Nom et prénoms", d.get("nom")), ("Matricule école", d.get("matricule_ecole")),
                 ("Matricule DESPS", d.get("matricule_desps") or "–"), ("Classe", d.get("classe")),
                 ("Sexe", {"M": "Masculin", "F": "Féminin"}.get(d.get("sexe"), "")),
                 ("Né(e) le", E.fmt(d.get("date_naissance"), "date") or "–"),
                 ("Redoublant", d.get("redoublant") or "NON"), ("Effectif classé", str(d.get("effectif") or 0))]
        photo = o.get("photo", True)
        larg = avail - (34 if photo else 0)
        t = self.d.add_table(4, 5 if photo else 4)
        _bordures(t, None)
        _largeurs(t, [larg * 0.18, larg * 0.32, larg * 0.18, larg * 0.32] + ([34] if photo else []))
        for i, (k, v) in enumerate(ident):
            r, c = i // 2, (i % 2) * 2
            _ecrire(t.cell(r, c), k, 8, couleur="404040")
            _ecrire(t.cell(r, c + 1), str(v or ""), 10, gras=True)
            _bordures_cellule(t.cell(r, c), bottom=(2, "C0C0C0"))
            _bordures_cellule(t.cell(r, c + 1), bottom=(2, "C0C0C0"))
        if photo:
            m = t.cell(0, 4).merge(t.cell(3, 4))
            for p in m.paragraphs[1:]:
                p._element.getparent().remove(p._element)
            p = _ecrire(m, align=WD_ALIGN_PARAGRAPH.CENTER)
            if not (b.photo and _image(p, b.photo, 30, 38)):
                _run(p, "PHOTO", 8, couleur="808080")
            _bordures_cellule(m, top=(4, "808080"), bottom=(4, "808080"), left=(4, "808080"), right=(4, "808080"))
        _para(self.d, "", apres=6)
        app = o.get("appreciations", True)
        if not annuel:
            coef = d["mode"] == "COEF"
            tete = ["MATIÈRE", "NOTE", "COEF." if coef else "SUR"] + (["APPRÉCIATION"] if app else [])
            lignes = []
            for n in d["notes"]:
                note = "Abs." if not d["present"] else (E.fmt(n["note"], "decimal") if n["note"] is not None else "–")
                lignes.append([n["matiere"], note, E.fmt(n["coef"] if coef else n["max"], "decimal")] +
                              ([n["appreciation"] or ""] if app else []))
            lignes.append(["TOTAL", E.fmt(d["total"], "decimal") if d["total"] is not None else "–",
                           "" if coef else E.fmt(d["total_max"], "decimal")] + ([""] if app else []))
            larg = [0.42, 0.14, 0.12] + ([0.32] if app else [])
            self._petit_tableau(tete, lignes, larg, total=True)
            res = [("MOYENNE", f"{E.fmt(d['moyenne'], 'decimal')} / {E.fmt(d['echelle'], 'entier')}"
                    if d["moyenne"] is not None else "–"),
                   ("RANG", f"{_rang(d['rang'])} sur {d['effectif']}" if d.get("rang") else "–"),
                   ("APPRÉCIATION", d.get("appreciation") or "–"), ("RÉSULTAT", d.get("resultat") or "–")]
        else:
            echelle = d["bareme"]
            tete = ["ÉVALUATION", "DATE", f"MOYENNE (/{E.fmt(echelle, 'entier')})"] + (["APPRÉCIATION"] if app else [])
            lignes = []
            for e in d["evaluations"]:
                m = "Abs. (0)" if not e["present"] else (E.fmt(e["moyenne"], "decimal")
                                                        if e["moyenne"] is not None else "–")
                lignes.append([e["evaluation"], E.fmt(e["date"], "date"), m] + ([e["appreciation"] or ""] if app else []))
            if d.get("regle_mga") == "STANDARD":
                lignes.append(["Moyenne des 3 compositions", "",
                               E.fmt(d.get("moy_compos"), "decimal") if d.get("moy_compos") is not None else "–"] +
                              ([""] if app else []))
            self._petit_tableau(tete, lignes, [0.40, 0.16, 0.16] + ([0.28] if app else []))
            regle = "MGA = (moyenne des 3 compositions + 2 × composition de passage) ÷ 3" \
                if d.get("regle_mga") == "STANDARD" else "MGA = moyenne des 4 évaluations"
            _para(self.d, regle, 7.5, italique=True, couleur="404040", avant=2, apres=4)
            suite = d.get("classe_suivante")
            dec = d.get("decision") or "–"
            if dec == "ADMIS":
                dec = f"ADMIS(E) en {suite}" if suite else "ADMIS(E) – entrée en 6ème"
            else:
                dec = {"EXCLU": "EXCLU(E)", "TRANSFERE": "TRANSFÉRÉ(E)"}.get(dec, dec)
            res = [("MOYENNE GÉNÉRALE ANNUELLE", f"{E.fmt(d['mga'], 'decimal')} / {E.fmt(echelle, 'entier')}"
                    if d.get("mga") is not None else "–"),
                   ("RANG", f"{_rang(d.get('rang'))} sur {d['effectif']}" if d.get("rang") else "–"),
                   ("DÉCISION", dec), ("OBSERVATION", d.get("observation") or "–")]
        t = self.d.add_table(2, 4)
        _bordures(t, 4, BLEU)
        _largeurs(t, [avail / 4] * 4)
        for i, (k, v) in enumerate(res):
            _ombre(t.cell(0, i), BLEU_CLAIR)
            _ecrire(t.cell(0, i), k, 8, couleur="404040", align=WD_ALIGN_PARAGRAPH.CENTER)
            _ecrire(t.cell(1, i), v, 12, gras=True, couleur=BLEU, align=WD_ALIGN_PARAGRAPH.CENTER)
        _hauteur(t.rows[1], 9)
        _para(self.d, "", apres=3)
        if o.get("statistiques_classe", True):
            ech = d.get("echelle") if not annuel else d.get("bareme")
            stat = (f"Classe : moyenne {E.fmt(d.get('moy_classe'), 'decimal') or '–'}  •  plus forte "
                    f"{E.fmt(d.get('moy_max'), 'decimal') or '–'}  •  plus faible "
                    f"{E.fmt(d.get('moy_min'), 'decimal') or '–'}  (sur {E.fmt(ech, 'entier')})  •  seuil "
                    f"d'admission {E.fmt(d.get('seuil'), 'decimal')}")
            _para(self.d, stat, 8.5, couleur="303030", align=WD_ALIGN_PARAGRAPH.CENTER, apres=6)
        if o.get("observations", True):
            self.zone(E.TextZone("Observations de l'enseignant(e) :", "", 3))
        sig = o.get("signatures") or []
        if sig:
            lieu = f"Fait à {doc.variables.get('localite', '')}, le {doc.variables.get('date_edition', '')}"
            _para(self.d, lieu, 9, apres=3)
            t = self.d.add_table(2, len(sig))
            _bordures(t, None)
            _largeurs(t, [avail / len(sig)] * len(sig))
            for i, x in enumerate(sig):
                _ecrire(t.cell(0, i), x, 9, gras=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _hauteur(t.rows[1], 18)

    def _petit_tableau(self, tete, lignes, proportions, total=False) -> None:
        t = self.d.add_table(1 + len(lignes), len(tete))
        _bordures(t, 4, BORD)
        s = sum(proportions)
        _largeurs(t, [self.avail * p / s for p in proportions])
        for i, h in enumerate(tete):
            _ombre(t.cell(0, i), BLEU)
            _ecrire(t.cell(0, i), h, 9, gras=True, couleur="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        _entete_repetee(t.rows[0])
        for r, ligne in enumerate(lignes, start=1):
            der = total and r == len(lignes)
            for i, x in enumerate(ligne):
                _ecrire(t.cell(r, i), x, 9.5, gras=der and i < 2,
                        align=WD_ALIGN_PARAGRAPH.LEFT if i in (0, 3) else WD_ALIGN_PARAGRAPH.CENTER)
                if der:
                    _ombre(t.cell(r, i), GRIS_FONCE)
            _hauteur(t.rows[r], 6)

    # ------------------------------------------------------------------ document
    def rendre(self) -> DocxDocument:
        doc = self.doc
        bulletins_seuls = all(isinstance(b, (E.Bulletin, E.PageBreak)) for b in doc.blocks) and doc.blocks
        if doc.blocks and all(isinstance(b, (E.Lettre, E.PageBreak)) for b in doc.blocks):
            for sec in self.d.sections:          # lettre : marges d'un courrier administratif
                sec.left_margin = sec.right_margin = Mm(MARGE_G + 16)
            self.avail -= 32
        if not bulletins_seuls:
            self.officiel()
        dernier_saut = True
        premier_bulletin = True
        for b in doc.blocks:
            if isinstance(b, tuple) and b[0] == "TITRE":
                self.titre(b[1], b[2])
            elif isinstance(b, E.Table):
                self.tableau(b)
            elif isinstance(b, E.Paragraph):
                if not b.text:
                    continue
                self.paragraphe(b)
            elif isinstance(b, E.KeyValues):
                self.indicateurs(b)
            elif isinstance(b, E.Chart):
                self.graphique(b)
            elif isinstance(b, E.TextZone):
                self.zone(b)
            elif isinstance(b, E.Signature):
                self.signature(b)
            elif isinstance(b, E.EnteteOfficiel):
                self.officiel(force=True)
            elif isinstance(b, E.Lettre):
                self.lettre(b)
            elif isinstance(b, E.Bulletin):
                if not premier_bulletin and not dernier_saut:
                    _saut_de_page(self.d)
                premier_bulletin = False
                self.bulletin(b)
            elif isinstance(b, E.PageBreak):
                if not dernier_saut:
                    _saut_de_page(self.d)
                    dernier_saut = True
                continue
            dernier_saut = False
        self._nettoyer()
        if not bulletins_seuls:
            self.pied()
        return self.d

    def _nettoyer(self) -> None:
        """Retire le paragraphe vide initial et les sauts de page en fin de document."""
        body = self.d.element.body
        premier = body[0] if len(body) else None
        if premier is not None and premier.tag == qn("w:p") and not "".join(premier.itertext()).strip() \
                and not premier.findall(".//" + qn("w:br")) and not premier.findall(".//" + qn("w:drawing")):
            body.remove(premier)
        while True:
            enfants = [x for x in body if x.tag in (qn("w:p"), qn("w:tbl"))]
            if not enfants or enfants[-1].tag != qn("w:p"):
                break
            p = enfants[-1]
            if "".join(p.itertext()).strip() or p.findall(".//" + qn("w:drawing")):
                break
            body.remove(p)
        blocs = [x for x in body if x.tag in (qn("w:p"), qn("w:tbl"))]
        if not blocs or blocs[-1].tag == qn("w:tbl"):
            # Word attend un paragraphe après un tableau final
            p = OxmlElement("w:p")
            sect = body.find(qn("w:sectPr"))
            if sect is not None:
                sect.addprevious(p)
            else:
                body.append(p)


_POLICES: dict = {}


def _mesure(texte: str, taille: float, gras: bool = False) -> float:
    """Largeur (mm) d'un texte en Arial (Liberation Sans, mêmes dimensions) à la taille donnée."""
    from PIL import ImageFont

    from epp.core.paths import resource_path
    cle = gras
    if cle not in _POLICES:
        nom = "LiberationSans-Bold.ttf" if gras else "LiberationSans-Regular.ttf"
        try:
            _POLICES[cle] = ImageFont.truetype(str(resource_path("resources", "fonts") / nom), 100)
        except OSError:
            _POLICES[cle] = None
    f = _POLICES[cle]
    if f is None:
        return len(str(texte)) * taille * 0.21
    return f.getlength(str(texte)) / 100 * taille * 0.3528


def _mots(texte: str) -> list[str]:
    """Mots insécables : un numéro de téléphone ou un nombre n'est jamais coupé."""
    from epp.reports.render_pdf import _words
    return _words(str(texte))


def _rang(r) -> str:
    if r is None:
        return "–"
    return "1er" if r == 1 else f"{r}e"


# =================================================================== bulletin coloré
class BulletinWord:
    """Bulletin coloré de l'école dans Word : mêmes textes et valeurs que le PDF (bulletin_colore.contenu)."""

    def __init__(self, r: _Rendu, b: E.Bulletin):
        from epp.reports.bulletin_colore import C, contenu
        self.r, self.b, self.d = r, b, r.d
        self.K = contenu(r.doc, b)
        self.C = {k: v.hexval()[2:].upper() for k, v in C.items()}
        self.T = self.K["couleurs"]
        self.W = r.avail

    def rendre(self) -> None:
        K, C, d = self.K, self.C, self.d
        if K["mention"]:
            _para(d, K["mention"], 6.5, gras=True, couleur=C["gris_txt"], align=WD_ALIGN_PARAGRAPH.CENTER, apres=2)
        self._entete()
        self._identite()
        self._disciplines()
        if K["stats"]:
            _para(d, K["stats"], 8, italique=True, couleur=C["texte"], align=WD_ALIGN_PARAGRAPH.CENTER, avant=2)
        _para(d, "", apres=3)
        self._bas()
        self._bandeau()

    def _entete(self) -> None:
        K, C, W = self.K, self.C, self.W
        g, dr = 32, 44
        t = self.d.add_table(1, 3)
        _bordures(t, None)
        _largeurs(t, [g, W - g - dr, dr])
        logos = K["logos"]
        p = _ecrire(t.cell(0, 0), align=WD_ALIGN_PARAGRAPH.CENTER)
        if not (logos and _image(p, logos[0], 30, 29)):
            _run(p, "★", 30, couleur=C["jaune"])
        c = t.cell(0, 1)
        _ecrire(c, K["titre"], 22 if len(K["titre"]) < 26 else 17, gras=True, couleur=C["marine"],
                align=WD_ALIGN_PARAGRAPH.CENTER)
        if K["slogan"]:
            _para(c, K["slogan"], 10, gras=True, italique=True, couleur=C["bleu_moyen"],
                  align=WD_ALIGN_PARAGRAPH.CENTER, apres=2)
        larg = min(110, W - g - dr - 4)
        trait = c.add_table(1, 3)
        trait.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures(trait, None)
        _largeurs(trait, [larg / 3] * 3)
        for i, coul in enumerate(("bleu_moyen", "jaune", "vert")):
            _ombre(trait.cell(0, i), C[coul])
            _ecrire(trait.cell(0, i), "", 2)
        _hauteur(trait.rows[0], 1.4, exacte=True)
        _para(c, "", 3)
        band = c.add_table(1, 1)
        band.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures(band, None)
        _largeurs(band, [min(108, W - g - dr - 4)])
        bc = band.cell(0, 0)
        T = self.T
        _ombre(bc, T["bande_fond"])
        if T["bande_bord"]:
            b = (12, T["bande_bord"])
            _bordures_cellule(bc, top=b, bottom=b, left=b, right=b)
        bt = T["bande_texte"].lstrip("#")
        _ecrire(bc, K["nom_bulletin"], 18, gras=True, couleur=bt, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(bc, f"Année scolaire : {K['annee'].replace('-', ' - ')}", 10, gras=True, couleur=bt,
              align=WD_ALIGN_PARAGRAPH.CENTER)
        c.paragraphs[-1].paragraph_format.space_after = Pt(0)
        droite = t.cell(0, 2)
        p = _ecrire(droite, align=WD_ALIGN_PARAGRAPH.CENTER, valign="haut")
        if len(logos) > 1:
            _image(p, logos[1], 40, 15)
        if K["devise"]:
            _para(droite, K["devise"], 10 if len(logos) < 2 else 9, gras=True, italique=True, couleur=C["marine"],
                  align=WD_ALIGN_PARAGRAPH.CENTER, avant=4)
            pp = _para(droite, "〰〰〰", 9, couleur=C["jaune"], align=WD_ALIGN_PARAGRAPH.CENTER)
            pp.paragraph_format.space_before = Pt(0)
        _para(self.d, "", 4)

    def _identite(self) -> None:
        K, C, W = self.K, self.C, self.W
        photo = K["photo"]
        wp = 24 if photo else 0
        xg = (W - wp) * 0.56
        t = self.d.add_table(1, 3 if photo else 2)
        _bordures(t, 8, C["bord"], interieur=False)
        _largeurs(t, [xg, W - wp - xg] + ([wp] if photo else []))
        for cell in t.rows[0].cells:
            _ombre(cell, self.T["identite"])
        _bordures_cellule(t.cell(0, 0), right=(4, C["bord"]))

        def champ(cell, premier, paires):
            p = _ecrire(cell, valign="haut") if premier else _para(cell)
            _format_para(p, 3, 1)
            for i, (lab, val) in enumerate(paires):
                if i:
                    _run(p, "     ", 9)
                _run(p, lab + " ", 8.5, gras=True, couleur=C["texte"])
                _run(p, val, 9.5, couleur=C["marine"])
            return p

        g = t.cell(0, 0)
        champ(g, True, [("Élève :", K["eleve"])])
        champ(g, False, [("Né(e) le :", K["naissance"]), ("Âge :", K["age"])])
        champ(g, False, [("Classe :", K["classe"]), ("Enseignant(e) :", K["enseignant"])])
        dcell = t.cell(0, 1)
        champ(dcell, True, [("Établissement :", K["ecole"])])
        champ(dcell, False, [("Année scolaire :", K["annee"])])
        p = champ(dcell, False, [(K["label_periode"], "")])
        for lib, coche in K["choix_periode"]:
            _run(p, ("☑ " if coche else "☐ "), 9, couleur=C["marine"], police="Segoe UI Symbol")
            _run(p, lib + "   ", 8, gras=coche, couleur=C["texte"])
        if photo:
            pc = t.cell(0, 2)
            pp = _ecrire(pc, align=WD_ALIGN_PARAGRAPH.CENTER)
            if not (self.b.photo and _image(pp, self.b.photo, 20, 22)):
                _run(pp, "PHOTO", 7, couleur=C["gris_txt"])
                _bordures_cellule(pc, top=(4, C["bord"]), bottom=(4, C["bord"]), left=(4, C["bord"]),
                                  right=(4, C["bord"]))
                _ombre(pc, "FFFFFF")
        _hauteur(t.rows[0], 24)
        _para(self.d, "", 4)

    def _disciplines(self) -> None:
        K, C, W = self.K, self.C, self.W
        entetes = list(K["entetes"]) or [("–", "ciel")]
        np_ = len(entetes)
        w_no, w_moy = 10, 25
        w_p = min(27, (W - w_no - w_moy - 55) / np_)
        w_disc = W - w_no - w_moy - w_p * np_
        heads = [("N°", "bleu"), ("DISCIPLINES", "bleu")] + entetes + [(K["titre_moyenne"], "marine")]
        n_lignes = len(K["lignes"]) + len(K["totaux"])
        t = self.d.add_table(1 + n_lignes, len(heads))
        _bordures(t, None)
        _marges_cellules(t, 0.3, 1.2)
        _largeurs(t, [w_no, w_disc] + [w_p] * np_ + [w_moy])
        T = self.T
        blanc = (8, "FFFFFF") if T["separateur"] == "#FFFFFF" else (4, T["separateur"])
        for i, (lab, coul) in enumerate(heads):
            cell = t.cell(0, i)
            fond, texte, trait = T["tete"][coul]
            _ombre(cell, fond)
            _ecrire(cell, lab, 9, gras=True, couleur=texte.lstrip("#"), align=WD_ALIGN_PARAGRAPH.CENTER)
            _bordures_cellule(cell, left=blanc, right=blanc, **({"top": (24, trait)} if trait else {}))
        _hauteur(t.rows[0], 10)
        _entete_repetee(t.rows[0])
        h = max(5.5, min(9.0, 150 / max(n_lignes, 1)))
        for k, m in enumerate(K["lignes"], start=1):
            tr = t.rows[k]
            fond = T["ligne_a"] if k % 2 else T["ligne_b"]
            vals = [str(k), None] + [x or "" for x in m["valeurs"]] + [m["moyenne"]]
            if not K["entetes"]:
                vals.insert(2, "")
            for i, x in enumerate(vals):
                cell = tr.cells[i]
                _ombre(cell, fond)
                _bordures_cellule(cell, left=blanc, right=blanc, bottom=(2, C["bord"]))
                if i == 1:
                    p = _ecrire(cell, m["libelle"], 9.5, couleur=C["texte"])
                    _run(p, f"  (/{m['max']})", 7, couleur=C["gris_txt"])
                else:
                    gras = i in (0, len(vals) - 1)
                    _ecrire(cell, x, 10 if i else 9.5, gras=gras, couleur=C["marine"] if gras else C["texte"],
                            align=WD_ALIGN_PARAGRAPH.CENTER)
            _hauteur(tr, h)
            _insecable(tr)
        for j, tot in enumerate(K["totaux"]):
            tr = t.rows[1 + len(K["lignes"]) + j]
            vals = ["", tot["libelle"]] + [x or "" for x in tot["valeurs"]] + [tot["fin"]]
            if not K["entetes"]:
                vals.insert(2, "")
            for i, x in enumerate(vals):
                cell = tr.cells[i]
                _ombre(cell, T["pied_tab"])
                _bordures_cellule(cell, left=blanc, right=blanc, bottom=(2, C["bord"]))
                _ecrire(cell, x, 10.5, gras=True, couleur=C["marine"],
                        align=WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER)
            _hauteur(tr, h)
            _insecable(tr)
        tblpr = t._tbl.tblPr
        b = tblpr.find(qn("w:tblBorders"))
        for cote in ("top", "left", "bottom", "right"):
            el = b.find(qn(f"w:{cote}"))
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "10")
            el.set(qn("w:color"), C["bleu_moyen"])

    def _cadre(self, conteneur, largeur, titre, coul_titre, fond_titre, bord, lignes_corps, taille_titre=10.5):
        t = conteneur.add_table(2, 1)
        _bordures(t, 8, bord, interieur=False)
        _largeurs(t, [largeur])
        h = t.cell(0, 0)
        _ombre(h, fond_titre)
        if fond_titre.lstrip("#").upper() == "FFFFFF":       # sans aplat : un filet sépare le titre
            _bordures_cellule(h, bottom=(4, bord))
        _ecrire(h, titre, taille_titre, gras=True, couleur=coul_titre)
        _hauteur(t.rows[0], 6.5)
        corps = t.cell(1, 0)
        _ecrire(corps, valign="haut")
        for i, (texte, kw) in enumerate(lignes_corps):
            p = corps.paragraphs[0] if i == 0 else corps.add_paragraph()
            _format_para(p, 3, 0)
            if texte:
                _run(p, texte, **kw)
        return t

    def _bas(self) -> None:
        K, C, W = self.K, self.C, self.W
        gw = W * 0.485
        rw = W - gw - W * 0.03
        t = self.d.add_table(1, 3)
        _bordures(t, None)
        _largeurs(t, [gw, W * 0.03, rw])
        _marges_cellules(t, 0, 0)
        g = t.cell(0, 0)
        _ecrire(g, "", 1, valign="haut")
        pts = "." * 95
        corps = [(l, dict(taille=10, gras=True, couleur=C["vert_fonce"])) for l in [K["appreciation"]] if l]
        corps += [(pts, dict(taille=8, couleur=C["gris"]))] * (6 - len(corps))
        self._cadre(g, gw, "Appréciation générale", self.T["app_texte"].lstrip("#"), self.T["app_fond"], C["vert"],
                    corps)
        _para(g, "", 1, apres=3)
        self._cadre(g, gw, "Observations des parents", C["marine"], self.T["obs_fond"], C["bord"],
                    [(pts, dict(taille=8, couleur=C["gris"]))] * 5, 10)
        dcell = t.cell(0, 2)
        _ecrire(dcell, "", 1, valign="haut")
        bil = dcell.add_table(2, 2)
        _bordures(bil, 8, C["jaune"], interieur=False)
        lw = rw * 0.52
        _largeurs(bil, [lw, rw - lw])
        tete = bil.cell(0, 0).merge(bil.cell(0, 1))
        _ombre(tete, self.T["bilan_fond"])
        if not self.T["aplats"]:
            _bordures_cellule(tete, bottom=(4, C["jaune"]))
        _ecrire(tete, K["titre_bilan"], 11, gras=True, couleur=C["marine"])
        _hauteur(bil.rows[0], 7)
        items = bil.cell(1, 0)
        _ecrire(items, valign="haut")
        _bordures_cellule(items, right=(4, C["bord"]))
        icones = {"fel": ("✔", "vert"), "enc": ("★", "orange"), "prog": ("◉", "bleu_moyen"), "suivi": ("▲", "rouge")}
        contours = {"fel": "○", "enc": "☆", "prog": "○", "suivi": "△"}      # choix non retenus, sans aplat
        for i, (k, lab) in enumerate(K["bilan"]):
            p = items.paragraphs[0] if i == 0 else items.add_paragraph()
            _format_para(p, 4, 2)
            actif = K["bilan_choix"] == k
            sym, coul = icones[k]
            if not (actif or self.T["aplats"]):
                sym = contours[k]
            _run(p, sym + "  ", 11, gras=True, couleur=C[coul] if actif else C["gris"], police="Segoe UI Symbol")
            _run(p, lab, 10, gras=actif, couleur=C["marine"] if actif else C["gris_txt"])
        dec = bil.cell(1, 1)
        _ecrire(dec, K["titre_decision"], 9, gras=True, couleur=C["marine"], align=WD_ALIGN_PARAGRAPH.CENTER,
                valign="haut").paragraph_format.space_before = Pt(4)
        boite = dec.add_table(1, 1)
        boite.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordures(boite, None if self.T["aplats"] else 4, C["bord"])
        _largeurs(boite, [rw - lw - 5])
        _ombre(boite.cell(0, 0), self.T["decision_fond"])
        _ecrire(boite.cell(0, 0), K["decision"], 11, gras=True, couleur=C["bleu_moyen"],
                align=WD_ALIGN_PARAGRAPH.CENTER)
        _hauteur(boite.rows[0], 13)
        _para(dec, "", 2)
        _para(dcell, "", 1, apres=3)
        sigs = K["signatures"]
        st = dcell.add_table(2, len(sigs))
        _bordures(st, None)
        _largeurs(st, [rw / len(sigs)] * len(sigs))
        for i, (lab, nom) in enumerate(sigs):
            c = st.cell(0, i)
            _ecrire(c, "Signature", 8, gras=True, couleur=C["texte"], align=WD_ALIGN_PARAGRAPH.CENTER, valign="haut")
            _para(c, lab, 8, couleur=C["texte"], align=WD_ALIGN_PARAGRAPH.CENTER)
            b = st.cell(1, i)
            _ecrire(b, "." * 22, 8, couleur=C["gris"], align=WD_ALIGN_PARAGRAPH.CENTER, valign="bas")
            if nom:
                _para(b, nom, 7, couleur=C["gris_txt"], align=WD_ALIGN_PARAGRAPH.CENTER)
            if i:
                for cell in (c, b):
                    _bordures_cellule(cell, left=(4, C["bord"]))
        _hauteur(st.rows[1], 24)
        _para(self.d, "", 4)

    def _bandeau(self) -> None:
        K, C, W = self.K, self.C, self.W
        t = self.d.add_table(1, 1)
        _bordures(t, None)
        _largeurs(t, [W])
        c = t.cell(0, 0)
        leger = self.T["pied"] == "leger"
        if leger:
            _bordures_cellule(c, top=(18, C["jaune"]), bottom=(8, C["marine"]))
        else:
            _ombre(c, C["marine"])
            _bordures_cellule(c, top=(24, C["jaune"]))
        if K["pied"]:
            _ecrire(c, f"—  {K['pied']}  —", 12, gras=True, italique=True, couleur=C["marine"] if leger else "FFFFFF",
                    align=WD_ALIGN_PARAGRAPH.CENTER)
        _hauteur(t.rows[0], 11)
        _insecable(t.rows[0])


# =================================================================== graphique (image)
def image_graphique(b: E.Chart, largeur_mm: float = 170) -> bytes:
    """Diagramme en barres (PNG) pour Word : mêmes données et couleurs que le PDF."""
    from PIL import Image as PILImage, ImageDraw, ImageFont

    from epp.core.paths import resource_path
    k = 6                                    # pixels par mm (≈ 150 ppp)
    W, H = int(largeur_mm * k), int(72 * k)
    im = PILImage.new("RGB", (W, H), "white")
    dr = ImageDraw.Draw(im)
    fonts = resource_path("resources", "fonts")
    try:
        f = ImageFont.truetype(str(fonts / "LiberationSans-Regular.ttf"), int(2.6 * k))
        fb = ImageFont.truetype(str(fonts / "LiberationSans-Bold.ttf"), int(3.5 * k))
    except OSError:
        f = fb = ImageFont.load_default()
    palette = ["#1F4E79", "#C55A11", "#548235"]
    dr.text((0, 0), b.title, font=fb, fill="#1F4E79")
    x0, y0, x1, y1 = 14 * k, 8 * k, W - 40 * k, H - 10 * k
    vals = [v or 0 for _, s in b.series for v in s] or [0]
    vmax = max(vals) or 1
    pas = _pas_axe(vmax)
    vmax = pas * max(1, -(-vmax // pas))
    nb = int(round(vmax / pas))
    for i in range(nb + 1):
        v = pas * i
        y = y1 - (y1 - y0) * v / vmax
        dr.line([(x0, y), (x1, y)], fill="#E0E0E0", width=1)
        lab = f"{v * 100:.0f} %" if b.percent else (f"{v:g}" if v < 1000 else f"{v:,.0f}".replace(",", " "))
        dr.text((x0 - 2 * k, y), lab, font=f, fill="#404040", anchor="rm")
    dr.line([(x0, y0), (x0, y1), (x1, y1)], fill="#404040", width=2)
    n = max(len(b.categories), 1)
    ns = max(len(b.series), 1)
    cw = (x1 - x0) / n
    bw = cw * 0.75 / ns
    for i, cat in enumerate(b.categories):
        for j, (_, s) in enumerate(b.series):
            v = s[i] if i < len(s) and s[i] is not None else 0
            xa = x0 + i * cw + cw * 0.125 + j * bw
            ya = y1 - (y1 - y0) * v / vmax
            dr.rectangle([xa, ya, xa + bw - 1, y1], fill=palette[j % 3])
        dr.text((x0 + i * cw + cw / 2, y1 + 1.5 * k), str(cat), font=f, fill="#202020", anchor="mt")
    for j, (nom, _) in enumerate(b.series):
        y = y0 + j * 5 * k
        dr.rectangle([x1 + 4 * k, y, x1 + 7 * k, y + 3 * k], fill=palette[j % 3])
        dr.text((x1 + 9 * k, y + 1.5 * k), str(nom), font=f, fill="#202020", anchor="lm")
    out = io.BytesIO()
    im.save(out, "PNG", dpi=(150, 150))
    return out.getvalue()


def _pas_axe(vmax: float) -> float:
    import math
    brut = vmax / 5
    e = 10 ** math.floor(math.log10(brut)) if brut > 0 else 1
    for m in (1, 2, 2.5, 5, 10):
        if brut <= m * e:
            return m * e
    return 10 * e


# =================================================================== points d'entrée
def render(doc: E.Document, out: str | Path | io.BytesIO) -> None:
    d = _Rendu(doc).rendre()
    d.save(str(out) if isinstance(out, Path) else out)


def render_bytes(doc: E.Document) -> bytes:
    buf = io.BytesIO()
    render(doc, buf)
    return buf.getvalue()

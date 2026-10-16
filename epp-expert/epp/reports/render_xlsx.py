"""Rendu Excel (openpyxl) d'un document produit par le moteur de rapports."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from epp.reports import engine as E

THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
GROUP_FILL = PatternFill("solid", fgColor="DDEBF7")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
TOT_FILL = PatternFill("solid", fgColor="D9D9D9")
ROSE_FILL = PatternFill("solid", fgColor="FCE4EC")
NUMFMT = {"entier": "#,##0", "decimal": "0.00", "pourcent": "0.0%", "monnaie": '#,##0 "FCFA"',
          "date": "DD/MM/YYYY"}
H_ALIGN = {"gauche": "left", "centre": "center", "droite": "right"}


def _excel_value(v, typ):
    if v is None:
        return None
    if typ == "date":
        from epp.domain.calculs import to_date
        return to_date(v) or v
    return v


def _entete_droite(ws, droite_rows: list, logos: list[bytes], maxcol: int, ligne_logos: int = 1) -> None:
    """Colonne de droite de l'en-tête officiel (alignée sur la dernière colonne) et logos de l'IEPP et de
    l'école sous les lignes de gauche (comme le prototype)."""
    last = max(maxcol, 4)
    for row, line in droite_rows:
        c = ws.cell(row, last, line)
        c.font = Font(bold=True, size=9)
        c.alignment = Alignment(horizontal="right")
    if not logos:
        return
    import io
    from openpyxl.drawing.image import Image as XLImage
    x = 0
    for data in logos:
        try:
            img = XLImage(io.BytesIO(data))
        except Exception:
            continue
        k = 54 / max(img.height, img.width, 1)          # environ 14 mm de côté
        img.width, img.height = img.width * k, img.height * k
        _ancre(ws, img, ligne_logos, 18 + x)
        x += img.width + 10


def _ancre(ws, img, ligne: int, decalage_px: float) -> None:
    """Ancre l'image en haut de la ligne donnée, décalée de decalage_px depuis le bord gauche de la colonne A."""
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU
    marqueur = AnchorMarker(col=0, colOff=pixels_to_EMU(int(decalage_px)), row=ligne - 1, rowOff=0)
    img.anchor = OneCellAnchor(_from=marqueur, ext=XDRPositiveSize2D(pixels_to_EMU(int(img.width)),
                                                                     pixels_to_EMU(int(img.height))))
    ws.add_image(img)


def render(doc: E.Document, out: str | Path, sheet_title: str = "Rapport") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Rapport")[:31].replace("/", "-")
    widths: dict[int, float] = {}
    r = 1
    maxcol = 1
    droite_rows, logos, ligne_logos = [], [], 1
    if doc.entete.get("officiel", True):
        gauche, droite, logos = E.entete_officiel(doc)
        h0 = r
        for line in gauche:
            ws.cell(r, 1, line).font = Font(bold=True, size=9)
            r += 1
        droite_rows = [(h0 + i, line) for i, line in enumerate(droite)]   # placées une fois la largeur connue
        ligne_logos = r
        r = max(r + (4 if logos else 0), h0 + len(droite)) + 1
    print_title_rows = None
    for b in doc.blocks:
        if isinstance(b, tuple) and b[0] == "TITRE":
            if b[1]:
                ws.cell(r, 1, b[1]).font = Font(bold=True, size=13, color="1F4E79")
                r += 1
            if b[2]:
                ws.cell(r, 1, b[2]).font = Font(italic=True)
                r += 1
            r += 1
        elif isinstance(b, E.Table):
            if b.title:
                ws.cell(r, 1, b.title).font = Font(bold=True, color="1F4E79")
                r += 1
            ncols = len(b.types)
            maxcol = max(maxcol, ncols)
            # en-têtes avec fusions
            occupied = set()
            h0 = r
            for hr, row in enumerate(b.header):
                c = 1
                for cell in row:
                    while (r + hr, c) in occupied:
                        c += 1
                    x = ws.cell(r + hr, c, cell.text)
                    x.font = Font(bold=True, color="FFFFFF")
                    x.fill = HEAD_FILL
                    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    for rr in range(cell.rowspan):
                        for cc in range(cell.colspan):
                            occupied.add((r + hr + rr, c + cc))
                            ws.cell(r + hr + rr, c + cc).border = BORDER
                            ws.cell(r + hr + rr, c + cc).fill = HEAD_FILL
                    if cell.rowspan > 1 or cell.colspan > 1:
                        ws.merge_cells(start_row=r + hr, start_column=c, end_row=r + hr + cell.rowspan - 1,
                                       end_column=c + cell.colspan - 1)
                    widths[c] = max(widths.get(c, 0), min(len(cell.text) + 2, 30) / max(cell.colspan, 1))
                    c += cell.colspan
            r += len(b.header)
            if print_title_rows is None:
                print_title_rows = f"{h0}:{r - 1}"
            for row in b.rows:
                if row.kind == "group":
                    x = ws.cell(r, 1, row.cells[0])
                    x.font = Font(bold=True)
                    for c in range(1, ncols + 1):
                        ws.cell(r, c).fill = GROUP_FILL
                        ws.cell(r, c).border = BORDER
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
                    r += 1
                    continue
                rtype = row.style.get("type")
                for i, val in enumerate(row.cells):
                    typ = b.types[i]
                    if rtype and i > 0 and not isinstance(val, str):
                        typ = rtype
                    x = ws.cell(r, i + 1, _excel_value(val, typ))
                    x.border = BORDER
                    if typ in NUMFMT and not isinstance(val, str):
                        x.number_format = NUMFMT[typ]
                    nom = i < len(b.noms) and b.noms[i]          # noms et prénoms : en gras, à gauche
                    x.alignment = Alignment(horizontal="left" if nom else H_ALIGN.get(b.aligns[i], "left"),
                                            vertical="center", wrap_text=True)
                    if nom or row.kind in ("subtotal", "total") or (b.first_col_header and i == 0):
                        x.font = Font(bold=True)
                    if row.kind == "subtotal":
                        x.fill = SUB_FILL
                    elif row.kind == "total":
                        x.fill = TOT_FILL
                    elif row.style.get("statut") == "ABANDON":
                        x.fill = ROSE_FILL
                    widths[i + 1] = max(widths.get(i + 1, 0), min(len(E.fmt(val, typ)) + 2, 45))
                r += 1
            r += 1
        elif isinstance(b, E.Paragraph) and b.text:
            x = ws.cell(r, 1, b.text)
            x.font = Font(bold=b.style == "titre2", italic=b.style == "note")
            r += 1
        elif isinstance(b, E.KeyValues):
            if b.title:
                ws.cell(r, 1, b.title).font = Font(bold=True, color="1F4E79")
                r += 1
            for k, val in b.items:
                ws.cell(r, 1, k)
                ws.cell(r, 2, val).font = Font(bold=True)
                r += 1
            r += 1
        elif isinstance(b, E.Chart):
            ws.cell(r, 1, b.title).font = Font(bold=True, color="1F4E79")
            r += 1
            ws.cell(r, 1, "Classe").font = Font(bold=True)
            for j, (n, _) in enumerate(b.series):
                ws.cell(r, 2 + j, n).font = Font(bold=True)
            r += 1
            for i, cat in enumerate(b.categories):
                ws.cell(r, 1, cat)
                for j, (_, vals) in enumerate(b.series):
                    x = ws.cell(r, 2 + j, vals[i])
                    if b.percent:
                        x.number_format = "0.0%"
                r += 1
            r += 1
        elif isinstance(b, E.TextZone):
            ws.cell(r, 1, b.title).font = Font(bold=True)
            r += 1
            ws.cell(r, 1, b.text or "")
            r += 2
        elif isinstance(b, E.Bulletin):
            d = b.data
            ws.cell(r, 1, d.get("nom")).font = Font(bold=True, size=12)
            ws.cell(r, 3, f"{d.get('matricule_ecole')} – {d.get('classe')}")
            r += 1
            if d["type"] == "annuel":
                lignes = [(e["evaluation"], e["moyenne"], e["appreciation"]) for e in d["evaluations"]]
                lignes += [("MGA", d.get("mga"), d.get("observation") or ""), ("Rang", d.get("rang"), ""),
                           ("Décision", d.get("decision") or "", "")]
            else:
                lignes = [(n["matiere"], n["note"], n["appreciation"]) for n in d["notes"]]
                lignes += [("Total", d.get("total"), ""), ("Moyenne", d.get("moyenne"), d.get("appreciation") or ""),
                           ("Rang", d.get("rang"), f"sur {d.get('effectif')}")]
            for lib, val, app in lignes:
                ws.cell(r, 1, lib).border = BORDER
                x = ws.cell(r, 2, val)
                x.border = BORDER
                if isinstance(val, float):
                    x.number_format = "0.00"
                ws.cell(r, 3, app).border = BORDER
                r += 1
            widths[1] = max(widths.get(1, 0), 30)
            widths[3] = max(widths.get(3, 0), 22)
            r += 1
        elif isinstance(b, E.Lettre):
            for ligne in b.entete:
                ws.cell(r, 1, ligne).font = Font(bold=True)
                r += 1
            for g, d in b.references:
                ws.cell(r, 1, g)
                ws.cell(r, 4, d)
                r += 1
            r += 1
            ws.cell(r, 1, b.titre).font = Font(bold=True, size=13)
            r += 2
            for x in b.destinataire:
                ws.cell(r, 1, x)
                r += 1
            r += 1
            for style, segs in b.corps:
                ws.cell(r, 1, "".join(s for s, _ in segs)).font = Font(bold=style == "intitule")
                r += 1
            r += 1
            ws.cell(r, 1, b.titre_signatures).font = Font(bold=True)
            r += 1
            for i, (tete, lignes) in enumerate(b.signatures):
                ws.cell(r, 1 + 2 * i, tete).font = Font(bold=True)
                ws.cell(r + 1, 1 + 2 * i, " / ".join(lignes))
            r += 3
            widths[1] = max(widths.get(1, 0), 40)
        elif isinstance(b, E.EnteteOfficiel):
            gauche, _, _ = E.entete_officiel(doc)
            for line in gauche:
                ws.cell(r, 1, line).font = Font(bold=True, size=9)
                r += 1
            r += 1
        elif isinstance(b, E.Signature):
            r += 1
            ws.cell(r, 1, b.lieu_date)
            ws.cell(r, max(3, maxcol - 2), b.fonction).font = Font(bold=True)
            r += 4
            ws.cell(r, max(3, maxcol - 2), b.nom).font = Font(bold=True)
            r += 2
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = max(6, min(w * 1.1, 50))
    ws.page_setup.orientation = "landscape" if doc.orientation == "paysage" else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    # rapport « sur une seule page » : la feuille entière tient sur une page à l'impression
    ws.page_setup.fitToHeight = 1 if getattr(doc, "une_page", False) and \
        not any(isinstance(b, E.PageBreak) for b in doc.blocks) else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.8 / 2.54
    if print_title_rows:
        ws.print_title_rows = print_title_rows
    if droite_rows or logos:
        _entete_droite(ws, droite_rows, logos, maxcol, ligne_logos)
    ws.oddFooter.center.text = "Page &P / &N"
    wb.save(out)

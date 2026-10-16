"""Import / export des élèves au format de l'export national (fichier « ELEVES », 22 colonnes).

Colonnes, dans cet ordre (COLONNES) :
    codedren;codeiep;iep;codeecole;ecole;matricule;nom;prenoms;sexe;jour;mois;annee;nationalite;
    localite;codesp;sp;pere;mere;nacte;lieuacte;residence;niveau

Règles reprises du prototype AKWA SCHOOL (buildImport / expStu) :
    - sexe MASCULIN / FEMININ (M, F, GARCON, FILLE… acceptés) ;
    - date de naissance éclatée en jour / mois / annee (ou une colonne « naissance ») ;
    - « AUCUN », « NEANT » ou « 0 » = vide (localité, sous-préfecture, lieu de l'acte, parents…) ;
    - nacte = 0 : pas d'extrait de naissance ; sinon numéro de l'acte (extrait = OUI) ;
    - pere / mere : nom complet en une seule colonne ;
    - niveau (CP1 … CM2) -> la classe de l'année dont le code commence par ce niveau (s'il y en a
      plusieurs, la moins chargée) ; aucune classe n'est créée : la ligne est signalée en erreur ;
    - doublons écartés : matricule déjà inscrit dans l'année, matricule répété dans le fichier,
      ou même nom et même date de naissance déjà inscrits.
Un élève déjà connu de la base (même matricule) mais pas encore inscrit dans l'année est réinscrit.

Fichiers acceptés : .xls (Excel 97-2003, lu sans bibliothèque), .xlsx et .csv (séparateur ; ou ,).

Fonctions pour l'interface :
    lire_fichier(chemin)                    -> list[list[str]]
    analyser(db, annee, chemin)             -> dict (aperçu ligne par ligne, rien n'est écrit)
    importer(db, session, annee, chemin)    -> dict (même bilan, après enregistrement)
    exporter(db, session, annee, chemin)    -> Path (.csv ou .xlsx)
    modele(chemin, db=None, exemple=False)  -> Path (.csv ou .xlsx : en-têtes seuls)
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

from epp.core.db import Database
from epp.core.security import Session, log
from epp.domain import calculs, referentiels, repo
from epp.domain.fiche_scolaire import niveau as niveau_de, separer_nom
from epp.domain.referentiels import norm
from epp.domain.repo import Annee, ValidationError

COLONNES = ["codedren", "codeiep", "iep", "codeecole", "ecole", "matricule", "nom", "prenoms", "sexe", "jour",
            "mois", "annee", "nationalite", "localite", "codesp", "sp", "pere", "mere", "nacte", "lieuacte",
            "residence", "niveau"]

# Synonymes acceptés dans la ligne d'en-tête (comparés sans accents, casse ni ponctuation)
_ALIAS = {
    "nom": ("nom",), "prenoms": ("prenoms", "prenom"), "niveau": ("niveau", "classe"),
    "sexe": ("sexe", "genre"), "naissance": ("naissance", "datedenaissance", "datenaissance"),
    "jour": ("jour",), "mois": ("mois",), "annee": ("annee",), "matricule": ("matricule", "matriculedesps"),
    "nacte": ("nacte", "actenum", "numeroacte", "ndelacte", "numeroextrait"),
    "lieuacte": ("lieuacte", "centreetatcivil", "centre"), "nationalite": ("nationalite",),
    "localite": ("localite", "lieu", "lieudenaissance"), "codesp": ("codesp",),
    "sp": ("sp", "sousprefecture"), "pere": ("pere",), "mere": ("mere",), "residence": ("residence", "domicile"),
    "codeecole": ("codeecole",), "ecole": ("ecole",), "codedren": ("codedren",), "codeiep": ("codeiep",),
}
OBLIGATOIRES = ("nom", "prenoms", "niveau", "sexe")
MASCULIN = {"m", "masculin", "garcon", "g", "h", "homme"}
FEMININ = {"f", "feminin", "fille", "femme"}


# =================================================================== lecture
def _texte(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(round(v, 6))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _lire_csv(data: bytes) -> list[list[str]]:
    try:
        t = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        t = data.decode("cp1252", "replace")
    premiere = t.splitlines()[0] if t else ""
    sep = ";" if premiere.count(";") >= premiere.count(",") else ","
    return [r for r in csv.reader(io.StringIO(t, newline=""), delimiter=sep)]


def _lire_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def lire_fichier(chemin: str | Path) -> list[list[str]]:
    """Lignes (première feuille) d'un fichier .xls, .xlsx ou .csv, chaque cellule en texte ; les lignes
    entièrement vides sont retirées. Le format est reconnu au contenu, pas à l'extension."""
    data = Path(chemin).read_bytes()
    if data[:4] == b"\xD0\xCF\x11\xE0":
        from epp.io.lecture_xls import lire_xls
        brut = lire_xls(data)
    elif data[:2] == b"PK":
        brut = _lire_xlsx(data)
    else:
        brut = _lire_csv(data)
    lignes = [[_texte(v) for v in r] for r in brut]
    lignes = [r for r in lignes if any(x for x in r)]
    for r in lignes:                          # cellules vides de fin de ligne retirées
        while r and r[-1] == "":
            r.pop()
    return lignes


# =================================================================== analyse
def vide0(v) -> str:
    """« AUCUN », « NEANT », « 0 » ou « - » -> chaîne vide."""
    s = str(v or "").strip()
    return "" if re.fullmatch(r"(?i)(aucun|aucune|neant|néant|0|-)", s) else s


def _colonnes(entete: list[str]) -> dict[str, int]:
    h = [norm(x) for x in entete]
    out = {}
    for cle, noms in _ALIAS.items():
        out[cle] = next((i for i, x in enumerate(h) if x in noms), -1)
    return out


def _sexe(v: str) -> str:
    n = norm(v)
    return "M" if n in MASCULIN else "F" if n in FEMININ else ""


def _date(jour: str, mois: str, an: str, brut: str) -> tuple[str | None, bool]:
    """(date ISO ou None, valeur illisible ?)."""
    if an:
        try:
            return date(int(float(an)), int(float(mois)), int(float(jour))).isoformat(), False
        except (TypeError, ValueError):
            return None, True
    if brut:
        d = calculs.to_date(brut)
        if d is None:
            m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", brut.strip())
            try:
                d = date(int(m[3]), int(m[2]), int(m[1])) if m else None
            except ValueError:
                d = None
        return (d.isoformat(), False) if d else (None, True)
    return None, False


def _effectifs(db: Database, annee: Annee) -> dict[str, int]:
    return {r["code"]: r["n"] for r in db.query(
        "SELECT c.code, COUNT(i.id) AS n FROM classe c LEFT JOIN inscription i ON i.classe_id=c.id "
        "WHERE c.annee_id=? GROUP BY c.id", (annee.id,))}


def classe_du_niveau(annee: Annee, niveau: str, effectifs: dict[str, int] | None = None) -> str | None:
    """Classe de l'année pour un niveau du fichier national (« CE1 » -> « CE1 » ou « CE1 A ») : code égal,
    sinon code commençant par le niveau ; s'il y en a plusieurs, la moins chargée. None si aucune."""
    n = re.sub(r"\s", "", str(niveau or "").upper())
    if not n:
        return None
    codes = annee.codes
    exact = [c for c in codes if re.sub(r"\s", "", c.upper()) == n]
    if exact:
        return exact[0]
    cand = [c for c in codes if re.sub(r"\s", "", c.upper()).startswith(n)]
    if not cand:
        return None
    eff = effectifs or {}
    return min(cand, key=lambda c: (eff.get(c, 0), codes.index(c)))


def analyser(db: Database, annee: Annee, chemin: str | Path) -> dict:
    """Lit le fichier et prépare l'import sans rien écrire.

    Renvoie {"fichier", "lignes", "a_importer", "a_reinscrire", "doublons", "erreurs", "avertissements": [str],
    "details": [{"ligne", "etat" (ok | reinscription | doublon | erreur), "motif", "matricule", "nom",
    "niveau", "classe", "donnees"}]} ; « ligne » est le numéro de ligne dans le fichier (en-tête = 1)."""
    try:
        lignes = lire_fichier(chemin)
    except (OSError, ValueError) as e:
        raise ValidationError(f"Fichier illisible : {e}") from None
    if len(lignes) < 2:
        raise ValidationError("Fichier vide ou sans ligne de données.")
    col = _colonnes(lignes[0])
    debut = 1
    if any(col[k] < 0 for k in OBLIGATOIRES):
        if len(lignes[0]) == len(COLONNES) and norm(lignes[0][6]) != "nom":
            col = {k: (COLONNES.index(k) if k in COLONNES else -1) for k in _ALIAS}   # fichier sans en-tête
            debut = 0
        else:
            manquantes = [k for k in OBLIGATOIRES if col[k] < 0]
            raise ValidationError("Colonnes obligatoires manquantes : " + ", ".join(manquantes)
                                  + ". Colonnes attendues : " + ";".join(COLONNES))

    def g(r: list[str], k: str) -> str:
        i = col.get(k, -1)
        return r[i].strip() if 0 <= i < len(r) else ""

    e = repo.etablissement(db)
    effectifs = _effectifs(db, annee)
    existants = {norm(r["matricule_desps"]): r["id"] for r in
                 db.query("SELECT id, matricule_desps FROM eleve WHERE IFNULL(matricule_desps,'')<>''")}
    inscrits = {r["eleve_id"] for r in db.query("SELECT eleve_id FROM inscription WHERE annee_id=?", (annee.id,))}
    nom_naiss = {(norm(r["nom"]), r["date_naissance"] or "") for r in
                 db.query("SELECT e.nom, e.date_naissance FROM inscription i JOIN eleve e ON e.id=i.eleve_id "
                          "WHERE i.annee_id=?", (annee.id,))}
    vus_mat: set[str] = set()
    vus_nom: dict[tuple[str, str], int] = {}
    details, ecoles = [], set()
    for k, r in enumerate(lignes[debut:], start=debut + 1):
        nom_f, prenoms = g(r, "nom").upper(), g(r, "prenoms").upper()
        nom = re.sub(r"\s+", " ", f"{nom_f} {prenoms}").strip()
        mat = re.sub(r"\s", "", g(r, "matricule")).upper()
        niv = re.sub(r"\s", "", g(r, "niveau").upper())
        classe = classe_du_niveau(annee, niv, effectifs)
        sexe = _sexe(g(r, "sexe"))
        dn, date_ko = _date(g(r, "jour"), g(r, "mois"), g(r, "annee"), g(r, "naissance"))
        nacte = vide0(g(r, "nacte"))
        sp = vide0(g(r, "sp"))
        if not sp and vide0(g(r, "codesp")):
            sp = referentiels.libelle_sp(g(r, "codesp"))
            sp = "" if sp == "INCONNU" else sp
        residence = vide0(g(r, "residence"))
        donnees = {
            "matricule_desps": mat or None, "nom": nom, "sexe": sexe, "date_naissance": dn,
            "nationalite": referentiels.nationalite_officielle(g(r, "nationalite")) or referentiels.NATIONALITE_DEFAUT,
            "localite": vide0(g(r, "localite")).upper() or None,
            "sous_prefecture": referentiels.sous_prefecture_officielle(sp) or None,
            "extrait": "OUI" if nacte else "NON", "acte_num": nacte or None,
            "centre_etat_civil": (vide0(g(r, "lieuacte")).upper() or None) if nacte else None,
            "pere_nom": vide0(g(r, "pere")).upper() or None, "mere_nom": vide0(g(r, "mere")).upper() or None,
            "pere_residence": residence or None, "mere_residence": residence or None,
        }
        d = {"ligne": k, "etat": "ok", "motif": "", "matricule": mat, "nom": nom, "niveau": niv,
             "classe": classe or "", "donnees": donnees, "eleve_id": None}
        if g(r, "codeecole") or g(r, "ecole"):
            ecoles.add((g(r, "codeecole"), g(r, "ecole")))
        km, kn = norm(mat), (norm(nom), dn or "")
        if not nom_f or not prenoms:
            d.update(etat="erreur", motif="nom ou prénoms manquant")
        elif not sexe:
            d.update(etat="erreur", motif=f"sexe illisible « {g(r, 'sexe')} »")
        elif not niv:
            d.update(etat="erreur", motif="niveau manquant")
        elif classe is None:
            d.update(etat="erreur", motif=f"aucune classe {niv} dans l'année {annee.libelle} : créez-la puis "
                                          f"relancez l'import")
        elif date_ko:
            d.update(etat="erreur", motif="date de naissance invalide")
        elif mat and not calculs.matricule_desps_valide(mat):
            d.update(etat="erreur", motif=f"matricule invalide « {mat} » : {calculs.MATRICULE_DESPS_AIDE}")
        elif km and km in vus_mat:
            d.update(etat="doublon", motif="matricule en double dans le fichier")
        elif km and km in existants and existants[km] in inscrits:
            d.update(etat="doublon", motif="matricule déjà inscrit cette année")
        elif kn in nom_naiss:
            d.update(etat="doublon", motif="même nom et même date de naissance déjà inscrits cette année")
        elif kn in vus_nom:
            d.update(etat="doublon", motif=f"même nom et même date de naissance qu'à la ligne {vus_nom[kn]} du "
                                          f"fichier : vérifiez (matricule différent ?)")
        else:
            try:
                repo.validate_eleve(donnees, annee)
            except ValidationError as ex:
                d.update(etat="erreur", motif=str(ex))
        if d["etat"] == "ok":
            if km and km in existants:
                d.update(etat="reinscription", motif="élève déjà connu : réinscrit pour l'année",
                         eleve_id=existants[km])
            effectifs[classe] = effectifs.get(classe, 0) + 1
            vus_nom[kn] = k
        if km:
            vus_mat.add(km)
        details.append(d)
    avert = []
    code_ecole = (e.get("ecole_code") or "").strip().upper()
    for code, nom_ecole in sorted(ecoles):
        if (code_ecole and code and code.strip().upper() != code_ecole) or \
                (not code_ecole and nom_ecole and norm(nom_ecole) != norm(e.get("ecole_nom"))):
            avert.append(f"Le fichier concerne l'école {nom_ecole or '?'} ({code or 'sans code'}) et non "
                         f"{e.get('ecole_nom', '')}{f' ({code_ecole})' if code_ecole else ''} : vérifiez le fichier.")
    return {
        "fichier": Path(chemin).name, "lignes": len(details),
        "a_importer": sum(d["etat"] == "ok" for d in details),
        "a_reinscrire": sum(d["etat"] == "reinscription" for d in details),
        "doublons": sum(d["etat"] == "doublon" for d in details),
        "erreurs": sum(d["etat"] == "erreur" for d in details),
        "avertissements": avert, "details": details,
    }


# =================================================================== import
def importer(db: Database, s: Session, annee: Annee, chemin: str | Path, sauvegarde: bool = True) -> dict:
    """Inscrit dans l'année les élèves du fichier national ; les doublons et les lignes en erreur sont écartés.

    Renvoie le bilan d'analyser() complété de {"importes", "reinscrits", "ignores", "sauvegarde"} ;
    chaque ligne de "details" garde son état final (ok, reinscription, doublon, erreur) et son motif."""
    s.require("import.excel")
    s.require("eleves.modifier")
    bilan = analyser(db, annee, chemin)
    valides = [d for d in bilan["details"] if d["etat"] in ("ok", "reinscription")]
    bilan.update(importes=0, reinscrits=0, sauvegarde=None)
    if valides and sauvegarde:
        from epp.core import backup
        bilan["sauvegarde"] = str(backup.sauvegarde_auto(db, "avant import ELEVES", login=s.login))
    with db.transaction():
        for d in valides:
            db.conn.execute("SAVEPOINT ligne_eleve")
            try:
                if d["etat"] == "reinscription":
                    ancien = repo.get_eleve(db, d["eleve_id"])
                    fiche = {k: d["donnees"].get(k) if d["donnees"].get(k) not in (None, "") else ancien.get(k)
                             for k in repo.ELEVE_CHAMPS}
                    eid = repo.save_eleve(db, s, d["eleve_id"], fiche, annee)
                    repo.inscrire(db, s, eid, annee, d["classe"], type_="REINSCRIPTION",
                                  date_inscription=annee.date_debut)
                    bilan["reinscrits"] += 1
                else:
                    eid = repo.save_eleve(db, s, None, d["donnees"], annee)
                    repo.inscrire(db, s, eid, annee, d["classe"], date_inscription=annee.date_debut)
                    bilan["importes"] += 1
                db.conn.execute("RELEASE ligne_eleve")
            except ValidationError as ex:
                db.conn.execute("ROLLBACK TO ligne_eleve")
                db.conn.execute("RELEASE ligne_eleve")
                d.update(etat="erreur", motif=str(ex))
    bilan["erreurs"] = sum(d["etat"] == "erreur" for d in bilan["details"])
    bilan["ignores"] = bilan["doublons"] + bilan["erreurs"]
    log(db, s.login, "IMPORT_ELEVES_NATIONAL", "eleve",
        f"{bilan['fichier']} : {bilan['importes']} importé(s), {bilan['reinscrits']} réinscrit(s), "
        f"{bilan['doublons']} doublon(s), {bilan['erreurs']} erreur(s)")
    return bilan


def resume(bilan: dict) -> str:
    """Bilan lisible d'une analyse ou d'un import (une ligne par élément, erreurs détaillées)."""
    out = [f"Fichier : {bilan.get('fichier', '')} ({bilan.get('lignes', 0)} ligne(s))"]
    if "importes" in bilan:
        out.append(f"Élèves importés : {bilan['importes']} ; réinscrits : {bilan['reinscrits']}")
    else:
        out.append(f"À importer : {bilan['a_importer']} ; à réinscrire : {bilan['a_reinscrire']}")
    out.append(f"Doublons écartés : {bilan['doublons']} ; lignes en erreur : {bilan['erreurs']}")
    out += [f"Attention : {a}" for a in bilan.get("avertissements", [])]
    for d in bilan.get("details", []):
        if d["etat"] in ("doublon", "erreur"):
            out.append(f" - ligne {d['ligne']} ({d['nom'] or '?'}) : {d['etat']} – {d['motif']}")
    return "\n".join(out)


# =================================================================== export
def _iepp_court(iepp: str) -> str:
    return re.sub(r"(?i)^\s*IEPP\s+", "", referentiels.nom_sous_prefecture(iepp or "")).strip()


def lignes_export(db: Database, annee: Annee) -> list[list[str]]:
    """Élèves présents de l'année, par classe puis par nom, au format des 22 colonnes (en-tête compris)."""
    e = repo.etablissement(db)
    tete = [e.get("code_dren", ""), e.get("code_iepp", ""), _iepp_court(e.get("iepp", "")),
            e.get("ecole_code", ""), e.get("ecole_nom", "")]
    rang = {c: i for i, c in enumerate(annee.codes)}
    rows = sorted(repo.inscriptions(db, annee, statut="PRESENT"),
                  key=lambda r: (rang.get(r["classe"], 999), r["nom"] or ""))
    out = [list(COLONNES)]
    for r in rows:
        nom, prenoms = separer_nom(r["nom"])
        dn = calculs.to_date(r.get("date_naissance"))
        extrait = (r.get("extrait") or "").upper() == "OUI" and bool(r.get("acte_num"))
        sp = referentiels.nom_sous_prefecture(r.get("sous_prefecture") or "").upper()
        out.append([
            *tete, r.get("matricule_desps") or "", nom.upper(), prenoms.upper(),
            "MASCULIN" if r["sexe"] == "M" else "FEMININ",
            str(dn.day) if dn else "", str(dn.month) if dn else "", str(dn.year) if dn else "",
            referentiels.nationalite_officielle(r.get("nationalite")) or referentiels.NATIONALITE_DEFAUT,
            (r.get("localite") or "").upper(), referentiels.code_sp(sp), sp,
            (r.get("pere_nom") or "").upper(), (r.get("mere_nom") or "").upper(),
            str(r["acte_num"]) if extrait else "0",
            (r.get("centre_etat_civil") or "").upper() if extrait else "",
            r.get("pere_residence") or r.get("mere_residence") or "", niveau_de(r["classe"]),
        ])
    return out


def _ecrire(lignes: list[list[str]], chemin: Path) -> Path:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    if chemin.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "ELEVES"
        for r in lignes:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0E6646")
        for row in ws.iter_rows():
            for c in row:
                c.number_format = "@"          # tout en texte : matricules et codes gardent leurs zéros
        for i, nom in enumerate(COLONNES, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(9, min(28, len(nom) + 4))
        ws.freeze_panes = "A2"
        wb.save(chemin)
    else:
        with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n").writerows(lignes)
    return chemin


def exporter(db: Database, s: Session, annee: Annee, chemin: str | Path) -> Path:
    """Exporte les élèves présents de l'année au format national (.csv séparé par « ; » ou .xlsx,
    selon l'extension du chemin)."""
    s.require("export.excel")
    s.require("eleves.lire")
    out = _ecrire(lignes_export(db, annee), Path(chemin))
    log(db, s.login, "EXPORT_ELEVES_NATIONAL", "eleve", out.name)
    return out


EXEMPLE = ["", "", "", "", "", "00000000A", "KOUADIO", "MARIE CLAIRE", "FEMININ", "14", "5", "2018", "IVOIRIENNE",
           "OUANINOU", "SP000725", "OUANINOU", "KOUADIO KOFFI", "YAO AYA", "245", "OUANINOU", "", "CE1"]


def modele(chemin: str | Path, db: Database | None = None, exemple: bool = False) -> Path:
    """Fichier modèle vide (.csv ou .xlsx) aux 22 colonnes ; avec exemple=True, une ligne d'exemple fictive
    (codes DREN / IEPP / école de l'établissement si db est fourni)."""
    lignes = [list(COLONNES)]
    if exemple:
        ligne = list(EXEMPLE)
        if db is not None:
            e = repo.etablissement(db)
            ligne[:5] = [e.get("code_dren", ""), e.get("code_iepp", ""), _iepp_court(e.get("iepp", "")),
                         e.get("ecole_code", ""), e.get("ecole_nom", "")]
        lignes.append(ligne)
    return _ecrire(lignes, Path(chemin))

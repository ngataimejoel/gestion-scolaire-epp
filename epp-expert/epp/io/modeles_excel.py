"""Import et export de données par modèles Excel.

Pour chaque type de données (élèves, personnel, classes, notes d'une évaluation, absences), l'application
fournit un classeur modèle : en-têtes, listes déroulantes, contrôles de saisie, feuille d'aide. Le même
classeur sert à l'export (rempli avec les données de l'année) et à l'import :

    1. analyser() lit le fichier sans rien écrire : action prévue et erreurs pour chaque ligne,
       doublons détectés (dans le fichier et avec la base) ;
    2. appliquer() enregistre les lignes valides, en une seule transaction, après une sauvegarde
       automatique ; les lignes en erreur sont ignorées et listées (ecrire_rapport()).
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from epp import APP_NAME
from epp.core.db import Database
from epp.core.security import Session, log
from epp.domain import calculs, repo
from epp.domain.repo import Annee, ValidationError

VERSION = 1
META = "_epp"
LISTES = "Listes"
ENTETE = PatternFill("solid", fgColor="0E6646")
ENTETE_OBL = PatternFill("solid", fgColor="0A4D35")
INFO = PatternFill("solid", fgColor="EEF3F0")
ERREUR = PatternFill("solid", fgColor="FDE2E1")
FIN = Side(style="thin", color="C9D3CE")


@dataclass
class Col:
    cle: str
    titre: str
    type: str = "texte"            # texte | date | nombre | liste | info (lu nulle part)
    obligatoire: bool = False
    liste: str | None = None       # nom d'une liste (paramètres) ou liste calculée
    largeur: int = 16
    aide: str = ""
    mini: float | None = None
    maxi: float | None = None


def _l(nom):
    return ("param", nom)


LISTES_CALCULEES: dict[str, Callable[[Database, Annee], list[str]]] = {
    "classes": lambda db, a: list(a.codes),
    "classes_ou_tiret": lambda db, a: ["-"] + list(a.codes),
    "grilles": lambda db, a: [g.code for g in a.grilles.values()],
    "calendriers": lambda db, a: sorted(a.evaluations) or ["STD", "CM2"],
    "regles": lambda db, a: ["STANDARD", "MOYENNE4"],
}

ELEVE_COLS = [
    Col("matricule_ecole", "Matricule école", aide="Vide pour un nouvel élève (attribué automatiquement). "
                                                   "Rempli : l'élève existant est mis à jour."),
    Col("matricule_desps", "Matricule DESPS"),
    Col("nom", "Nom et prénoms", obligatoire=True, largeur=30),
    Col("sexe", "Sexe", "liste", True, "SEXE", 7),
    Col("date_naissance", "Date de naissance", "date", largeur=13),
    Col("classe", "Classe", "liste", True, "classes", 9),
    Col("redoublant", "Redoublant", "liste", liste="OUI_NON", largeur=11),
    Col("statut", "Statut", "liste", liste="STATUT", largeur=12),
    Col("nationalite", "Nationalité", "liste", liste="NATIONALITE", largeur=14),
    Col("localite", "Localité"), Col("sous_prefecture", "Sous-préfecture"),
    Col("extrait", "Extrait de naissance", "liste", liste="OUI_NON", largeur=11),
    Col("acte_num", "N° acte"), Col("acte_date", "Date de l'acte", "date", largeur=13),
    Col("centre_etat_civil", "Centre d'état civil"),
    Col("orphelin", "Orphelin", "liste", liste="OUI_NON", largeur=10),
    Col("orphelin_de", "Orphelin de", "liste", liste="ORPHELIN_DE", largeur=13),
    Col("pere_nom", "Père", largeur=24), Col("pere_profession", "Profession du père"),
    Col("pere_residence", "Résidence du père"), Col("pere_tel", "Téléphone du père", largeur=15),
    Col("mere_nom", "Mère", largeur=24), Col("mere_profession", "Profession de la mère"),
    Col("mere_residence", "Résidence de la mère"), Col("mere_tel", "Téléphone de la mère", largeur=15),
    Col("tuteur_nom", "Tuteur", largeur=24), Col("tuteur_profession", "Profession du tuteur"),
    Col("tuteur_residence", "Résidence du tuteur"), Col("tuteur_tel", "Téléphone du tuteur", largeur=15),
    Col("observations", "Observations", largeur=30),
]
PERSONNEL_COLS = [
    Col("matricule", "Matricule", obligatoire=True, largeur=13,
        aide="Clé de l'agent : un matricule existant met la fiche à jour."),
    Col("nom", "Nom et prénoms", obligatoire=True, largeur=30),
    Col("sexe", "Sexe", "liste", True, "SEXE", 7),
    Col("date_naissance", "Date de naissance", "date", largeur=13),
    Col("fonction", "Fonction", "liste", liste="FONCTION", largeur=24),
    Col("grade", "Grade / emploi"), Col("diplome", "Diplôme"),
    Col("date_prise_service", "Prise de service", "date", largeur=13),
    Col("date_arrivee", "Arrivée dans l'école", "date", largeur=13),
    Col("telephone", "Téléphone", largeur=15),
    Col("situation_matrimoniale", "Situation matrimoniale", "liste", liste="SITUATION_MATRIMONIALE", largeur=16),
    Col("classe_tenue", "Classe tenue", "liste", liste="classes_ou_tiret", largeur=10),
    Col("observations", "Observations", largeur=30),
]
CLASSE_COLS = [
    Col("code", "Code", obligatoire=True, largeur=9, aide="Code existant : la classe est mise à jour."),
    Col("libelle", "Libellé", largeur=18), Col("ordre", "Ordre", "nombre", largeur=7, mini=0, maxi=99),
    Col("niveau_suivant", "Classe suivante", largeur=12, aide="Vide pour la dernière classe (sortie)."),
    Col("seuil", "Seuil d'admission", "nombre", True, largeur=10, mini=0, maxi=100),
    Col("bareme", "Barème", "nombre", True, largeur=9, mini=1, maxi=100),
    Col("grille", "Grille de notes", "liste", True, "grilles", 14),
    Col("calendrier", "Calendrier", "liste", liste="calendriers", largeur=11),
    Col("regle_mga", "Règle de MGA", "liste", liste="regles", largeur=13),
]


def _abs_cols(kind):
    cible = ([Col("matricule_ecole", "Matricule école", obligatoire=True, largeur=14)] if kind == "eleve" else
             [Col("matricule", "Matricule", obligatoire=True, largeur=13)])
    return [Col("date", "Date", "date", True, largeur=12), *cible,
            Col("nom", "Nom (pour information)", "info", largeur=28),
            Col("nature", "Nature", "liste", True, "NATURE_ABSENCE", 11),
            Col("jours", "Jours d'absence", "nombre", largeur=10, mini=0, maxi=31, aide="0,5 = demi-journée"),
            Col("retard_min", "Retard (minutes)", "nombre", largeur=10, mini=0, maxi=600),
            Col("motif", "Motif", "liste", liste="MOTIF_ABSENCE", largeur=22),
            Col("justifie", "Justifié", "liste", liste="OUI_NON", largeur=9),
            Col("observations", "Observations", largeur=30)]


TYPES = {
    "eleves": ("Élèves", "Élèves et inscriptions de l'année"),
    "personnel": ("Personnel", "Personnel de l'école et classe tenue"),
    "classes": ("Classes", "Classes de l'année (seuils, barèmes, grilles)"),
    "notes": ("Notes", "Notes d'une évaluation (une feuille par classe)"),
    "absences_eleves": ("Absences élèves", "Retards et absences des élèves"),
    "absences_personnel": ("Absences personnel", "Retards et absences du personnel"),
}


def colonnes(kind: str, annee: Annee | None = None, classe: str | None = None, num: int | None = None) -> list[Col]:
    if kind == "eleves":
        return ELEVE_COLS
    if kind == "personnel":
        return PERSONNEL_COLS
    if kind == "classes":
        return CLASSE_COLS
    if kind == "absences_eleves":
        return _abs_cols("eleve")
    if kind == "absences_personnel":
        return _abs_cols("personnel")
    if kind == "notes":
        cl = annee.classe(classe)
        g = annee.grilles[cl.grille_id]
        cols = [Col("matricule_ecole", "Matricule école", obligatoire=True, largeur=14),
                Col("nom", "Nom (pour information)", "info", largeur=28),
                Col("present", "Présent", "liste", liste="OUI_NON", largeur=9,
                    aide="NON : élève absent à l'évaluation (ses notes sont effacées).")]
        for m in g.matieres.get(num, []):
            mx = repo.max_note(g, m)
            cols.append(Col(f"m:{m['id']}", f"{m['libelle']} (/{mx:g})", "nombre", largeur=14, mini=0, maxi=mx,
                            aide=f"Note sur {mx:g}. Case vide : la note existante n'est pas modifiée."))
        return cols
    raise ValueError(kind)


# =================================================================== génération / export
def _valeurs_liste(db: Database, annee: Annee, nom: str) -> list[str]:
    if nom in LISTES_CALCULEES:
        return LISTES_CALCULEES[nom](db, annee)
    return repo.liste(db, nom) or []


def _feuille(wb, db, annee, titre, cols: list[Col], lignes: list[dict], listes_ws, listes_idx):
    ws = wb.create_sheet(titre[:31])
    bord = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
    for j, c in enumerate(cols, start=1):
        x = ws.cell(1, j, c.titre + (" *" if c.obligatoire else ""))
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = ENTETE_OBL if c.obligatoire else ENTETE
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.border = bord
        if c.aide:
            x.comment = Comment(c.aide, "EPP Expert")
        ws.column_dimensions[get_column_letter(j)].width = c.largeur
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    for i, ligne in enumerate(lignes, start=2):
        for j, c in enumerate(cols, start=1):
            v = ligne.get(c.cle)
            if c.type == "date":
                d = calculs.to_date(v)
                v = d if d else v
            x = ws.cell(i, j, v)
            if c.type == "date":
                x.number_format = "DD/MM/YYYY"
            if c.type == "info":
                x.fill = INFO
    derniere = max(len(lignes) + 1, 500)
    for j, c in enumerate(cols, start=1):
        lettre = get_column_letter(j)
        plage = f"{lettre}2:{lettre}{derniere}"
        dv = None
        if c.type == "liste" and c.liste:
            if c.liste not in listes_idx:
                vals = _valeurs_liste(db, annee, c.liste)
                col = len(listes_idx) + 1
                listes_ws.cell(1, col, c.liste)
                for k, v in enumerate(vals, start=2):
                    listes_ws.cell(k, col, v)
                listes_idx[c.liste] = (get_column_letter(col), len(vals) + 1)
            lc, n = listes_idx[c.liste]
            dv = DataValidation(type="list", formula1=f"={LISTES}!${lc}$2:${lc}${max(n, 2)}", allow_blank=True)
            dv.error, dv.errorTitle = "Choisissez une valeur de la liste.", "Valeur non prévue"
        elif c.type == "date":
            dv = DataValidation(type="date", operator="between", formula1="DATE(1940,1,1)",
                                formula2="DATE(2100,12,31)", allow_blank=True)
            dv.error, dv.errorTitle = "Saisissez une date (jj/mm/aaaa).", "Date invalide"
            for r in range(2, derniere + 1):
                ws.cell(r, j).number_format = "DD/MM/YYYY"
        elif c.type == "nombre" and c.maxi is not None:
            dv = DataValidation(type="decimal", operator="between", formula1=str(c.mini or 0),
                                formula2=str(c.maxi), allow_blank=True)
            dv.error, dv.errorTitle = f"Nombre entre {c.mini or 0:g} et {c.maxi:g}.", "Valeur hors limites"
        if dv:
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(plage)
    return ws


def _lignes_export(db: Database, annee: Annee, kind: str, classe: str | None = None, num: int | None = None):
    if kind == "eleves":
        return [r | {"classe": r["classe"]} for r in repo.inscriptions(db, annee)]
    if kind == "personnel":
        return repo.personnel(db, annee)
    if kind == "classes":
        return [{"code": c.code, "libelle": c.libelle, "ordre": c.ordre, "niveau_suivant": c.niveau_suivant or "",
                 "seuil": c.seuil, "bareme": c.bareme, "grille": annee.grilles[c.grille_id].code if c.grille_id in
                 annee.grilles else "", "calendrier": c.calendrier, "regle_mga": c.regle_mga} for c in annee.classes]
    if kind == "absences_eleves":
        return repo.absences_eleves(db, annee)
    if kind == "absences_personnel":
        return repo.absences_personnel(db, annee)
    if kind == "notes":
        cl = annee.classe(classe)
        _, lignes = repo.notes_grille(db, annee, cl.grille_id, num, cl.code)
        return [{"matricule_ecole": r["matricule_ecole"], "nom": r["nom"], "present": "OUI" if r["present"] else "NON",
                 **{f"m:{mid}": v for mid, v in r["notes"].items()}} for r in lignes]
    raise ValueError(kind)


def creer_modele(db: Database, annee: Annee, kind: str, chemin: str | Path, avec_donnees: bool = False,
                 num: int | None = None, classes: list[str] | None = None) -> Path:
    """Classeur modèle vide (ou rempli avec les données de l'année si avec_donnees : export)."""
    if kind not in TYPES:
        raise ValueError(kind)
    if kind == "notes" and num is None:
        raise ValidationError("Choisissez l'évaluation.", "evaluation")
    wb = Workbook()
    wb.remove(wb.active)
    listes_ws = wb.create_sheet(LISTES)
    listes_idx: dict = {}
    if kind == "notes":
        for code in classes or annee.codes:
            cl = annee.classe(code)
            g = annee.grilles.get(cl.grille_id)
            if g is None or num not in g.matieres:
                continue
            _feuille(wb, db, annee, code, colonnes(kind, annee, code, num),
                     _lignes_export(db, annee, kind, code, num) if avec_donnees else
                     [{"matricule_ecole": r["matricule_ecole"], "nom": r["nom"]}
                      for r in repo.inscriptions(db, annee, code)], listes_ws, listes_idx)
    else:
        _feuille(wb, db, annee, TYPES[kind][0], colonnes(kind),
                 _lignes_export(db, annee, kind) if avec_donnees else [], listes_ws, listes_idx)
    aide = wb.create_sheet("Aide")
    textes = [f"Modèle « {TYPES[kind][1]} » – {APP_NAME} – année {annee.libelle}", "",
              "• Les colonnes marquées * sont obligatoires ; les listes déroulantes donnent les valeurs admises.",
              "• Les dates s'écrivent jj/mm/aaaa.",
              "• Ne renommez pas les feuilles et ne supprimez pas la feuille cachée _epp.",
              "• À l'import, l'application montre d'abord ce qui sera créé, modifié ou refusé (avec la raison), "
              "puis n'enregistre que les lignes valides, après une sauvegarde automatique."]
    if kind == "notes":
        textes.append("• Une case de note vide ne modifie pas la note déjà enregistrée.")
    for i, t in enumerate(textes, start=1):
        aide.cell(i, 1, t).font = Font(bold=i == 1, size=12 if i == 1 else 10)
    aide.column_dimensions["A"].width = 120
    meta = wb.create_sheet(META)
    for i, (k, v) in enumerate((("type", kind), ("annee", annee.libelle), ("version", VERSION),
                                ("evaluation", num or "")), start=1):
        meta.cell(i, 1, k)
        meta.cell(i, 2, v)
    meta.sheet_state = "hidden"
    listes_ws.sheet_state = "hidden"
    wb.active = 1                        # première feuille de données (la feuille Listes est cachée)
    chemin = Path(chemin)
    wb.save(chemin)
    return chemin


def exporter(db: Database, s: Session, annee: Annee, kind: str, chemin: str | Path, num: int | None = None) -> Path:
    s.require("export.excel")
    out = creer_modele(db, annee, kind, chemin, avec_donnees=True, num=num)
    log(db, s.login, "EXPORT_MODELE_EXCEL", kind, Path(chemin).name)
    return out


# =================================================================== analyse
@dataclass
class Ligne:
    feuille: str
    numero: int
    donnees: dict
    action: str = ""               # creer | modifier | inscrire | ignorer
    erreurs: list[str] = field(default_factory=list)
    avertissements: list[str] = field(default_factory=list)
    cible: dict = field(default_factory=dict)

    @property
    def valide(self) -> bool:
        return not self.erreurs and self.action != "ignorer"


@dataclass
class Analyse:
    kind: str
    annee: str
    evaluation: int | None
    lignes: list[Ligne] = field(default_factory=list)
    erreurs_fichier: list[str] = field(default_factory=list)

    @property
    def corrections(self) -> int:
        """Notes déjà saisies que l'import remplacerait (mot de passe administrateur exigé)."""
        return sum(l.cible.get("corrections", 0) for l in self.lignes if l.valide)

    def compte(self) -> dict:
        c = {"creer": 0, "modifier": 0, "inscrire": 0, "ignorer": 0, "erreurs": 0}
        for l in self.lignes:
            if l.erreurs:
                c["erreurs"] += 1
            else:
                c[l.action] = c.get(l.action, 0) + 1
        return c

    @property
    def importable(self) -> bool:
        return not self.erreurs_fichier and any(l.valide for l in self.lignes)


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _texte(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)                       # téléphone ou matricule saisi comme nombre
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).strip()
    return s or None


def _date(v) -> tuple[str | None, bool]:
    """(date ISO, valide)."""
    if v in (None, ""):
        return None, True
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat(), True
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", s)
    if m:
        try:
            return date(int(m[3]), int(m[2]), int(m[1])).isoformat(), True
        except ValueError:
            return None, False
    d = calculs.to_date(s)
    return (d.isoformat(), True) if d else (None, False)


def _lire_feuille(ws, cols: list[Col], an: Analyse) -> list[Ligne]:
    entetes = [_norm(str(c.value or "").replace("*", "")) for c in ws[1]]
    pos = {}
    for c in cols:
        k = _norm(c.titre)
        if k in entetes:
            pos[c.cle] = entetes.index(k)
        elif c.obligatoire:
            an.erreurs_fichier.append(f"Feuille {ws.title} : colonne obligatoire « {c.titre} » absente.")
    if an.erreurs_fichier:
        return []
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in row):
            continue
        brut = {cle: row[p] if p < len(row) else None for cle, p in pos.items()}
        l = Ligne(ws.title, i, {})
        for c in cols:
            if c.cle not in brut or c.type == "info":
                l.donnees[c.cle] = _texte(brut.get(c.cle)) if c.type == "info" else None
                continue
            v = brut[c.cle]
            if c.type == "date":
                iso, ok = _date(v)
                if not ok:
                    l.erreurs.append(f"{c.titre} : date invalide « {v} ».")
                l.donnees[c.cle] = iso
            elif c.type == "nombre":
                if v in (None, ""):
                    l.donnees[c.cle] = None
                else:
                    try:
                        x = float(str(v).replace(",", ".")) if not isinstance(v, (int, float)) else float(v)
                        if c.maxi is not None and not (c.mini or 0) <= x <= c.maxi:
                            l.erreurs.append(f"{c.titre} : {x:g} hors limites ({c.mini or 0:g} à {c.maxi:g}).")
                        l.donnees[c.cle] = x
                    except ValueError:
                        if c.cle.startswith("m:") and _norm(v) in ("abs", "absent"):
                            l.donnees[c.cle] = "ABS"
                        else:
                            l.erreurs.append(f"{c.titre} : nombre attendu, « {v} » saisi.")
                            l.donnees[c.cle] = None
            else:
                t = _texte(v)
                if c.type == "liste" and t and c.liste in ("SEXE", "OUI_NON", "STATUT", "NATURE_ABSENCE"):
                    t = t.upper()
                    t = {"PRÉSENT": "PRESENT", "TRANSFÉRÉ": "TRANSFERE", "TRANSFERE(E)": "TRANSFERE",
                         "G": "M", "GARCON": "M", "FILLE": "F"}.get(t, t)
                l.donnees[c.cle] = t
            if c.obligatoire and l.donnees.get(c.cle) in (None, ""):
                l.erreurs.append(f"{c.titre} : obligatoire.")
        out.append(l)
    return out


def analyser(db: Database, annee: Annee, fichier: str | Path, kind: str | None = None) -> Analyse:
    """Lit un classeur rempli et prévoit, ligne par ligne, ce que fera l'import. N'écrit rien."""
    try:
        wb = load_workbook(fichier, data_only=True)
    except Exception as e:
        a = Analyse(kind or "", annee.libelle, None)
        a.erreurs_fichier.append(f"Fichier Excel illisible : {e}")
        return a
    meta = {}
    if META in wb.sheetnames:
        meta = {str(r[0]): r[1] for r in wb[META].iter_rows(values_only=True) if r and r[0]}
    if kind and meta.get("type") and meta["type"] != kind:
        a = Analyse(kind, annee.libelle, None)
        a.erreurs_fichier.append(f"Ce fichier est un modèle « {TYPES.get(meta['type'], (meta['type'],))[0]} », "
                                 f"pas « {TYPES.get(kind, (kind,))[0]} ».")
        return a
    kind = kind or meta.get("type")
    if not kind:
        noms = {v[0]: k for k, v in TYPES.items()}
        kind = next((noms[n] for n in wb.sheetnames if n in noms), None)
    num = int(meta["evaluation"]) if str(meta.get("evaluation") or "").isdigit() else None
    an = Analyse(kind or "", annee.libelle, num)
    if kind not in TYPES:
        an.erreurs_fichier.append("Type de modèle non reconnu : utilisez un modèle fourni par l'application.")
        return an
    if meta.get("annee") and meta["annee"] != annee.libelle:
        an.erreurs_fichier.append(f"Ce fichier concerne l'année {meta['annee']}, l'année ouverte est {annee.libelle}.")
        return an
    if annee.cloturee:
        an.erreurs_fichier.append(f"L'année {annee.libelle} est clôturée : import impossible.")
        return an
    if kind == "notes":
        if num is None:
            an.erreurs_fichier.append("Évaluation non précisée dans le fichier.")
            return an
        for nom in wb.sheetnames:
            cl = annee.classe(nom)
            if nom in (META, LISTES, "Aide") or cl is None:
                continue
            g = annee.grilles.get(cl.grille_id)
            if g is None or num not in g.matieres:
                continue
            for l in _lire_feuille(wb[nom], colonnes("notes", annee, nom, num), an):
                l.cible["classe"] = nom
                an.lignes.append(l)
    else:
        nom = TYPES[kind][0]
        ws = wb[nom] if nom in wb.sheetnames else wb.worksheets[0]
        an.lignes = _lire_feuille(ws, colonnes(kind), an)
    if not an.erreurs_fichier:
        {"eleves": _an_eleves, "personnel": _an_personnel, "classes": _an_classes, "notes": _an_notes,
         "absences_eleves": _an_absences, "absences_personnel": _an_absences}[kind](db, annee, an)
    return an


def _an_eleves(db, annee, an):
    par_mat = {r["matricule_ecole"]: r for r in repo.inscriptions(db, annee)}
    par_nom = {(_norm(r["nom"]), r["date_naissance"]): r for r in par_mat.values()}
    vus_mat, vus_nom, vus_desps = {}, {}, {}
    for l in an.lignes:
        d = l.donnees
        if d.get("classe") and annee.classe(d["classe"]) is None:
            l.erreurs.append(f"Classe inconnue : {d['classe']}.")
        try:
            repo.validate_eleve(d, annee)
        except ValidationError as e:
            l.erreurs.append(str(e))
        for k, ok in (("redoublant", ("OUI", "NON")), ("statut", ("PRESENT", "ABANDON", "TRANSFERE"))):
            if d.get(k) and d[k] not in ok:
                l.erreurs.append(f"{k.capitalize()} : {' ou '.join(ok)}.")
        cle_nom = (_norm(d.get("nom")), d.get("date_naissance"))
        mat = d.get("matricule_ecole")
        desps = (d.get("matricule_desps") or "").replace(" ", "") or None
        if mat and mat in vus_mat:
            l.erreurs.append(f"Matricule {mat} déjà présent ligne {vus_mat[mat]}.")
        if desps and desps in vus_desps:
            l.erreurs.append(f"Matricule DESPS {desps} déjà présent ligne {vus_desps[desps]}.")
        if d.get("nom") and cle_nom in vus_nom:
            msg = f"même nom{' et date de naissance' if cle_nom[1] else ''} que la ligne {vus_nom[cle_nom]}"
            if cle_nom[1]:
                l.erreurs.append(f"Doublon : {msg}.")
            else:
                l.avertissements.append(f"Doublon possible : {msg} (date de naissance absente).")
        if mat:
            vus_mat.setdefault(mat, l.numero)
        if desps:
            vus_desps.setdefault(desps, l.numero)
        vus_nom.setdefault(cle_nom, l.numero)
        autre = desps and db.scalar("SELECT id FROM eleve WHERE matricule_desps=?", (desps,))
        if mat and mat in par_mat:
            eid = par_mat[mat]["eleve_id"]
            l.action, l.cible = "modifier", {"inscription_id": par_mat[mat]["inscription_id"], "eleve_id": eid}
            if autre and autre != eid:
                l.erreurs.append(f"Matricule DESPS {desps} déjà attribué à un autre élève.")
            continue
        if not mat and cle_nom in par_nom:
            r = par_nom[cle_nom]
            msg = (f"Élève déjà inscrit ({r['matricule_ecole']}, même nom et date de naissance) : "
                   "indiquez son matricule école pour le modifier.")
            if cle_nom[1]:
                l.erreurs.append(msg)
                continue
            l.avertissements.append(f"Un élève du même nom est déjà inscrit ({r['matricule_ecole']}) : "
                                    "vérifiez qu'il ne s'agit pas d'un doublon.")
        if autre:
            ins = db.scalar("SELECT matricule_ecole FROM inscription WHERE eleve_id=? AND annee_id=?",
                            (autre, annee.id))
            if ins:
                l.erreurs.append(f"Matricule DESPS déjà inscrit cette année ({ins}).")
            else:
                l.action, l.cible = "inscrire", {"eleve_id": autre}
                l.avertissements.append("Élève déjà connu (même matricule DESPS) : sa fiche sera mise à jour et "
                                        "il sera inscrit pour l'année.")
            continue
        if mat and db.scalar("SELECT 1 FROM inscription WHERE annee_id=? AND matricule_ecole=?", (annee.id, mat)):
            l.erreurs.append(f"Matricule école {mat} déjà utilisé.")
        l.action = "creer"


def _an_personnel(db, annee, an):
    existants = {p["matricule"]: p for p in repo.personnel(db, annee, actifs=False)}
    vus = {}
    for l in an.lignes:
        d = l.donnees
        m = d.get("matricule")
        if m in vus:
            l.erreurs.append(f"Matricule {m} déjà présent ligne {vus[m]}.")
        vus.setdefault(m, l.numero)
        if d.get("sexe") and d["sexe"] not in ("M", "F"):
            l.erreurs.append("Sexe : M ou F.")
        ct = d.get("classe_tenue")
        if ct and ct != "-" and annee.classe(ct) is None:
            l.erreurs.append(f"Classe tenue inconnue : {ct}.")
        l.action = "modifier" if m in existants else "creer"
        if m in existants:
            l.cible = {"personnel_id": existants[m]["id"]}


def _an_classes(db, annee, an):
    grilles = {g.code: g.id for g in annee.grilles.values()}
    codes = {l.donnees.get("code", "").upper() for l in an.lignes if l.donnees.get("code")} | set(annee.codes)
    vus = {}
    for l in an.lignes:
        d = l.donnees
        code = (d.get("code") or "").upper()
        d["code"] = code
        if code in vus:
            l.erreurs.append(f"Classe {code} déjà présente ligne {vus[code]}.")
        vus.setdefault(code, l.numero)
        if d.get("grille") and d["grille"] not in grilles:
            l.erreurs.append(f"Grille de notes inconnue : {d['grille']}.")
        if d.get("seuil") is not None and d.get("bareme") and d["seuil"] > d["bareme"]:
            l.erreurs.append("Le seuil dépasse le barème.")
        if d.get("niveau_suivant") and d["niveau_suivant"].upper() not in codes:
            l.erreurs.append(f"Classe suivante inconnue : {d['niveau_suivant']}.")
        cl = annee.classe(code)
        l.action = "modifier" if cl else "creer"
        l.cible = {"classe_id": cl.id if cl else None, "grille_id": grilles.get(d.get("grille"))}


def _an_notes(db, annee, an):
    for code in {l.cible["classe"] for l in an.lignes}:
        ins = {r["matricule_ecole"]: r for r in repo.inscriptions(db, annee, code)}
        cl = annee.classe(code)
        deja = {r["inscription_id"]: r["notes"] for r in
                repo.notes_grille(db, annee, cl.grille_id, an.evaluation, code)[1]} if cl and cl.grille_id in annee.grilles else {}
        vus = {}
        for l in (x for x in an.lignes if x.cible["classe"] == code):
            d = l.donnees
            m = d.get("matricule_ecole")
            if m in vus:
                l.erreurs.append(f"Matricule {m} déjà présent ligne {vus[m]}.")
            vus.setdefault(m, l.numero)
            if m not in ins:
                l.erreurs.append(f"Matricule {m} inconnu en {code}.")
                continue
            l.cible["inscription_id"] = ins[m]["inscription_id"]
            notes = {k: v for k, v in d.items() if k.startswith("m:") and v is not None}
            absent = d.get("present") == "NON" or any(v == "ABS" for v in notes.values())
            if d.get("present") not in (None, "OUI", "NON"):
                l.erreurs.append("Présent : OUI ou NON.")
            if absent and any(v != "ABS" for v in notes.values()):
                l.erreurs.append("Élève noté absent mais des notes sont saisies.")
            if not notes and d.get("present") is None:
                l.action = "ignorer"
                l.avertissements.append("Aucune note saisie.")
            else:
                l.action = "modifier"
                avant = deja.get(l.cible["inscription_id"], {})
                n = len(avant) if absent else sum(1 for k, v in notes.items()
                                                  if avant.get(int(k[2:])) not in (None, v))
                if n:
                    l.cible["corrections"] = n
                    l.avertissements.append(f"{n} note(s) déjà saisie(s) corrigée(s) : mot de passe "
                                            "administrateur demandé.")


def _an_absences(db, annee, an):
    eleve = an.kind == "absences_eleves"
    if eleve:
        cibles = {r["matricule_ecole"]: r["inscription_id"] for r in repo.inscriptions(db, annee)}
        deja = {(r["inscription_id"], r["date"], r["nature"]) for r in repo.absences_eleves(db, annee)}
        cle = "matricule_ecole"
    else:
        cibles = {p["matricule"]: p["id"] for p in repo.personnel(db, annee, actifs=False)}
        deja = {(r["personnel_id"], r["date"], r["nature"]) for r in repo.absences_personnel(db, annee)}
        cle = "matricule"
    vus = {}
    for l in an.lignes:
        d = l.donnees
        cid = cibles.get(d.get(cle))
        if d.get(cle) and cid is None:
            l.erreurs.append(f"Matricule inconnu : {d[cle]}.")
        try:
            repo._validate_absence(d, annee)
        except ValidationError as e:
            l.erreurs.append(str(e))
        k = (cid, d.get("date"), d.get("nature"))
        if cid and k in vus:
            l.erreurs.append(f"Même absence que la ligne {vus[k]}.")
        vus.setdefault(k, l.numero)
        l.cible = {"inscription_id" if eleve else "personnel_id": cid}
        if k in deja:
            l.action = "ignorer"
            l.avertissements.append("Déjà enregistrée (même personne, même date, même nature).")
        else:
            l.action = "creer"


# =================================================================== application
def appliquer(db: Database, s: Session, annee: Annee, an: Analyse) -> dict:
    """Enregistre les lignes valides (tout ou rien en cas d'incident), après une sauvegarde."""
    from epp.core import backup
    s.require("import.excel")
    if an.erreurs_fichier:
        raise ValidationError(" ".join(an.erreurs_fichier))
    valides = [l for l in an.lignes if l.valide]
    bilan = {"enregistrees": 0, "ignorees": sum(1 for l in an.lignes if l.action == "ignorer" and not l.erreurs),
             "refusees": sum(1 for l in an.lignes if l.erreurs), "sauvegarde": None}
    if not valides:
        return bilan
    bilan["sauvegarde"] = backup.sauvegarde_auto(db, f"avant import {an.kind}", login=s.login)
    fn = {"eleves": _ap_eleve, "personnel": _ap_personnel, "classes": _ap_classe, "notes": _ap_notes,
          "absences_eleves": _ap_absence, "absences_personnel": _ap_absence}[an.kind]
    with db.transaction():
        for l in valides:
            try:
                fn(db, s, annee, an, l)
            except ValidationError as e:
                raise ValidationError(f"Ligne {l.numero} ({l.feuille}) : {e}. Rien n'a été importé.") from None
            bilan["enregistrees"] += 1
    log(db, s.login, "IMPORT_MODELE_EXCEL", an.kind, f"{bilan['enregistrees']} ligne(s)")
    return bilan


def _ap_eleve(db, s, annee, an, l):
    d = dict(l.donnees)
    if l.action == "modifier":
        ancien = repo.get_eleve(db, l.cible["eleve_id"])
        fiche = {k: d[k] if d.get(k) not in (None, "") else ancien.get(k) for k in repo.ELEVE_CHAMPS}
        repo.save_eleve(db, s, l.cible["eleve_id"], fiche, annee)
        ins = db.one("SELECT * FROM inscription WHERE id=?", (l.cible["inscription_id"],))
        repo.update_inscription(db, s, ins["id"], annee, d.get("classe") or annee.classe_par_id(ins["classe_id"]).code,
                                d.get("redoublant") or ins["redoublant"], d.get("statut") or ins["statut"],
                                ins["observations"])
        return
    if l.action == "inscrire":
        ancien = repo.get_eleve(db, l.cible["eleve_id"])
        fiche = {k: d[k] if d.get(k) not in (None, "") else ancien.get(k) for k in repo.ELEVE_CHAMPS}
        eid = repo.save_eleve(db, s, l.cible["eleve_id"], fiche, annee)
        typ = "REINSCRIPTION"
    else:
        eid = repo.save_eleve(db, s, None, d, annee)
        typ = "INSCRIPTION"
    repo.inscrire(db, s, eid, annee, d["classe"], d.get("redoublant") or "NON", d.get("statut") or "PRESENT", typ,
                  matricule=d.get("matricule_ecole") or None)


def _ap_personnel(db, s, annee, an, l):
    d = dict(l.donnees)
    pid = l.cible.get("personnel_id")
    if pid:
        ancien = dict(db.one("SELECT * FROM personnel WHERE id=?", (pid,)))
        d = {k: d[k] if d.get(k) not in (None, "") else ancien.get(k) for k in repo.PERSONNEL_CHAMPS} | (
            {"classe_tenue": d["classe_tenue"]} if d.get("classe_tenue") else {})
    if not d.get("classe_tenue"):
        d.pop("classe_tenue", None)
    repo.save_personnel(db, s, pid, d, annee)


def _ap_classe(db, s, annee, an, l):
    d = dict(l.donnees)
    cl = annee.classe(d["code"])
    data = {"code": d["code"], "libelle": d.get("libelle") or (cl.libelle if cl else d["code"]),
            "ordre": d.get("ordre") if d.get("ordre") is not None else (cl.ordre if cl else 0),
            "niveau_suivant": (d.get("niveau_suivant") or "").upper() or None, "seuil": d["seuil"],
            "bareme": d["bareme"], "grille_id": l.cible["grille_id"],
            "calendrier": d.get("calendrier") or (cl.calendrier if cl else "STD"),
            "regle_mga": d.get("regle_mga") or (cl.regle_mga if cl else "STANDARD")}
    repo.save_classe(db, s, l.cible["classe_id"], annee.id, data)


def _ap_notes(db, s, annee, an, l):
    iid = l.cible["inscription_id"]
    cl = annee.classe(l.cible["classe"])
    notes = {k: v for k, v in l.donnees.items() if k.startswith("m:") and v is not None}
    absent = l.donnees.get("present") == "NON" or any(v == "ABS" for v in notes.values())
    deja_present = bool(db.scalar("SELECT present FROM presence_eval WHERE inscription_id=? AND eval_num=?",
                                  (iid, an.evaluation), 1))
    if absent or not deja_present:
        repo.set_presence(db, s, iid, an.evaluation, not absent, cl.grille_id)
    if not absent:
        for k, v in notes.items():
            repo.set_note(db, s, annee, iid, int(k[2:]), v)


def _ap_absence(db, s, annee, an, l):
    kind = "eleve" if an.kind == "absences_eleves" else "personnel"
    repo.save_absence(db, s, kind, None, annee, l.donnees | l.cible)


def ecrire_rapport(an: Analyse, chemin: str | Path) -> Path:
    """Rapport d'analyse : une ligne par ligne du fichier, avec l'action prévue et les erreurs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport d'import"
    ws.append(["Feuille", "Ligne", "Action prévue", "Erreurs", "Remarques"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), ENTETE
    lib = {"creer": "Création", "modifier": "Mise à jour", "inscrire": "Inscription", "ignorer": "Ignorée"}
    for e in an.erreurs_fichier:
        ws.append(["", "", "Fichier refusé", e, ""])
    for l in an.lignes:
        ws.append([l.feuille, l.numero, "Refusée" if l.erreurs else lib.get(l.action, l.action),
                   " ; ".join(l.erreurs), " ; ".join(l.avertissements)])
        if l.erreurs:
            for c in ws[ws.max_row]:
                c.fill = ERREUR
    for col, w in zip("ABCDE", (16, 8, 14, 70, 50)):
        ws.column_dimensions[col].width = w
    chemin = Path(chemin)
    wb.save(chemin)
    return chemin
